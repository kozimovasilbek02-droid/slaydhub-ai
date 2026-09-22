from pptx import Presentation

prs = Presentation(r'output\academic_studio\Kvant_kompyuterlari_va_asimmetrik_kripto_Customized.pptx')
with open('scratch_all_slides.txt', 'w', encoding='utf-8') as out:
    for idx, s in enumerate(prs.slides):
        out.write(f"\n==================== SLIDE {idx+1} ({s.slide_layout.name}) ====================\n")
        for i, shp in enumerate(s.shapes):
            l = round(shp.left / 914400 * 72, 1)
            t = round(shp.top / 914400 * 72, 1)
            w = round(shp.width / 914400 * 72, 1)
            h = round(shp.height / 914400 * 72, 1)
            txt = shp.text_frame.text if shp.has_text_frame else "NO TEXT"
            out.write(f"Shape {i}: {shp.name} (type={shp.shape_type}) pos=({l},{t}) size=({w},{h})\n")
            out.write(f"   TEXT: {repr(txt)}\n")

print("Wrote scratch_all_slides.txt successfully")
