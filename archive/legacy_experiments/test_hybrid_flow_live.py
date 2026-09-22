# -*- coding: utf-8 -*-
"""
End-to-end test script for the Content-First Hybrid Workflow.
Simulates:
1. User provides Topic and 10-slide NotebookLM Markdown outline.
2. NotebookLMMarkdownParser parses 10 slides and infers optimal archetypes.
3. AcademicMatcher / Catalog assigns matching native PPTX template slides.
4. UniversalSlideSplicer / TemplateDeckAssembler builds the PPTX with Title-Theme Harmonization.
5. Verifies resulting PPTX integrity and slide diversity.
"""

import os
import sys
from pptx import Presentation

from core.academic_catalog import get_academic_catalog
from core.academic_matcher import get_academic_matcher
from core.notebooklm_markdown_parser import NotebookLMMarkdownParser
from core.template_pptx_builder import get_academic_pptx_builder
from tests.test_notebooklm_markdown_parser import QUANTUM_PRESENTATION_MARKDOWN


def run_hybrid_workflow_test():
    print("=== STARTING HYBRID CONTENT-FIRST WORKFLOW TEST ===")
    
    topic = "Kvant kompyuterlari va asimmetrik kriptografiya"
    meta = {
        "author_name": "Alisher Vohidov",
        "organization": "O'zbekiston Milliy Universiteti",
        "supervisor": "Ilmiy rahbar: DSc., prof. A. Karimov",
        "email": "alisher@univ.edu.uz | Himoya 2026"
    }

    # Step 1: Parse Markdown Outline
    print("\n[Step 1] Parsing NotebookLM Markdown Outline...")
    parsed_content = NotebookLMMarkdownParser.parse(QUANTUM_PRESENTATION_MARKDOWN, topic=topic, default_metadata=meta)
    slides = parsed_content.get("slides", [])
    print(f"-> Successfully parsed {len(slides)} slides.")
    
    archetypes = [s["layout_type"] for s in slides]
    print(f"-> Inferred Archetypes: {archetypes}")
    unique_archetypes = set(archetypes)
    print(f"-> Unique Archetypes count: {len(unique_archetypes)} / {len(slides)}")
    assert len(slides) == 10, f"Expected 10 slides, got {len(slides)}"
    assert len(unique_archetypes) >= 5, f"Expected at least 5 distinct archetypes, got {len(unique_archetypes)}"

    # Step 2: Build Blueprint and Match Templates
    print("\n[Step 2] Building Dynamic Narrative Blueprint...")
    matcher = get_academic_matcher()
    catalog = get_academic_catalog()
    
    # Build baseline blueprint
    bp = matcher.build_blueprint(
        topic=topic,
        slide_count=len(slides),
        language="uz",
        selected_category="cs_ai_cyber",
        presentation_mode="scientific_defense_bmi",
        presenter_info=meta
    )
    
    # Sync blueprint slide archetypes with parsed content archetypes
    for i, s in enumerate(slides):
        bp["slides"][i]["layout_type"] = s["layout_type"]
        bp["slides"][i]["title_hint"] = s["title"]
        # Match template slide for this archetype
        matches = catalog.search_slides_by_archetype(s["layout_type"], category="cs_ai_cyber", topic=topic, limit=1)
        if matches:
            bp["slides"][i]["matched_slide_ref"] = matches[0]
            bp["slides"][i]["pptx_path"] = matches[0]["pptx_path"]
            bp["slides"][i]["template_title"] = matches[0]["template_title"]
            bp["slides"][i]["template_slide_number"] = matches[0]["slide_index"] + 1

    print(f"-> Primary Deck: {bp['primary_deck_title']}")

    # Step 3: Assemble Presentation with Harmonized Engine
    print("\n[Step 3] Assembling Presentation with Harmonized Engine...")
    builder = get_academic_pptx_builder()
    output_path = builder.create_presentation(parsed_content, bp, engine_mode="harmonized")
    
    print(f"-> Output PPTX generated: {output_path}")
    assert os.path.exists(output_path), "Generated PPTX does not exist!"
    
    # Step 4: Verify Generated PPTX
    prs = Presentation(output_path)
    slide_count = len(prs.slides)
    print(f"-> Verified PPTX slide count: {slide_count}")
    assert slide_count == 10, f"Expected 10 slides in PPTX, got {slide_count}"

    for i, slide in enumerate(prs.slides, 1):
        text_count = sum(1 for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip())
        print(f"   Slide {i:02d}: text shapes={text_count}, total shapes={len(slide.shapes)}")
        assert len(slide.shapes) > 0, f"Slide {i} is empty!"

    print("\n=== HYBRID CONTENT-FIRST WORKFLOW TEST PASSED 100% ===")
    return output_path


if __name__ == "__main__":
    out_file = run_hybrid_workflow_test()
    print(f"Result file: {out_file}")
