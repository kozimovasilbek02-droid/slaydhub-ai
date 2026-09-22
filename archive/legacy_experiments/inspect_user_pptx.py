# -*- coding: utf-8 -*-
import os
import sys
from pptx import Presentation

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

candidates = [
    r"C:\Users\user\Downloads\Kvant_kompyuterlari_va_asimmetrik_kripto_Customized.pptx",
    r"c:\Users\user\Desktop\Antigravity\Power Point\output\academic_studio\Kvant_kompyuterlari_va_asimmetrik_kripto_Customized.pptx"
]

target_file = None
for c in candidates:
    if os.path.exists(c):
        target_file = c
        break

if not target_file:
    print("File not found in candidates, searching Downloads...")
    import glob
    dl_ppts = glob.glob(r"C:\Users\user\Downloads\*Kvant*.pptx")
    if dl_ppts:
        target_file = dl_ppts[0]

if not target_file:
    print("Could not find file!")
    sys.exit(1)

print(f"Inspecting file: {target_file}")
print(f"File size: {os.path.getsize(target_file):,} bytes")

prs = Presentation(target_file)
print(f"Total slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides):
    print(f"\n================ SLIDE {i+1} ================")
    text_shapes = []
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                font_sizes = []
                font_names = []
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size:
                            font_sizes.append(round(r.font.size.pt, 1))
                        if r.font.name:
                            font_names.append(r.font.name)
                text_shapes.append((s_idx, shape, txt, font_sizes, font_names))
    
    print(f"Total text shapes found: {len(text_shapes)}")
    for s_idx, shape, txt, font_sizes, font_names in text_shapes:
        print(f"  [Shape #{s_idx}] Top={shape.top}, Left={shape.left}, W={shape.width}, H={shape.height}")
        print(f"     Fonts={set(font_names)}, Sizes={set(font_sizes)}")
        print(f"     Text={repr(txt)}")
