# -*- coding: utf-8 -*-
import os
import sys
import glob
from pptx import Presentation

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

sn_files = glob.glob(r"G:\Мой диск\Slaydlar\Slidenest\*.pptx")
if not sn_files:
    print("No Slidenest files found")
    sys.exit(0)

sample = sn_files[0]
print(f"Testing template: {os.path.basename(sample)}")

prs = Presentation(sample)
print(f"Total slides in template: {len(prs.slides)}")

slides_list = list(prs.slides)
for i, slide in enumerate(slides_list[:3]):
    print(f"\n--- SLIDE {i+1} ---")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    font_name = p.runs[0].font.name if p.runs and p.runs[0].font.name else "Inherited"
                    font_size = p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else "Auto"
                    print(f"  Shape {s_idx} [Font: {font_name}, Size: {font_size}]: \"{t[:50]}\"")
