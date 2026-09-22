# -*- coding: utf-8 -*-
import os
import sys
from pptx import Presentation
from pptx.util import Inches

from core.academic_matcher import get_academic_matcher
matcher = get_academic_matcher()
bp = matcher.build_blueprint("Kvant kompyuterlari va asimmetrik kriptografiya", slide_count=10, language="uz")
template_path = bp["primary_pptx_path"]

print(f"Template path: {template_path}")
if not os.path.exists(template_path):
    print("Template not found!")
    sys.exit()

prs = Presentation(template_path)
sw = prs.slide_width / Inches(1)
sh = prs.slide_height / Inches(1)
print(f"Total slides in template: {len(prs.slides)}, Dimensions: {sw:.2f} x {sh:.2f} inches")

for s_idx in range(min(12, len(prs.slides))):
    slide = prs.slides[s_idx]
    print(f"\n==================== SLIDE {s_idx+1} ====================")
    for sh_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.replace("\n", " ").strip()
            if not txt:
                continue
            fs_list = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
            max_fs = max(fs_list) if fs_list else 0
            l = shape.left / Inches(1)
            t = shape.top / Inches(1)
            w = shape.width / Inches(1)
            h = shape.height / Inches(1)
            print(f"  Shape {sh_idx:02d} [{shape.name}]: pos=({l:.2f}\", {t:.2f}\") size=({w:.2f}\" x {h:.2f}\") max_fs={max_fs:.1f}pt | text: {repr(txt[:65])}")
