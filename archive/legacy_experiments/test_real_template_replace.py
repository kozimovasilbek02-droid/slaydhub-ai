# -*- coding: utf-8 -*-
import os
import sys
import glob
from pptx import Presentation

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from core.pptx_text_replacer import replace_presentation_content

# Find a real template PPTX
sn_files = glob.glob(r"G:\Мой диск\Slaydlar\Slidenest\*.pptx")
if not sn_files:
    print("No template files found")
    sys.exit(1)

sample_template = sn_files[0]
print(f"Testing authentic template: {os.path.basename(sample_template)}")

sample_content = {
    "topic": "Sun'iy intellekt va Neyron tarmoqlari",
    "slides": [
        {
            "slide_number": 1,
            "layout_type": "cover",
            "title": "Sun'iy intellekt va Neyron tarmoqlari",
            "subtitle": "Zamonaviy chuqur o'rganish arxitekturasi va modellari | 2026-yil"
        },
        {
            "slide_number": 2,
            "layout_type": "agenda",
            "points": [
                "1. Neyron tarmoqlari nazariyasi",
                "2. Chuqur arxitektura va transformatorlar",
                "3. Amaliy natijalar va xulosalar",
                "4. Kelajak istiqbollari"
            ]
        },
        {
            "slide_number": 3,
            "layout_type": "theory_concept",
            "title": "Chuqur O'rganish Asoslari",
            "concept_name": "Deep Neural Networks",
            "definition": "Ko'p qatlamli neyron tarmoqlari yordamida murakkab qonuniyatlarni o'rganish texnologiyasi.",
            "key_formula_or_law": "y = sigma(W*x + b)",
            "importance": "Generativ AI va kompyuter ko'rishining asosi."
        }
    ]
}

out_file = os.path.abspath(r"output\academic_studio\test_customized_native.pptx")
replace_presentation_content(sample_template, sample_content, out_file)

print(f"✅ SUCCESS! Created customized presentation at: {out_file}")
print(f"File size: {os.path.getsize(out_file):,} bytes")

prs = Presentation(out_file)
print(f"Total slides in customized PPTX: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    texts = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    print(f"\n--- SLIDE {i+1} ({len(texts)} texts) ---")
    for t in texts:
        print(f"  * {t[:60]}")
