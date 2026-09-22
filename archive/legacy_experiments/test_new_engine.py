# -*- coding: utf-8 -*-
import os
import sys
import re
import math
from typing import Dict, Any, List, Tuple, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from core.academic_matcher import get_academic_matcher

matcher = get_academic_matcher()
bp = matcher.build_blueprint("Kvant kompyuterlari va asimmetrik kriptografiya", slide_count=10, language="uz")
template_path = bp["primary_pptx_path"]

print(f"Loaded Template: {template_path}")
prs = Presentation(template_path)
sw = prs.slide_width
sh = prs.slide_height

print(f"Slide Dimensions: {sw/Inches(1):.2f}\" x {sh/Inches(1):.2f}\" ({sw.pt:.1f} x {sh.pt:.1f} pt)")

# Let's inspect slide classification for each template slide
def classify_slide_detailed(slide, s_idx: int) -> Dict[str, Any]:
    text_shapes = []
    has_step = False
    has_timeline = False
    has_comparison = False
    has_agenda = False
    has_credits = False
    has_table = False
    
    title_candidates = []
    card_boxes = []
    
    slide_w_pt = sw.pt
    slide_h_pt = sh.pt
    
    for s_idx_shape, s in enumerate(slide.shapes):
        if not s.has_text_frame:
            continue
        t = s.text_frame.text.strip()
        if not t:
            continue
            
        t_low = t.lower()
        if any(k in t_low for k in ["table of contents", "agenda", "contents", "index"]):
            has_agenda = True
        if any(k in t_low for k in ["step 01", "step 1", "01.", "step infographic", "steps infographic"]):
            has_step = True
        if any(k in t_low for k in ["timeline", "2026", "2027", "2028", "2029", "2030", "roadmap"]):
            has_timeline = True
        if any(k in t_low for k in ["comparison", "vs", "versus"]):
            has_comparison = True
        if any(k in t_low for k in ["table slide"]):
            has_table = True
        if any(k in t_low for k in ["credits", "fonts", "images", "let’s keep in touch", "contact", "thank you", "thanks"]):
            has_credits = True
            
        # Check if shape is in top area (potential slide title)
        is_top = s.top < (sh * 0.35)
        text_shapes.append({
            "shape": s,
            "text": t,
            "left": s.left,
            "top": s.top,
            "width": s.width,
            "height": s.height,
            "is_top": is_top,
            "is_step": bool(re.match(r'^(0?[1-9]|1[0-9])[\.\:\)]?$', t) or re.match(r'^step\s*0?[1-9]$', t, re.I) or re.match(r'^[A-D][\.\:\)]?$', t)),
            "font_size": max([r.font.size.pt for p in s.text_frame.paragraphs for r in p.runs if r.font.size] or [14.0])
        })
        
    layout_type = "general"
    if s_idx == 0:
        layout_type = "cover"
    elif has_table:
        layout_type = "table_data"
    elif has_credits:
        layout_type = "conclusion_sources"
    elif has_agenda:
        layout_type = "agenda"
    elif has_timeline:
        layout_type = "timeline_steps"
    elif has_step:
        layout_type = "timeline_steps"
    elif has_comparison:
        layout_type = "comparison_vs"
    elif len(text_shapes) >= 8:
        layout_type = "cards_grid"
    elif len(text_shapes) in [2, 3, 4] and any(s["top"] > sh * 0.4 for s in text_shapes):
        layout_type = "theory_concept"
    else:
        layout_type = "cards_grid"
        
    return {
        "slide_index": s_idx,
        "layout_type": layout_type,
        "text_shape_count": len(text_shapes),
        "text_shapes": text_shapes
    }

classified = [classify_slide_detailed(slide, i) for i, slide in enumerate(prs.slides)]

for c in classified:
    print(f"Slide {c['slide_index']+1:02d}: {c['layout_type']:<20} ({c['text_shape_count']} text shapes)")
