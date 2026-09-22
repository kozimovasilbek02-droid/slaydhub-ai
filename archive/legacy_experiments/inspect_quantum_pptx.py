# -*- coding: utf-8 -*-
import os
import sys
from pptx import Presentation

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

q_path = r"G:\Мой диск\Slaydlar\Slidenest\Quantum Computing Presentation.pptx"
if not os.path.exists(q_path):
    print("Not found")
    sys.exit(1)

prs = Presentation(q_path)
print(f"Total slides: {len(prs.slides)}")

for i, slide in enumerate(list(prs.slides)[:6]):
    print(f"\n--- SLIDE {i+1} ({len(slide.shapes)} shapes) ---")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                print(f"  [#{s_idx}] Top={shape.top}, Left={shape.left}, W={shape.width}, H={shape.height}")
                print(f"      Text: {repr(txt[:60])}")
