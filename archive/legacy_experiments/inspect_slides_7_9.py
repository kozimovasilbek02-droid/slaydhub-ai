# -*- coding: utf-8 -*-
import os
import sys
from pptx import Presentation
from pptx.util import Inches

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from core.academic_matcher import get_academic_matcher
matcher = get_academic_matcher()
bp = matcher.build_blueprint("Kvant kompyuterlari va asimmetrik kriptografiya", slide_count=10, language="uz")
template_path = bp["primary_pptx_path"]
prs = Presentation(template_path)

for s_idx in [12]:
    slide = prs.slides[s_idx]
    print(f"\n=======================================================")
    print(f"ORIGINAL TEMPLATE SLIDE {s_idx+1} ({len(slide.shapes)} total shapes)")
    print(f"=======================================================")
    for sh_idx, shape in enumerate(slide.shapes):
        l = shape.left / Inches(1)
        t = shape.top / Inches(1)
        w = shape.width / Inches(1)
        h = shape.height / Inches(1)
        has_tf = shape.has_text_frame
        txt = shape.text_frame.text.replace("\n", " \\n ").strip() if has_tf else ""
        shp_type = shape.shape_type
        print(f"Shape {sh_idx:02d} [{shape.name}]: pos=({l:.2f}\", {t:.2f}\") size=({w:.2f}\" x {h:.2f}\") | text: {repr(txt[:75])}")
