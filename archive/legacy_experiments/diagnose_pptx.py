import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Inches

pptx_path = r"C:\Users\user\Downloads\Kvant_kompyuterlari_va_asimmetrik_kripto_Customized.pptx"
if not os.path.exists(pptx_path):
    print("File not found at", pptx_path)
    sys.exit()

prs = Presentation(pptx_path)
sw = prs.slide_width / Inches(1)
sh = prs.slide_height / Inches(1)
print(f"Slide dimensions: {sw:.2f} x {sh:.2f} inches ({sw*72:.1f} x {sh*72:.1f} pt)")

for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx+1} ====================")
    text_shapes = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    print(f"Total non-empty text shapes: {len(text_shapes)}")
    
    for s_idx, shape in enumerate(slide.shapes):
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue
        l = shape.left / Inches(1)
        t = shape.top / Inches(1)
        w = shape.width / Inches(1)
        h = shape.height / Inches(1)
        r = l + w
        b = t + h
        
        # Calculate theoretical text height at current font size
        fs_list = [run.font.size.pt for p in shape.text_frame.paragraphs for run in p.runs if run.font.size]
        avg_fs = sum(fs_list)/len(fs_list) if fs_list else 12.0
        
        # Estimate rendered lines
        # Usable width in points:
        usable_w_pt = max(10, (w * 72) - 8)
        char_w = avg_fs * 0.51
        chars_per_line = max(1, int(usable_w_pt / char_w))
        total_rendered_lines = sum(max(1, (len(line) + chars_per_line - 1) // chars_per_line) for line in text.splitlines() if line.strip())
        estimated_text_height_pt = total_rendered_lines * (avg_fs * 1.25)
        box_height_pt = h * 72
        
        overflow_ratio = estimated_text_height_pt / max(1, box_height_pt)
        is_overflowing_box = overflow_ratio > 1.05
        is_outside_slide = (r > sw + 0.1) or (b > sh + 0.1) or (l < -0.1) or (t < -0.1)
        
        status = []
        if is_overflowing_box:
            status.append(f"BOX OVERFLOW ({estimated_text_height_pt:.1f}pt needed vs {box_height_pt:.1f}pt box, ratio={overflow_ratio:.2f}x)")
        if is_outside_slide:
            status.append(f"OUTSIDE SLIDE (right={r:.2f}, bottom={b:.2f})")
            
        status_str = " -> " + " | ".join(status) if status else " -> OK"
        
        print(f" Shape {s_idx:02d} [{shape.name}]: pos=({l:.2f}\", {t:.2f}\") size=({w:.2f}\" x {h:.2f}\") font={avg_fs:.1f}pt lines={total_rendered_lines}{status_str}")
        for l_num, line in enumerate(text.splitlines()[:4]):
            print(f"    [{l_num+1}] {line[:90]}")
        if len(text.splitlines()) > 4:
            print(f"    ... + {len(text.splitlines()) - 4} more lines (total {len(text)} chars)")
