# -*- coding: utf-8 -*-
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Inches

from core.academic_matcher import get_academic_matcher
matcher = get_academic_matcher()
bp = matcher.build_blueprint("Kvant kompyuterlari va asimmetrik kriptografiya", slide_count=10, language="uz")
template_path = bp["primary_pptx_path"]
prs = Presentation(template_path)

def inspect_slide_detail(slide_idx):
    slide = prs.slides[slide_idx]
    print(f"\n==================== SLIDE {slide_idx+1} DETAIL ====================")
    text_shapes = []
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text_frame.text.replace("\n", " ").strip()
            l = shape.left / Inches(1)
            t = shape.top / Inches(1)
            w = shape.width / Inches(1)
            h = shape.height / Inches(1)
            text_shapes.append((s_idx, shape.name, l, t, w, h, txt))
    
    # Sort by top, then left
    text_shapes.sort(key=lambda x: (x[3], x[2]))
    for s in text_shapes:
        print(f"  Shape {s[0]:02d} [{s[1]}]: pos=({s[2]:.2f}\", {s[3]:.2f}\") size=({s[4]:.2f}\" x {s[5]:.2f}\") | text: {repr(s[6][:50])}")

for idx in [0, 1, 8, 11, 13, 19, 20, 24]:
    inspect_slide_detail(idx)
