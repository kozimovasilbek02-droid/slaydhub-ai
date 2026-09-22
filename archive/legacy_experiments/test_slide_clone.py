# -*- coding: utf-8 -*-
import os
import copy
from pptx import Presentation
from pptx.util import Inches

from core.academic_matcher import get_academic_matcher
matcher = get_academic_matcher()
bp = matcher.build_blueprint("Kvant kompyuterlari va asimmetrik kriptografiya", slide_count=10, language="uz")
template_path = bp["primary_pptx_path"]

prs = Presentation(template_path)
print(f"Loaded template with {len(prs.slides)} slides.")

# Target sequence of slide indices (0-indexed)
chosen_indices = [0, 1, 2, 8, 11, 13, 19, 20, 21, 24] # Cover, Agenda, Concept, Cards(9), Comparison(12), Cards(14), Steps(20), Timeline(21), Map(22), Conclusion(25)
print(f"Chosen template slide indices: {chosen_indices}")

# Let's test how to prune or reorder slides in python-pptx
# If we delete slides that are NOT in chosen_indices:
# Note: when deleting by index, indices shift, so we can mark which slides to keep and delete the rest in reverse order!
to_keep = set(chosen_indices)
to_delete = [i for i in range(len(prs.slides)) if i not in to_keep]

print(f"Deleting {len(to_delete)} unused slides...")
for idx in sorted(to_delete, reverse=True):
    rId = prs.slides._sldIdLst[idx].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[idx]

print(f"Remaining slides count: {len(prs.slides)}")

test_out = r"output\academic_studio\test_pruned.pptx"
prs.save(test_out)
print(f"Saved pruned presentation: {test_out} (size: {os.path.getsize(test_out):,} bytes)")

# Verify that all 10 remaining slides can be opened and have shapes
prs_check = Presentation(test_out)
print(f"Verification: {len(prs_check.slides)} slides preserved successfully!")
for i, sl in enumerate(prs_check.slides):
    t_shapes = len([s for s in sl.shapes if s.has_text_frame and s.text_frame.text.strip()])
    print(f"  Slide {i+1}: {len(sl.shapes)} shapes ({t_shapes} text shapes)")
