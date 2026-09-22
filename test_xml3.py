from pptx import Presentation

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
table = slide.shapes.add_table(2, 2, 0, 0, 100000, 100000).table

tbl = table._tbl
for tc in tbl.xpath('.//a:tc'):
    print("Found tc")
