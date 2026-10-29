# -*- coding: utf-8 -*-
from typing import List, Dict, Any
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _set_east_asia_and_cs_fonts(style, font_name: str = "Times New Roman") -> None:
    """Kirill, Lotin va maxsus belgilar barcha Word versiyalarida Times New Roman bo'lishini ta'minlaydi."""
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)


def _enable_auto_update_fields(doc: Document) -> None:
    """Word faylni ochganda Mundarija (TOC) va sahifa raqamlarini avtomatik yangilash bayrog'i."""
    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def _add_page_number_to_footer(footer_para) -> None:
    """Pastki kolontitulga avtomatik OpenXML PAGE maydonini qo'shadi."""
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_para.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")

    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)


def _remove_table_borders(table) -> None:
    """Titul varag'idagi Bajardi/Qabul qildi jadvali chegaralarini ko'rinmas qiladi."""
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "none")
        borders.append(border)
    tblPr.append(borders)


def _insert_native_toc(doc: Document, toc_preview_entries: List[Dict[str, Any]]) -> None:
    """
    Word va WPS Office uchun Native TOC maydoni + oldindan hisoblangan fallback ro'yxat.
    """
    h_para = doc.add_paragraph()
    h_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h_para.paragraph_format.first_line_indent = Mm(0)
    h_para.paragraph_format.space_after = Pt(12)
    h_run = h_para.add_run("MUNDARIJA")
    h_run.bold = True
    h_run.font.name = "Times New Roman"
    h_run.font.size = Pt(14)

    # OpenXML Complex Field: TOC \o "1-3" \h \z \u
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Mm(0)
    run = p.add_run()
    r = run._r

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    fld_begin.set(qn("w:dirty"), "true")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)

    # Fallback preview lines (Mobil Telegram ichida ochilganda ham mundarija to'liq ko'rinishi uchun)
    for item in toc_preview_entries:
        line_run = p.add_run(f"{item['title']}\t{item.get('approx_page', '')}\n")
        line_run.font.name = "Times New Roman"
        line_run.font.size = Pt(14)

    end_run = p.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_end)

    doc.add_page_break()


def build_otm_academic_docx(spec: Dict[str, Any], out_path: str) -> str:
    """
    O'zbekiston OTM standartlariga 100% mos Referat, Mustaqil ish, Kurs ishi va IMRAD maqola generatori.
    """
    doc = Document()
    _enable_auto_update_fields(doc)

    # 1. Sahifa o'lchami va GOST/OTM hoshiyalari (Chap 3sm, O'ng 1.5sm, Yuqori/Past 2sm)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.different_first_page_header_footer = True  # Titul varag'ida sahifa raqami bo'lmaydi

    # 2-betdan boshlab pastki sahifa raqami
    _add_page_number_to_footer(section.footer.paragraphs[0])

    # 2. Asosiy Normal stili (Times New Roman, 14pt, 1.5 interval, 1.25sm abzas)
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    _set_east_asia_and_cs_fonts(normal_style, "Times New Roman")

    pf = normal_style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Mm(12.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 3. Rasmiy OTM Titul Varag'i
    meta = spec.get("meta", {})
    ministry = meta.get(
        "ministry",
        "O'ZBEKISTON RESPUBLIKASI OLIY TA'LIM, FAN VA INNOVATSIYALAR VAZIRLIGI",
    )
    university = meta.get("university", "TOSHKENT DAVLAT UNIVERSITETI")
    faculty_dept = meta.get("department", "«Axborot texnologiyalari» kafedrasi")
    doc_type = meta.get("doc_type", "MUSTAQIL ISH").upper()
    topic = spec.get("title", "Ilmiy-tadqiqot mavzusi")
    student_name = meta.get("student", "Talaba F.I.Sh.")
    group_name = meta.get("group", "Guruh: ___")
    advisor_name = meta.get("advisor", "Ilmiy rahbar: dots. ___________")
    city_year = meta.get("city_year", "Toshkent – 2026")

    for header_line in (ministry, university, faculty_dept):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(header_line)
        r.bold = True
        r.font.size = Pt(14)

    # Markaziy blok: Hujjat turi va Mavzu
    spacer1 = doc.add_paragraph()
    spacer1.paragraph_format.space_before = Pt(72)

    p_type = doc.add_paragraph()
    p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_type.paragraph_format.first_line_indent = Mm(0)
    r_type = p_type.add_run(doc_type)
    r_type.bold = True
    r_type.font.size = Pt(20)

    p_topic = doc.add_paragraph()
    p_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_topic.paragraph_format.first_line_indent = Mm(0)
    p_topic.paragraph_format.space_before = Pt(12)
    p_topic.paragraph_format.space_after = Pt(72)
    r_topic = p_topic.add_run(f"Mavzu: «{topic}»")
    r_topic.bold = True
    r_topic.font.size = Pt(16)

    # Bajardi / Qabul qildi (Ko'rinmas jadval yordamida o'ng tomonda tekis blok)
    sign_table = doc.add_table(rows=2, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _remove_table_borders(sign_table)

    right_cell_1 = sign_table.cell(0, 1).paragraphs[0]
    right_cell_1.paragraph_format.first_line_indent = Mm(0)
    right_cell_1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_s1 = right_cell_1.add_run(f"Bajardi: {group_name} talabasi\n{student_name}")
    r_s1.font.size = Pt(14)

    right_cell_2 = sign_table.cell(1, 1).paragraphs[0]
    right_cell_2.paragraph_format.first_line_indent = Mm(0)
    right_cell_2.paragraph_format.space_before = Pt(14)
    right_cell_2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_s2 = right_cell_2.add_run(f"Qabul qildi:\n{advisor_name}")
    r_s2.font.size = Pt(14)

    # Shahar va yil
    p_city = doc.add_paragraph()
    p_city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_city.paragraph_format.first_line_indent = Mm(0)
    p_city.paragraph_format.space_before = Pt(96)
    r_city = p_city.add_run(city_year)
    r_city.bold = True
    r_city.font.size = Pt(14)

    doc.add_page_break()

    # 4. Mundarija (2-sahifa)
    chapters = spec.get("chapters", [])
    toc_preview = [{"title": "KIRISH", "approx_page": 3}]
    running_page = 4
    for idx, ch in enumerate(chapters, 1):
        toc_preview.append({"title": f"{idx}. {ch.get('heading', '')}", "approx_page": running_page})
        running_page += max(2, len(ch.get("paragraphs", [])) // 3)
    toc_preview.append({"title": "XULOSA VA TAKLIFLAR", "approx_page": running_page})
    toc_preview.append({"title": "FOYDALANILGAN ADABIYOTLAR RO'YXATI", "approx_page": running_page + 1})

    _insert_native_toc(doc, toc_preview)

    # 5. Kirish, Asosiy Boblar va Xulosa
    def _write_section_block(heading_text: str, paragraphs: List[str], page_break_after: bool = True):
        h = doc.add_heading(level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.first_line_indent = Mm(0)
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(12)
        run_h = h.add_run(heading_text.upper())
        run_h.bold = True
        run_h.font.name = "Times New Roman"
        run_h.font.size = Pt(14)
        run_h.font.color.rgb = RGBColor(0, 0, 0)

        for para_txt in paragraphs:
            if not str(para_txt).strip():
                continue
            p_body = doc.add_paragraph(str(para_txt).strip())
            p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_body.paragraph_format.first_line_indent = Mm(12.5)

        if page_break_after:
            doc.add_page_break()

    _write_section_block("KIRISH", spec.get("introduction", []), page_break_after=True)

    for idx, ch in enumerate(chapters, 1):
        _write_section_block(
            f"{idx}. {ch.get('heading', f'BOB {idx}')}",
            ch.get("paragraphs", []),
            page_break_after=True,
        )

    _write_section_block("XULOSA VA TAKLIFLAR", spec.get("conclusion", []), page_break_after=True)

    # 6. Foydalanilgan Adabiyotlar Ro'yxati (GOST 7.1-2003 / OTM)
    references = spec.get("references", [])
    ref_lines = [f"{i}. {ref}" for i, ref in enumerate(references, 1)]
    _write_section_block("FOYDALANILGAN ADABIYOTLAR RO'YXATI", ref_lines, page_break_after=False)

    doc.save(out_path)
    return out_path
