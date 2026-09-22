from pptx.oxml.xmlchemy import OxmlElement
from pptx import Presentation

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
table = slide.shapes.add_table(2, 2, 0, 0, 100000, 100000)

tbl = table._element.graphic.graphicData.tbl
for tc in tbl.xpath('.//a:tc', namespaces={'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}):
    print("Found tc")

