# -*- coding: utf-8 -*-
"""
Academic PPTX Presentation Builder (Native Template Customizer)
Customizes authentic, pre-designed PPTX templates in-place, preserving
all 3D graphics, vector artwork, background cards, and typography.
"""

import os
import re
import logging
from pathlib import Path
from typing import Dict, Any, List, Optional
from pptx import Presentation
from pptx.util import Inches, Pt

from core.pptx_text_replacer import replace_presentation_content, PPTXTextReplacer

from core.config import config

logger = logging.getLogger(__name__)

OUTPUT_DIR = str(config.OUTPUT_STUDIO_DIR)


class AcademicPPTXBuilder:
    """
    Renders academic slides directly into authentic PPTX templates.
    """

    def __init__(self):
        pass

    def create_presentation(
        self,
        content_data: Dict[str, Any],
        blueprint: Optional[Dict[str, Any]] = None,
        engine_mode: Optional[str] = None
    ) -> str:
        """
        Creates/customizes a presentation from structured JSON content and blueprint,
        strictly preserving authentic 3D graphics, vector artwork, and pre-designed PPTX templates.
        """
        topic = content_data.get("topic", "Akademik Taqdimot")
        safe_topic = re.sub(r'[^a-zA-Z0-9_\-\u0400-\u04FF]', '_', topic)[:40].strip('_')
        out_path = os.path.join(OUTPUT_DIR, f"{safe_topic}_Customized.pptx")

        # 1. Find primary authentic PPTX template from blueprint
        template_pptx = None
        if blueprint and blueprint.get("primary_pptx_path") and os.path.exists(blueprint["primary_pptx_path"]):
            template_pptx = blueprint["primary_pptx_path"]
        elif blueprint and blueprint.get("slides"):
            for s in blueprint["slides"]:
                if s.get("pptx_path") and os.path.exists(s["pptx_path"]):
                    template_pptx = s["pptx_path"]
                    break

        # 2. If not in blueprint, search catalog for authentic PPTX template
        if not template_pptx or not os.path.exists(template_pptx):
            from core.academic_catalog import get_academic_catalog
            catalog = get_academic_catalog()
            matches = catalog.search_templates(topic, limit=5)
            if not matches:
                matches = catalog.templates[:5]
            if matches:
                template_pptx = matches[0]["pptx_path"]

        # 3. Always customize authentic template in-place
        if template_pptx and os.path.exists(template_pptx):
            chosen_mode = engine_mode or (blueprint.get("engine_mode") if blueprint else None)
            logger.info(f"Customizing authentic template in-place: {template_pptx} (engine: {chosen_mode})")
            replace_presentation_content(template_pptx, content_data, out_path, blueprint=blueprint, engine_mode=chosen_mode)
            return out_path

        # 4. Ultimate fallback (only if no template exists at all)
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        slides = content_data.get("slides", [])
        for idx, slide_data in enumerate(slides, start=1):
            blank_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank_layout)
            PPTXTextReplacer.replace_slide_content(slide, slide_data, topic)

        prs.save(out_path)
        return out_path


# Singleton instance
_builder_instance: Optional[AcademicPPTXBuilder] = None

def get_academic_pptx_builder() -> AcademicPPTXBuilder:
    global _builder_instance
    if _builder_instance is None:
        _builder_instance = AcademicPPTXBuilder()
    return _builder_instance
