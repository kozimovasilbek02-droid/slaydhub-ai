import os
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def hex_to_rgb(hex_code):
    h = hex_code.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

def add_header(slide, title_text, subtitle_text=None, align=PP_ALIGN.CENTER, x=0.8, y=0.50, w=11.733):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Montserrat'
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('#222222')
    p.alignment = align
    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.name = 'Segoe UI'
        p2.font.size = Pt(13)
        p2.font.color.rgb = hex_to_rgb('#666666')
        p2.alignment = align

# 1. Folded Tiered Cards (Slide 1 style)
def draw_folded_tiered_cards(slide, cards_data, y_axis=1.95, y_tab=2.45, y_body=3.20, cw=2.15, h_body=3.40):
    # Axis line
    ax = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.80), Inches(y_axis), Inches(11.733), Inches(0.025))
    ax.fill.solid()
    ax.fill.fore_color.rgb = hex_to_rgb('#BDC3C7')
    ax.line.fill.background()
    
    for title, bullets, col_head, col_body, x in cards_data:
        # Stem Connector
        c_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + cw/2 - 0.015), Inches(y_axis + 0.10), Inches(0.03), Inches(y_tab - y_axis - 0.10))
        c_line.fill.solid()
        c_line.fill.fore_color.rgb = hex_to_rgb(col_body)
        c_line.line.fill.background()
        
        # Body
        body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y_body), Inches(cw), Inches(h_body))
        body.fill.solid()
        body.fill.fore_color.rgb = hex_to_rgb(col_body)
        body.line.fill.background()
        tf_b = body.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = Inches(0.18)
        tf_b.margin_right = Inches(0.18)
        tf_b.margin_top = Inches(0.35)
        for i, b_text in enumerate(bullets):
            p_b = tf_b.paragraphs[0] if i == 0 else tf_b.add_paragraph()
            p_b.text = '•  ' + b_text
            p_b.font.name = 'Segoe UI'
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = hex_to_rgb('#FFFFFF')
            p_b.space_after = Pt(8)
            
        # Top Tab
        tab = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y_tab), Inches(cw), Inches(y_body - y_tab + 0.20))
        tab.fill.solid()
        tab.fill.fore_color.rgb = hex_to_rgb(col_head)
        tab.line.fill.background()
        tf_tab = tab.text_frame
        tf_tab.word_wrap = True
        tf_tab.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_t = tf_tab.paragraphs[0]
        p_t.text = title.replace('\n', ' ')
        p_t.font.name = 'Montserrat'
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = hex_to_rgb('#FFFFFF')
        p_t.alignment = PP_ALIGN.CENTER
        
        # Dot on Axis
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + cw/2 - 0.12), Inches(y_axis - 0.12), Inches(0.24), Inches(0.24))
        dot.fill.solid()
        dot.fill.fore_color.rgb = hex_to_rgb(col_body)
        dot.line.fill.background()

# 2. Interlocking Rounded Cards (Slide 4 style)
def draw_interlocking_cards(slide, cards_data, start_x=0.75, card_w=1.96, card_h=1.75, y_card=2.45):
    tab_w = 0.22
    tab_h = 0.46
    for i, (d_txt, sub_d, sub_b, col) in enumerate(cards_data):
        x = start_x + i * card_w
        c_rgb = hex_to_rgb(col)
        
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y_card), Inches(card_w - 0.02), Inches(card_h))
        card.fill.solid()
        card.fill.fore_color.rgb = c_rgb
        card.line.fill.background()
        
        tab_pts = [
            (Inches(x + card_w - 0.03), Inches(y_card + card_h/2 - tab_h/2)),
            (Inches(x + card_w + tab_w), Inches(y_card + card_h/2)),
            (Inches(x + card_w - 0.03), Inches(y_card + card_h/2 + tab_h/2))
        ]
        fb_tab = slide.shapes.build_freeform(tab_pts[0][0], tab_pts[0][1])
        fb_tab.add_line_segments(tab_pts[1:], close=True)
        tab_shape = fb_tab.convert_to_shape()
        tab_shape.fill.solid()
        tab_shape.fill.fore_color.rgb = c_rgb
        tab_shape.line.fill.background()
        
        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_c.margin_left = Inches(0.10 if i == 0 else 0.22)
        tf_c.margin_right = Inches(0.10)
        p1 = tf_c.paragraphs[0]
        p1.text = d_txt
        p1.font.name = 'Montserrat'
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = hex_to_rgb('#FFFFFF')
        p1.alignment = PP_ALIGN.CENTER
        p2 = tf_c.add_paragraph()
        p2.text = sub_d
        p2.font.name = 'Segoe UI'
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = hex_to_rgb('#FFFFFF')
        p2.alignment = PP_ALIGN.CENTER
        
        c_stem = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + card_w/2 - 0.05), Inches(y_card + card_h), Inches(0.10), Inches(0.40))
        c_stem.fill.solid()
        c_stem.fill.fore_color.rgb = c_rgb
        c_stem.line.fill.background()
        
        tb_sub = slide.shapes.add_textbox(Inches(x - 0.05), Inches(y_card + card_h + 0.45), Inches(card_w + 0.10), Inches(1.10))
        tf_sub = tb_sub.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_right = tf_sub.margin_top = tf_sub.margin_bottom = 0
        p_s = tf_sub.paragraphs[0]
        p_s.text = sub_b.replace('\n', ' ')
        p_s.font.name = 'Segoe UI'
        p_s.font.size = Pt(10.5)
        p_s.font.color.rgb = hex_to_rgb('#222222')
        p_s.alignment = PP_ALIGN.CENTER
        
        b_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(6.15), Inches(card_w - 0.02), Inches(0.14))
        b_bar.fill.solid()
        b_bar.fill.fore_color.rgb = c_rgb
        b_bar.line.fill.background()
