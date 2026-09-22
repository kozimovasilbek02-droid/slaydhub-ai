# -*- coding: utf-8 -*-
"""
core/utils/pptx_helpers.py
Reusable PowerPoint shape, font, and layout helper utilities.
"""

from typing import Optional, Union, Tuple
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from core.utils.colors import hex_to_rgb


def safe_set_font(
    run,
    font_name: Optional[str] = None,
    font_size_pt: Optional[float] = None,
    font_color: Optional[Union[str, RGBColor, Tuple[int, int, int]]] = None,
    bold: Optional[bool] = None,
    italic: Optional[bool] = None
) -> None:
    """
    Safely applies font attributes to a PPTX TextRun without crashing if attributes are unset.
    """
    if font_name and hasattr(run.font, "name"):
        run.font.name = font_name
    if font_size_pt and hasattr(run.font, "size"):
        run.font.size = Pt(font_size_pt)
    if bold is not None and hasattr(run.font, "bold"):
        run.font.bold = bold
    if italic is not None and hasattr(run.font, "italic"):
        run.font.italic = italic
    if font_color and hasattr(run.font, "color"):
        rgb = hex_to_rgb(font_color)
        run.font.color.rgb = rgb


def format_text_frame(
    tf,
    margin_in: float = 0.0,
    word_wrap: bool = True
) -> None:
    """
    Sets zero or uniform margins and enables word wrap on a TextFrame.
    """
    tf.word_wrap = word_wrap
    m = Inches(margin_in)
    tf.margin_left = m
    tf.margin_top = m
    tf.margin_right = m
    tf.margin_bottom = m


def add_clean_textbox(
    slide,
    x_in: float,
    y_in: float,
    w_in: float,
    h_in: float,
    text: str = "",
    font_size_pt: float = 14,
    font_color: Optional[Union[str, RGBColor]] = None,
    font_name: Optional[str] = "Calibri",
    bold: bool = False,
    align: Optional[PP_ALIGN] = PP_ALIGN.LEFT,
    margin_in: float = 0.05
):
    """
    Helper to add a cleanly formatted textbox with text in a single call.
    """
    box = slide.shapes.add_textbox(Inches(x_in), Inches(y_in), Inches(w_in), Inches(h_in))
    tf = box.text_frame
    format_text_frame(tf, margin_in=margin_in, word_wrap=True)

    if text:
        p = tf.paragraphs[0]
        p.text = text
        if align:
            p.alignment = align
        if p.runs:
            safe_set_font(p.runs[0], font_name=font_name, font_size_pt=font_size_pt, font_color=font_color, bold=bold)

    return box
