from typing import Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from core.models import (
    PresentationProject, SlideData, ShapeElement, TextBoxElement, 
    Position, ColorRGB, MilestoneNode
)
from core.templates import build_roadmap_timeline, hex_to_rgb

SHAPE_MAPPING = {
    'RECTANGLE': MSO_SHAPE.RECTANGLE,
    'ROUNDED_RECTANGLE': MSO_SHAPE.ROUNDED_RECTANGLE,
    'OVAL': MSO_SHAPE.OVAL,
    'CIRCLE': MSO_SHAPE.OVAL,
    'CHEVRON': MSO_SHAPE.CHEVRON,
    'RIGHT_ARROW': MSO_SHAPE.RIGHT_ARROW,
    'LINE': MSO_SHAPE.RECTANGLE,
    'DONUT': MSO_SHAPE.DONUT,
    'BADGE': MSO_SHAPE.ROUNDED_RECTANGLE,
    'CONNECTOR': MSO_SHAPE.RECTANGLE,
}

ALIGN_MAPPING = {
    'LEFT': PP_ALIGN.LEFT,
    'CENTER': PP_ALIGN.CENTER,
    'RIGHT': PP_ALIGN.RIGHT,
    'JUSTIFY': PP_ALIGN.JUSTIFY,
}

class PPTXGenerator:
    def __init__(self, aspect_ratio: str = '16:9'):
        self.prs = Presentation()
        if aspect_ratio == '16:9':
            self.width_in = 13.333
            self.height_in = 7.5
        else:
            self.width_in = 10.0
            self.height_in = 7.5
        self.prs.slide_width = Inches(self.width_in)
        self.prs.slide_height = Inches(self.height_in)
        self.blank_slide_layout = self.prs.slide_layouts[6]

    def pct_to_inches(self, pos: Position):
        left = Inches((pos.left_pct / 100.0) * self.width_in)
        top = Inches((pos.top_pct / 100.0) * self.height_in)
        width = Inches((pos.width_pct / 100.0) * self.width_in)
        height = Inches((pos.height_pct / 100.0) * self.height_in)
        return left, top, width, height

    def add_background(self, slide, bg_color: str):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), 
            Inches(self.width_in), Inches(self.height_in)
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = hex_to_rgb(bg_color)
        bg.line.fill.background()
        return bg

    def add_custom_shape(self, slide, shape_elem: ShapeElement):
        left, top, width, height = self.pct_to_inches(shape_elem.position)
        st = SHAPE_MAPPING.get(shape_elem.shape_type, MSO_SHAPE.RECTANGLE)
        shape = slide.shapes.add_shape(st, left, top, width, height)

        if shape_elem.fill_color:
            shape.fill.solid()
            shape.fill.fore_color.rgb = hex_to_rgb(shape_elem.fill_color)
        else:
            shape.fill.background()

        if shape_elem.line_color and shape_elem.line_width_pt > 0:
            shape.line.color.rgb = hex_to_rgb(shape_elem.line_color)
            shape.line.width = Pt(shape_elem.line_width_pt)
        else:
            shape.line.fill.background()

        if shape_elem.text:
            tf = shape.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.text = shape_elem.text
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(shape_elem.text_size_pt)
            p.font.bold = shape_elem.text_bold
            p.font.name = 'Segoe UI'
            if shape_elem.text_color:
                p.font.color.rgb = hex_to_rgb(shape_elem.text_color)
        return shape

    def add_custom_textbox(self, slide, tb_elem: TextBoxElement):
        left, top, width, height = self.pct_to_inches(tb_elem.position)
        tx_box = slide.shapes.add_textbox(left, top, width, height)
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        for i, para in enumerate(tb_elem.paragraphs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN_MAPPING.get(para.alignment, PP_ALIGN.LEFT)
            for run in para.runs:
                r = p.add_run()
                r.text = run.text
                r.font.size = Pt(run.font_size_pt)
                r.font.bold = run.font_bold
                r.font.italic = run.font_italic
                r.font.name = run.font_name
                if run.font_color:
                    r.font.color.rgb = hex_to_rgb(run.font_color)
        return tx_box

    def render_slide(self, slide_data: SlideData):
        slide = self.prs.slides.add_slide(self.blank_slide_layout)
        
        if slide_data.background_color:
            self.add_background(slide, slide_data.background_color)

        if slide_data.layout_type == 'ROADMAP_TIMELINE' and (slide_data.milestones or 'Roadmap' in (slide_data.title or '')):
            build_roadmap_timeline(slide, slide_data, self.width_in, self.height_in)
            return slide

        if slide_data.title:
            tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(self.width_in - 1.6), Inches(0.8))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = slide_data.title
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.name = 'Segoe UI'
            p.font.color.rgb = hex_to_rgb('#1E293B')

        for s in slide_data.shapes:
            self.add_custom_shape(slide, s)

        for tb in slide_data.text_boxes:
            self.add_custom_textbox(slide, tb)

        return slide

    def generate_presentation(self, project: PresentationProject, output_path: str):
        for slide_data in project.slides:
            self.render_slide(slide_data)
        self.prs.save(output_path)
        return output_path
