# Core High Precision Slide Builder Module
import os
from PIL import Image
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from core.utils.colors import hex_to_rgb

class SlideBuilder:
    def __init__(self, prs):
        self.prs = prs
        self.charcoal = '#1A252C'
        self.grey_desc = '#556270'

    def create_title(self, slide, text, x=1.0, y=0.40, w=11.333, h=0.80, align=PP_ALIGN.LEFT, font_size=28, bold=False):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = 'Georgia'
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = hex_to_rgb(self.charcoal)
        p.alignment = align
        return tb

    def build_slide_1(self, slide, img, w_px, h_px):
        top_accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0), Inches(1.8), Inches(0.08))
        top_accent.fill.solid()
        top_accent.fill.fore_color.rgb = hex_to_rgb('#009688')
        top_accent.line.fill.background()

        crop = img.crop((int(0.653 * w_px), 0, w_px, h_px))
        p_path = 'temp_s1_photo.png'
        crop.save(p_path)
        slide.shapes.add_picture(p_path, Inches(8.70), Inches(0), width=Inches(4.633), height=Inches(7.50))
        if os.path.exists(p_path): os.remove(p_path)

        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.70), Inches(1.42), Inches(4.633), Inches(4.65))
        card.fill.solid()
        card.fill.fore_color.rgb = hex_to_rgb('#009688')
        card.line.fill.background()
        tf_c = card.text_frame
        tf_c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p1 = tf_c.paragraphs[0]
        p1.text = 'Creative agenda\npowerpoint slides'
        p1.font.name = 'Georgia'
        p1.font.size = Pt(28)
        p1.font.color.rgb = hex_to_rgb('#FFFFFF')
        p1.alignment = PP_ALIGN.CENTER

        l1 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.75), Inches(0), Inches(0.025), Inches(1.95))
        l1.fill.solid()
        l1.fill.fore_color.rgb = hex_to_rgb(self.charcoal)
        l1.line.fill.background()
        d1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.68), Inches(1.88), Inches(0.16), Inches(0.16))
        d1.fill.solid()
        d1.fill.fore_color.rgb = hex_to_rgb(self.charcoal)
        d1.line.fill.background()

        d2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.68), Inches(3.48), Inches(0.16), Inches(0.16))
        d2.fill.solid()
        d2.fill.fore_color.rgb = hex_to_rgb(self.charcoal)
        d2.line.fill.background()
        l2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.75), Inches(3.55), Inches(0.025), Inches(3.95))
        l2.fill.solid()
        l2.fill.fore_color.rgb = hex_to_rgb(self.charcoal)
        l2.line.fill.background()

        chev = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(0), Inches(2.26), Inches(3.92), Inches(1.00))
        chev.fill.solid()
        chev.fill.fore_color.rgb = hex_to_rgb('#009688')
        chev.line.fill.background()
        p_ch = chev.text_frame.paragraphs[0]
        p_ch.text = 'Agenda Slide'
        p_ch.font.name = 'Georgia'
        p_ch.font.size = Pt(24)
        p_ch.font.color.rgb = hex_to_rgb('#FFFFFF')
        p_ch.alignment = PP_ALIGN.CENTER

        items = [
            ('01', '#174052', 'Agenda 01', 'This slide is an editable\nslide with all your needs.'),
            ('02', '#20B2AA', 'Agenda 02', 'This slide is an editable\nslide with all your needs.'),
            ('03', '#16A085', 'Agenda 03', 'This slide is an editable\nslide with all your needs.'),
            ('04', '#20B2AA', 'Agenda 04', 'This slide is an editable\nslide with all your needs.'),
            ('05', '#16A085', 'Agenda 05', 'This slide is an editable\nslide with all your needs.'),
        ]
        for idx, (num, col_hex, title, desc) in enumerate(items):
            cy = 0.70 + idx * 1.15
            sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.70), Inches(cy + 0.03), Inches(0.82), Inches(0.82))
            sh.fill.solid()
            sh.fill.fore_color.rgb = hex_to_rgb('#D0D7DE')
            sh.line.fill.background()

            c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.70), Inches(cy), Inches(0.80), Inches(0.80))
            c.fill.solid()
            c.fill.fore_color.rgb = hex_to_rgb(col_hex)
            c.line.fill.background()
            pc = c.text_frame.paragraphs[0]
            pc.text = num
            pc.font.name = 'Georgia'
            pc.font.size = Pt(17)
            pc.font.color.rgb = hex_to_rgb('#FFFFFF')
            pc.alignment = PP_ALIGN.CENTER

            tb = slide.shapes.add_textbox(Inches(5.68), Inches(cy - 0.08), Inches(2.95), Inches(0.95))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            pt = tf.paragraphs[0]
            pt.text = title
            pt.font.name = 'Georgia'
            pt.font.size = Pt(14.5)
            pt.font.bold = True
            pt.font.color.rgb = hex_to_rgb(self.charcoal)
            pt.space_after = Pt(2)
            for line in desc.split('\n'):
                pd = tf.add_paragraph()
                pd.text = line
                pd.font.name = 'Georgia'
                pd.font.size = Pt(10)
                pd.font.color.rgb = hex_to_rgb(self.grey_desc)
