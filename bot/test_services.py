# -*- coding: utf-8 -*-
import asyncio
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from pptx import Presentation
from bot.services.notebooklm_service import NotebookLMService
from bot.services.slide_builder import SlideBuilderService


async def test():
    print("[*] 1. Testing NotebookLM service...")
    service = NotebookLMService()
    spec = await service.generate_slide_content("Kriptografiya asoslari", slide_count=3)
    print(f"[+] Content generated: {spec.get('title')} ({len(spec.get('slides', []))} slides)")

    print("[*] 2. Testing SlideBuilder service...")
    out_path = SlideBuilderService.build_presentation(spec)
    print(f"[+] Presentation created at: {out_path}")

    prs = Presentation(out_path)
    print(f"[+] Verified slide count in PPTX: {len(prs.slides)} slides")
    assert len(prs.slides) == 3, f"Expected 3 slides, got {len(prs.slides)}"
    print("[🎉] INTEGRATION TEST PASSED!")


if __name__ == "__main__":
    asyncio.run(test())
