from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
shape = slide.shapes.title
shape.text = "Hello Ozbekiston"

r = shape.text_frame.paragraphs[0].runs[0]
rPr = r._r.get_or_add_rPr()

latins = rPr.xpath('./a:latin')
if not latins:
    latin = OxmlElement('a:latin')
    rPr.insert(0, latin)
else:
    latin = latins[0]
latin.set('typeface', 'Montserrat')

css = rPr.xpath('./a:cs')
if not css:
    cs = OxmlElement('a:cs')
    rPr.insert(1, cs)
else:
    cs = css[0]
cs.set('typeface', 'Montserrat')

print(rPr.xml)
