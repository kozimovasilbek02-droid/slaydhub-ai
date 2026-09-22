import glob
from pptx import Presentation

for p in glob.glob('output/**/*.pptx', recursive=True):
    try:
        prs = Presentation(p)
        for idx, s in enumerate(prs.slides):
            all_t = " ".join([shp.text_frame.text for shp in s.shapes if shp.has_text_frame])
            if "Superpozitsiya" in all_t or "Bloch" in all_t:
                print(f"MATCH IN {p} SLIDE {idx+1}")
                print(f"  Slide layout: {s.slide_layout.name}")
                for sidx, shp in enumerate(s.shapes):
                    l = round(shp.left / 914400 * 72, 1)
                    t = round(shp.top / 914400 * 72, 1)
                    w = round(shp.width / 914400 * 72, 1)
                    h = round(shp.height / 914400 * 72, 1)
                    txt = shp.text_frame.text[:60].replace('\n', ' ') if shp.has_text_frame else 'NO TEXT'
                    print(f"    Shape {sidx}: {shp.name} pos=({l},{t}) size=({w},{h}) txt='{txt}'")
    except Exception as e:
        pass
