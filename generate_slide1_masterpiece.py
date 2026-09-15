import os
import win32com.client
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def hex_to_rgb(hex_code):
    h = hex_code.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

s1 = prs.slides.add_slide(prs.slide_layouts[6])

# 1. Header (Title + Subtitle)
tb_h = s1.shapes.add_textbox(Inches(0.80), Inches(0.80), Inches(11.733), Inches(1.10))
tf_h = tb_h.text_frame
tf_h.word_wrap = True
tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = 0

p_t = tf_h.paragraphs[0]
p_t.text = '100 Days Action Plan'
p_t.font.name = 'Montserrat'
p_t.font.size = Pt(28)
p_t.font.bold = True
p_t.font.color.rgb = hex_to_rgb('#2D3748')
p_t.alignment = PP_ALIGN.CENTER

p_sub = tf_h.add_paragraph()
p_sub.text = 'In 5 Phases'
p_sub.font.name = 'Segoe UI'
p_sub.font.size = Pt(13)
p_sub.font.color.rgb = hex_to_rgb('#718096')
p_sub.alignment = PP_ALIGN.CENTER
p_sub.space_before = Pt(4)

s1_data = [
    ('First 20\nDays', ['Understand the company culture', 'Outline current year commitment plans', 'Create yearly agenda with stakeholders'], '#8CC63F', 0.77),
    ('20-40\nDays', ['Meetings with internal teams', 'Build sales funnel'], '#29B6D8', 3.17),
    ('40-60\nDays', ['Executing plan', 'Formal meeting with clients', 'Dialogue with suitable prospects'], '#FFAA00', 5.57),
    ('60-80\nDays', ['Full engagement with client & team on existing opportunities'], '#EA4335', 7.97),
    ('80-100\nDays', ['Working towards mile stones'], '#AF26FA', 10.37),
]

cw = 2.18
ch = 3.90
y_top = 2.45

for title, bullets, col, x in s1_data:
    c_rgb = hex_to_rgb(col)
    
    # 1. Soft Drop Shadow Base (Offset layer for realistic depth)
    shadow = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.03), Inches(y_top + 0.04), Inches(cw), Inches(ch))
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = hex_to_rgb('#E2E8F0')
    shadow.line.fill.background()
    
    # 2. Main White Card
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y_top), Inches(cw), Inches(ch))
    card.fill.solid()
    card.fill.fore_color.rgb = hex_to_rgb('#FFFFFF')
    card.line.color.rgb = hex_to_rgb('#EDF2F7')
    card.line.width = Pt(1.0)
    
    # 3. Downward Pointed Top Banner (Single clean polygon, zero seams)
    pts_banner = [
        (Inches(x), Inches(y_top)),
        (Inches(x + cw), Inches(y_top)),
        (Inches(x + cw), Inches(y_top + 1.15)),
        (Inches(x + cw/2), Inches(y_top + 1.45)),
        (Inches(x), Inches(y_top + 1.15))
    ]
    fb = s1.shapes.build_freeform(pts_banner[0][0], pts_banner[0][1])
    fb.add_line_segments(pts_banner[1:], close=True)
    ban = fb.convert_to_shape()
    ban.fill.solid()
    ban.fill.fore_color.rgb = c_rgb
    ban.line.fill.background()
    
    # Banner Text
    tb_b = s1.shapes.add_textbox(Inches(x), Inches(y_top + 0.12), Inches(cw), Inches(0.95))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0
    lines = title.split('\n')
    p1 = tf_b.paragraphs[0]
    p1.text = lines[0]
    p1.font.name = 'Montserrat'
    p1.font.size = Pt(13.5)
    p1.font.bold = True
    p1.font.color.rgb = hex_to_rgb('#FFFFFF')
    p1.alignment = PP_ALIGN.CENTER
    if len(lines) > 1:
        p2 = tf_b.add_paragraph()
        p2.text = lines[1]
        p2.font.name = 'Montserrat'
        p2.font.size = Pt(13.5)
        p2.font.bold = True
        p2.font.color.rgb = hex_to_rgb('#FFFFFF')
        p2.alignment = PP_ALIGN.CENTER
        
    # Card Body Content
    tb_t = s1.shapes.add_textbox(Inches(x + 0.16), Inches(y_top + 1.65), Inches(cw - 0.32), Inches(ch - 1.80))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    for i, b in enumerate(bullets):
        p = tf_t.paragraphs[0] if i == 0 else tf_t.add_paragraph()
        
        # Checkmark in matching color + Text in dark charcoal
        run_chk = p.add_run()
        run_chk.text = '✓  '
        run_chk.font.name = 'Segoe UI'
        run_chk.font.size = Pt(10)
        run_chk.font.bold = True
        run_chk.font.color.rgb = c_rgb
        
        run_txt = p.add_run()
        run_txt.text = b
        run_txt.font.name = 'Segoe UI'
        run_txt.font.size = Pt(10)
        run_txt.font.color.rgb = hex_to_rgb('#374151')
        p.space_after = Pt(8)

out_file = os.path.abspath('output/action_s1_masterpiece.pptx')
prs.save(out_file)
print('Saved perfected Slide 1 to:', out_file)
