import io
import re
from pptx import Presentation
from PIL import Image

prs = Presentation(r'output\academic_studio\Rejim2_Shablonlar_Bazasi.pptx')

def get_slide_bg_luminance(part):
    if hasattr(part, "_cached_bg_lum"):
        return part._cached_bg_lum

    lum = 255.0 # default light
    containers = []
    if hasattr(part, "slide"):
        containers.append(part.slide)
    if hasattr(part, "slide_layout"):
        containers.append(part.slide_layout)
        if hasattr(part.slide_layout, "slide_master"):
            containers.append(part.slide_layout.slide_master)

    found = False
    for container in containers:
        if found:
            break
        try:
            xml = container._element.xml if hasattr(container, "_element") else ""
            bg_matches = re.findall(r'<p:bg>.*?</p:bg>', xml, re.DOTALL) or re.findall(r'<p:bgPr>.*?</p:bgPr>', xml, re.DOTALL)
            for bg_xml in bg_matches:
                m_blip = re.search(r'r:embed="([^"]+)"', bg_xml)
                if m_blip:
                    rid = m_blip.group(1)
                    c_part = container.part if hasattr(container, "part") else part
                    if rid in c_part.rels:
                        target = c_part.rels[rid].target_part
                        img = Image.open(io.BytesIO(target.blob))
                        img_small = img.resize((15, 15)).convert('L')
                        pixels = list(img_small.getdata())
                        lum = sum(pixels) / float(len(pixels))
                        found = True
                        break
                m_color = re.search(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', bg_xml)
                if m_color:
                    hex_str = m_color.group(1)
                    r, g, b = int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16)
                    lum = 0.299 * r + 0.587 * g + 0.114 * b
                    found = True
                    break
        except Exception:
            pass

    # Check full canvas picture shape
    if not found and hasattr(part, "slide"):
        try:
            slide = part.slide
            sw = part.package.presentation_part.presentation.slide_width
            sh = part.package.presentation_part.presentation.slide_height
            for s in slide.shapes:
                if s.shape_type == 13 and s.width >= sw * 0.90 and s.height >= sh * 0.90:
                    img = Image.open(io.BytesIO(s.image.blob))
                    img_small = img.resize((15, 15)).convert('L')
                    pixels = list(img_small.getdata())
                    lum = sum(pixels) / float(len(pixels))
                    found = True
                    break
        except Exception:
            pass

    part._cached_bg_lum = lum
    return lum

def is_shape_on_dark_background(shape) -> bool:
    # 1. Check shape own fill
    try:
        s_xml = shape._element.xml
        m_spPr = re.search(r'<p:spPr>(.*?)</p:spPr>', s_xml, re.DOTALL)
        if m_spPr:
            spPr = m_spPr.group(1)
            if not ('<a:noFill' in spPr):
                m_clr = re.search(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', spPr)
                if m_clr:
                    hex_str = m_clr.group(1)
                    r, g, b = int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16)
                    lum = 0.299 * r + 0.587 * g + 0.114 * b
                    return lum < 135.0
    except Exception:
        pass

    # 2. Transparent or inherited: check slide background
    try:
        part = shape.part
        lum = get_slide_bg_luminance(part)
        return lum < 135.0
    except Exception:
        pass

    return False

for idx, slide in enumerate(prs.slides):
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip():
            dark = is_shape_on_dark_background(shp)
            print(f"Slide {idx+1} | Shape: {shp.name[:20]:20} | DarkBG: {str(dark):5}")
