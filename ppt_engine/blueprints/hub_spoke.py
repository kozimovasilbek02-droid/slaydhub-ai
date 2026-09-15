from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from ppt_engine.theme import Colors, Fonts, Sizes, hex_to_rgb
from ppt_engine.typography import add_slide_header, add_card_text
from ppt_engine.cards import create_floating_card, create_circular_badge
from ppt_engine.layout import distribute_vertical

def build_center_hub_slide(prs, title: str, center_hero_path: str, left_cards: list, right_cards: list):
    """
    Builds a central hub slide flanked by left and right floating cards with connecting spoke lines.
    left_cards/right_cards format: [(title, desc, color_hex, icon_or_num), ...]
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, title)
    
    # Place center hero graphic
    if center_hero_path:
        slide.shapes.add_picture(center_hero_path, Inches(4.35), Inches(1.80), width=Inches(4.65))
        
    ys = distribute_vertical(len(left_cards), card_height=0.78, start_y=1.50, end_y=6.48)
    
    # Left Cards
    for i, (c_title, c_desc, col_hex, badge_txt) in enumerate(left_cards):
        y = ys[i]
        c_rgb = hex_to_rgb(col_hex)
        
        # Spoke Line
        ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.00), Inches(y + 0.38), Inches(0.85), Inches(0.02))
        ln.fill.solid()
        ln.fill.fore_color.rgb = hex_to_rgb(Colors.BORDER_LIGHT)
        ln.line.fill.background()
        
        # Floating Card
        create_floating_card(slide, 0.80, y, 3.50, 0.78, border_color=col_hex, border_width=1.5)
        
        # Right circular badge disc
        create_circular_badge(slide, 3.65, y + 0.09, 0.60, bg_color='#FFFFFF', border_color=col_hex, text=badge_txt, text_color=col_hex, font_size=11)
        
        # Text
        add_card_text(
            slide, 0.95, y + 0.08, 2.65, 0.62,
            title=c_title, body=c_desc,
            align=PP_ALIGN.RIGHT,
            title_color=Colors.TEXT_TITLE,
            body_color=Colors.TEXT_MUTED,
            title_size=10.5, body_size=8.5
        )
        
    # Right Cards
    for i, (c_title, c_desc, col_hex, badge_txt) in enumerate(right_cards):
        y = ys[i]
        c_rgb = hex_to_rgb(col_hex)
        
        # Spoke Line
        ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.50), Inches(y + 0.38), Inches(0.85), Inches(0.02))
        ln.fill.solid()
        ln.fill.fore_color.rgb = hex_to_rgb(Colors.BORDER_LIGHT)
        ln.line.fill.background()
        
        # Floating Card
        create_floating_card(slide, 9.00, y, 3.50, 0.78, border_color=col_hex, border_width=1.5)
        
        # Left circular badge disc
        create_circular_badge(slide, 9.08, y + 0.09, 0.60, bg_color='#FFFFFF', border_color=col_hex, text=badge_txt, text_color=col_hex, font_size=11)
        
        # Text
        add_card_text(
            slide, 9.75, y + 0.08, 2.65, 0.62,
            title=c_title, body=c_desc,
            align=PP_ALIGN.LEFT,
            title_color=Colors.TEXT_TITLE,
            body_color=Colors.TEXT_MUTED,
            title_size=10.5, body_size=8.5
        )
        
    return slide
