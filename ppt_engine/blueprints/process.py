from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from ppt_engine.theme import Colors, Fonts, Sizes, hex_to_rgb
from ppt_engine.typography import add_slide_header, add_card_text
from ppt_engine.cards import create_two_tone_card, create_circular_badge
from ppt_engine.layout import distribute_horizontal

def build_two_row_process_slide(prs, title: str, row1_steps: list, row2_steps: list):
    """
    Builds a 10-step process slide with 2 rows of 5 two-tone block cards connected by a ribbon.
    row1_steps / row2_steps format: [(title, col_top, col_bot, num), ...]
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, title)
    
    # Connecting horizontal ribbon
    rb = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.80), Inches(3.60), Inches(11.733), Inches(0.12))
    rb.fill.solid()
    rb.fill.fore_color.rgb = hex_to_rgb(Colors.BORDER_LIGHT)
    rb.line.fill.background()
    
    xs1 = distribute_horizontal(len(row1_steps), card_width=1.85, slide_width=13.333, margin_x=1.10)
    xs2 = distribute_horizontal(len(row2_steps), card_width=1.85, slide_width=13.333, margin_x=1.10)
    
    # Row 1 (Top Cards)
    for i, (st_title, col_top, col_bot, num) in enumerate(row1_steps):
        x = xs1[i]
        top_s, bot_s = create_two_tone_card(slide, x, 1.85, 1.85, 2.00, col_top, col_bot, top_ratio=0.55)
        
        # Text in Top
        tb = slide.shapes.add_textbox(Inches(x + 0.10), Inches(1.95), Inches(1.65), Inches(0.90))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = st_title
        p.font.name = Fonts.TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb(Colors.TEXT_WHITE)
        p.alignment = PP_ALIGN.CENTER
        
        # Number in Bot
        p_n = bot_s.text_frame.paragraphs[0]
        p_n.text = num
        p_n.font.name = Fonts.TITLE
        p_n.font.size = Pt(18)
        p_n.font.bold = True
        p_n.font.color.rgb = hex_to_rgb(Colors.TEXT_WHITE)
        p_n.alignment = PP_ALIGN.CENTER
        
    # Row 2 (Bottom Cards)
    for i, (st_title, col_top, col_bot, num) in enumerate(row2_steps):
        x = xs2[i]
        top_s, bot_s = create_two_tone_card(slide, x, 4.30, 1.85, 2.00, col_top, col_bot, top_ratio=0.45)
        
        # Number in Top
        p_n = top_s.text_frame.paragraphs[0]
        p_n.text = num
        p_n.font.name = Fonts.TITLE
        p_n.font.size = Pt(18)
        p_n.font.bold = True
        p_n.font.color.rgb = hex_to_rgb(Colors.TEXT_WHITE)
        p_n.alignment = PP_ALIGN.CENTER
        
        # Text in Bot
        tb = slide.shapes.add_textbox(Inches(x + 0.10), Inches(5.30), Inches(1.65), Inches(0.90))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = st_title
        p.font.name = Fonts.TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb(Colors.TEXT_WHITE)
        p.alignment = PP_ALIGN.CENTER
        
    return slide
