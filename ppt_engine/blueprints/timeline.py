from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from ppt_engine.theme import Colors, Fonts, Sizes, hex_to_rgb
from ppt_engine.geometry import create_seamless_callout, create_chevron, create_downward_banner
from ppt_engine.typography import add_slide_header, add_card_text
from ppt_engine.cards import create_circular_badge, create_floating_card
from ppt_engine.layout import distribute_horizontal

def build_callout_timeline_slide(prs, title: str, steps_data: list, baseline_y: float = 5.30):
    """
    Builds a timeline slide with seamless callout cards pointing to baseline milestone beads.
    steps_data format: [(title, desc, color_hex, icon_or_badge_text), ...]
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, title)
    
    # Baseline
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.80), Inches(baseline_y), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = hex_to_rgb(Colors.BORDER_LIGHT)
    line.line.fill.background()
    
    count = len(steps_data)
    card_w = 2.45 if count == 4 else 2.05
    card_h = 2.80
    xs = distribute_horizontal(count, card_w, slide_width=13.333, margin_x=0.80)
    
    for i, (st_title, st_desc, col_hex, badge_txt) in enumerate(steps_data):
        x = xs[i]
        c_rgb = hex_to_rgb(col_hex)
        
        # Single seamless speech callout shape (with integrated bottom pointer notch)
        callout = create_seamless_callout(
            slide, x, 1.80, card_w, card_h,
            notch_w=0.35, notch_h=0.25,
            bg_color=Colors.BG_CARD,
            border_color=col_hex,
            border_width=1.5,
            has_shadow=True
        )
        
        # Badge / Icon
        if badge_txt:
            tb_i = slide.shapes.add_textbox(Inches(x), Inches(2.05), Inches(card_w), Inches(0.50))
            p_i = tb_i.text_frame.paragraphs[0]
            p_i.text = badge_txt
            p_i.font.name = Fonts.TITLE
            p_i.font.size = Pt(20)
            p_i.font.bold = True
            p_i.font.color.rgb = c_rgb
            p_i.alignment = PP_ALIGN.CENTER
            
        # Card Text
        text_y = 2.70 if badge_txt else 2.10
        add_card_text(
            slide, x + 0.15, text_y, card_w - 0.30, card_h - 1.0,
            title=st_title, body=st_desc,
            align=PP_ALIGN.CENTER,
            title_color=Colors.TEXT_TITLE,
            body_color=Colors.TEXT_MUTED,
            title_size=12, body_size=9
        )
        
        # Milestone bead on baseline
        create_circular_badge(
            slide, x + (card_w - 0.28) / 2.0, baseline_y - 0.13, 0.28,
            bg_color=col_hex, has_shadow=False
        )
        
        # Year label below baseline
        tb_yr = slide.shapes.add_textbox(Inches(x), Inches(baseline_y + 0.30), Inches(card_w), Inches(0.40))
        p_yr = tb_yr.text_frame.paragraphs[0]
        p_yr.text = '20XX'
        p_yr.font.name = Fonts.TITLE
        p_yr.font.size = Pt(11)
        p_yr.font.bold = True
        p_yr.font.color.rgb = hex_to_rgb(Colors.TEXT_TITLE)
        p_yr.alignment = PP_ALIGN.CENTER
        
    return slide

def build_chevron_timeline_slide(prs, title: str, steps_data: list):
    """
    Builds an alternating chevron timeline slide.
    steps_data format: [(title, desc, color_hex, step_num), ...]
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(slide, title)
    
    count = len(steps_data)
    chv_w = 2.45
    chv_h = 0.65
    y_chv = 3.65
    xs = distribute_horizontal(count, chv_w, slide_width=13.333, margin_x=0.80)
    
    for i, (st_title, st_desc, col_hex, num) in enumerate(steps_data):
        x = xs[i]
        is_top = (i % 2 == 0)
        c_rgb = hex_to_rgb(col_hex)
        
        # Chevron
        create_chevron(slide, x, y_chv, chv_w, chv_h, fill_color=col_hex, text='20XX', font_size=11)
        
        # Disc position
        y_disc = 2.00 if is_top else 4.75
        
        # Vertical connector
        ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.55), Inches(2.95 if is_top else 4.30), Inches(0.02), Inches(0.70))
        ln.fill.solid()
        ln.fill.fore_color.rgb = hex_to_rgb(Colors.BORDER_LIGHT)
        ln.line.fill.background()
        
        # Circular Badge
        create_circular_badge(slide, x + 0.05, y_disc, 1.05, bg_color=col_hex, text=num, font_size=16)
        
        # Card Text
        add_card_text(
            slide, x + 1.20, y_disc + 0.05, 1.50, 0.95,
            title=st_title, body=st_desc,
            align=PP_ALIGN.LEFT,
            title_color=Colors.TEXT_TITLE,
            body_color=Colors.TEXT_MUTED,
            title_size=11, body_size=8.5
        )
    return slide
