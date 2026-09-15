from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from ppt_engine.theme import hex_to_rgb, Colors, Fonts
from ppt_engine.geometry import add_soft_shadow
from ppt_engine.typography import add_card_text

def create_floating_card(slide, x: float, y: float, w: float, h: float, 
                         border_color: str = Colors.BORDER_LIGHT,
                         bg_color: str = Colors.BG_CARD,
                         border_width: float = 1.5,
                         has_shadow: bool = True):
    """Creates a sleek modern floating card with soft ambient drop shadow."""
    if has_shadow:
        add_soft_shadow(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, 0.03, 0.03)
        
    crd = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    crd.fill.solid()
    crd.fill.fore_color.rgb = hex_to_rgb(bg_color)
    crd.line.color.rgb = hex_to_rgb(border_color)
    crd.line.width = Pt(border_width)
    return crd

def create_pill_card(slide, x: float, y: float, w: float, h: float,
                     accent_color: str = Colors.BLUE,
                     is_cap_left: bool = True,
                     has_shadow: bool = True):
    """Creates a pill card with an accent cap on one side."""
    if has_shadow:
        add_soft_shadow(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, 0.03, 0.03)
        
    crd = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    crd.fill.solid()
    crd.fill.fore_color.rgb = hex_to_rgb(Colors.BG_CARD)
    crd.line.color.rgb = hex_to_rgb(accent_color)
    crd.line.width = Pt(1.5)
    
    cap_x = x if is_cap_left else (x + w - 0.40)
    cap = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cap_x), Inches(y), Inches(0.40), Inches(h))
    cap.fill.solid()
    cap.fill.fore_color.rgb = hex_to_rgb(accent_color)
    cap.line.fill.background()
    return crd

def create_two_tone_card(slide, x: float, y: float, w: float, h: float,
                         top_color: str, bot_color: str,
                         top_ratio: float = 0.55,
                         has_shadow: bool = True):
    """Creates a two-tone matrix block card (e.g. top header block + bottom number block)."""
    if has_shadow:
        add_soft_shadow(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, 0.03, 0.03)
        
    h_top = h * top_ratio
    h_bot = h * (1.0 - top_ratio)
    
    # Top Half
    top_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h_top))
    top_shape.fill.solid()
    top_shape.fill.fore_color.rgb = hex_to_rgb(top_color)
    top_shape.line.fill.background()
    
    # Bottom Half
    bot_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + h_top), Inches(w), Inches(h_bot))
    bot_shape.fill.solid()
    bot_shape.fill.fore_color.rgb = hex_to_rgb(bot_color)
    bot_shape.line.fill.background()
    return top_shape, bot_shape

def create_circular_badge(slide, x: float, y: float, size: float,
                          bg_color: str, border_color: str = '',
                          text: str = '', text_color: str = '#FFFFFF',
                          font_size: float = 12, has_shadow: bool = True):
    """Creates a circular badge for milestone nodes, numbers, or icons."""
    if has_shadow:
        add_soft_shadow(slide, MSO_SHAPE.OVAL, x, y, size, size, 0.03, 0.03)
        
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    badge.fill.solid()
    badge.fill.fore_color.rgb = hex_to_rgb(bg_color)
    
    if border_color:
        badge.line.color.rgb = hex_to_rgb(border_color)
        badge.line.width = Pt(1.5)
    else:
        badge.line.fill.background()
        
    if text:
        p = badge.text_frame.paragraphs[0]
        p.text = text
        p.font.name = Fonts.TITLE
        p.font.size = Pt(font_size)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb(text_color)
        p.alignment = PP_ALIGN.CENTER
    return badge
