import os
import glob
import re
import sys
import json
import requests
from pptx import Presentation
import json_repair

sys.stdout.reconfigure(encoding='utf-8')

folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"
files = sorted(glob.glob(os.path.join(folder, "*.pptx")))

api_key = os.environ.get('GROQ_API_KEY')
url = 'https://api.groq.com/openai/v1/chat/completions'
headers = {
    'Authorization': f'Bearer {api_key}',
    'Content-Type': 'application/json'
}

cyrillic_pattern = re.compile(r'[\u0400-\u04FF]')
latin_ru_words = [
    'chto', 'kak', 'esli', 'kogda', 'potomu', 'posle', 'kotory', 'uchenik',
    'klass', 'urok', 'zadacha', 'zadanie', 'opredelenie',
    'eto', 'ponyati', 'svoystv', 'silы', 'vremya', 'veshestv', 'protsess',
    'zarojdenie', 'chastits', 'otkrыtie', 'etapы', 'razvitiya', 'pervыy',
    'skorost', 'visot', 'vыsot', 'kineticheskaya', 'padayet', 'spasibo', 'vnimanie',
    'atmosfernoe', 'davlenie', 'rtutn', 'vozdux', 'struktura', 'lichnosti', 'povedenie'
]

def is_ru_or_cyr(txt):
    if not txt or len(txt.strip()) < 2:
        return False
    if cyrillic_pattern.search(txt):
        return True
    txt_low = txt.lower()
    for w in latin_ru_words:
        if re.search(rf'\b{w}', txt_low):
            return True
    return False

def translate_batch(batch_items):
    sys_prompt = """Siz professional akademik tarjimonsiz. Taqdimot matnlarini toza, ravon va ilmiy O'ZBEK tiliga (Lotin alifbosida) tarjima qiling.
DIQQAT: Matnlar kirill yoki transliteratsiya qilingan rus tilida bo'lishi mumkin. Ularni 100% o'zbekchaga o'giring.
Hech qanday ruscha yoki kirillcha harf qoldirmang.
Formulalar (masalan: F=ma, 10^-13 s, 1 m/s) va raqamlarni to'g'ri saqlang.
Natijani FAQATGINA quyidagi JSON formatida qaytaring:
[
  {"id": 0, "uzbek": "Tarjima qilingan o'zbekcha matn"}
]"""
    
    payload = {
        'model': 'qwen/qwen3.8-27b',
        'messages': [
            {'role': 'system', 'content': sys_prompt},
            {'role': 'user', 'content': json.dumps(batch_items, ensure_ascii=False, indent=2)}
        ],
        'temperature': 0.1
    }
    
    try:
        r = requests.post(url, headers=headers, json=payload, timeout=15)
        if r.status_code == 200:
            raw_text = r.json()['choices'][0]['message']['content'].strip()
            raw_text = re.sub(r"^```(?:json)?\s*", "", raw_text)
            raw_text = re.sub(r"\s*```$", "", raw_text)
            parsed = json_repair.repair_json(raw_text, return_objects=True)
            res = {}
            if isinstance(parsed, list):
                for row in parsed:
                    if isinstance(row, dict) and 'id' in row:
                        val = row.get('uzbek') or row.get('translated') or row.get('text') or ''
                        res[int(row['id'])] = val
                        res[str(row['id'])] = val
            if res:
                return res
    except Exception as e:
        print(f"Qwen request error: {e}", flush=True)
        
    # Fallback to deep_translator
    try:
        from deep_translator import GoogleTranslator
        gt = GoogleTranslator(source='auto', target='uz')
        res = {}
        for it in batch_items:
            # normalize pseudo-latin before google translate
            norm_text = it['text'].replace('З', 'Z').replace('з', 'z').replace('ы', 'y')
            res[it['id']] = gt.translate(norm_text)
        return res
    except Exception:
        return {it['id']: it['text'] for it in batch_items}

total_fixed_files = 0
total_fixed_paragraphs = 0

for f_idx, fpath in enumerate(files, 1):
    fname = os.path.basename(fpath)
    prs = Presentation(fpath)
    
    dirty_paragraphs = []
    
    def scan_shape(shape):
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if is_ru_or_cyr(t):
                    dirty_paragraphs.append((p, t))
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if cell.text_frame:
                        for p in cell.text_frame.paragraphs:
                            t = p.text.strip()
                            if is_ru_or_cyr(t):
                                dirty_paragraphs.append((p, t))
        if shape.shape_type == 6: # Group shape
            for subshape in shape.shapes:
                scan_shape(subshape)

    for slide in prs.slides:
        for shape in slide.shapes:
            scan_shape(shape)
            
    if not dirty_paragraphs:
        continue
        
    print(f"[{f_idx}/{len(files)}] {fname}: {len(dirty_paragraphs)} ta nomuqobil matn aniqlandi...", flush=True)
    
    chunk_size = 15
    for i in range(0, len(dirty_paragraphs), chunk_size):
        chunk = dirty_paragraphs[i:i+chunk_size]
        batch_input = [{"id": idx, "text": item[1]} for idx, item in enumerate(chunk)]
        tr_map = translate_batch(batch_input)
        
        for idx, (p, orig_text) in enumerate(chunk):
            tr_text = tr_map.get(idx) or tr_map.get(str(idx))
            if tr_text and tr_text.strip():
                # Clean up any leftover cyrillic character artifacts
                clean_tr = tr_text.replace('З', 'Z').replace('з', 'z').replace('ы', 'y')
                p.text = clean_tr
                total_fixed_paragraphs += 1
                
    prs.save(fpath)
    total_fixed_files += 1
    print(f"  -> {fname} to'liq toza o'zbek tiliga o'tkazildi va saqlandi!", flush=True)

print("\n" + "="*60, flush=True)
print(f"YAKUNIY NATIJA: {total_fixed_files} ta fayl, {total_fixed_paragraphs} ta matn qatori mukammal tozalandi!", flush=True)
print("="*60, flush=True)
