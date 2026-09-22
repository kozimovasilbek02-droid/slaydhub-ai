from pptx import Presentation
import os

folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"

# 1. Fizika - Moddiy nuqta va sanoq sistemasi (9-sinf).pptx
f1 = os.path.join(folder, "Fizika - Moddiy nuqta va sanoq sistemasi (9-sinf).pptx")
prs1 = Presentation(f1)
for s in prs1.slides:
    for shape in s.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "davlenie" in p.text:
                    p.text = p.text.replace("davlenie", "bosim")
prs1.save(f1)

# 2. Geografiya - Atmosfera bosimi va barometrlar (7-sinf).pptx
f2 = os.path.join(folder, "Geografiya - Atmosfera bosimi va barometrlar (7-sinf).pptx")
prs2 = Presentation(f2)
for s in prs2.slides:
    for shape in s.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Vpervye vesomost vozduxa" in p.text or "vesomost vozduxa" in p.text:
                    p.text = "1638-yilda Florensiyadagi saroy bog'ida favvora qurish urinishi muvaffaqiyatsiz tugagach, havoning og'irligi birinchi marta olimlarni hayratda qoldirdi."
prs2.save(f2)

# 3. Menejment - Nizolar va stresslarni boshqarish.pptx
f3 = os.path.join(folder, "Menejment - Nizolar va stresslarni boshqarish.pptx")
prs3 = Presentation(f3)
for s in prs3.slides:
    for shape in s.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if "Faza trevogi" in p.text:
                    p.text = "Tashvish fazasi – yaqinlashib kelayotgan xavf-xatar oldidagi noaniqlik va qo'rquv holati."
                elif "Krolik" in p.text and "passivnaya" in p.text:
                    p.text = "Quyon (Krolik) – stress holatiga passiv reaksiya. Stress qarshilik ko'rsatish imkoniyatidan mahrum qiladi, inson o'zini nochor his qiladi."
prs3.save(f3)

print("Final polish completed!")
