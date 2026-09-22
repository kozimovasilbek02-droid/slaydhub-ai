import os
import requests
import json
import re
import sys
from pptx import Presentation
import json_repair

sys.stdout.reconfigure(encoding='utf-8')

folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"
target_files = [
    "Fizika - Elektr maydoni ishi va potensiali (10-sinf).pptx",
    "Fizika - Elementar zarralar fizikasining paydo bo'lishi (11-sinf).pptx",
    "Fizika - Energiya saqlanish qonuni bo'yicha masalalar (10-sinf).pptx",
    "Fizika - Ichki energiya va uni o'zgartirish usullari (10-sinf).pptx",
    "Fizika - Moddiy nuqta va sanoq sistemasi (9-sinf).pptx",
    "Geografiya - Atmosfera bosimi va barometrlar (7-sinf).pptx",
    "Iqtisodiyot - Xodimlar motivatsiyasi.pptx",
    "Menejment - Nizolar va stresslarni boshqarish.pptx"
]

api_key = os.environ.get('GROQ_API_KEY')
url = 'https://api.groq.com/openai/v1/chat/completions'
headers = {
    'Authorization': f'Bearer {api_key}',
    'Content-Type': 'application/json'
}

prompt_template = """Siz O'zbek tili va fizika/iqtisodiyot/geografiya fanlari bo'yicha oliy toifali mutaxassissiz.
Quyidagi ro'yxatda berilgan matnlarni toza, ravon va ilmiy O'ZBEK tiliga (Lotin yozuvida) tarjima qiling.
DIQQAT: Matnlar rus tilida, kirillda yoki lotin harflari bilan yozilgan ruscha transliteratsiyada bo'lishi mumkin (masalan: 'ЗAROJDENIE FIЗIKI Elementarnыx chastits', 'DOMAShNEE ЗADANIE', 'Rtutnыy barometr', 'Snijaetsya kontsentratsiya vnimaniya', 'SPASIBO ZA VNIMANIE', 'Metallы obladayut samoy vыsokoy teploprovodnostyu').
Ularni 100% o'zbek tilidagi ekvivalentiga aylantiring (masalan: 'Elementar zarralar fizikasining paydo bo\\'lishi', 'Uyga vazifa', 'Simobli barometr', 'Diqqat jamlanishi pasayadi', 'E\\'tiboringiz uchun rahmat!', 'Metalllar eng yuqori issiqlik o\\'tkazuvchanlikka ega').
Formulalar (masalan: E=mc^2, F=ma, 10^-13 s) va sonlarni to'g'ri saqlang.

Javobni FAQAT JSON formatda bering:
[
  {{"id": 0, "uzbek": "tarjima matni"}}
]
"""

def translate_texts(texts_with_ids):
    payload = {
        'model': 'openai/gpt-oss-120b',
        'messages': [
            {'role': 'system', 'content': prompt_template},
            {'role': 'user', 'content': json.dumps(texts_with_ids, ensure_ascii=False, indent=2)}
        ],
        'temperature': 0.1
    }
    for model in ['openai/gpt-oss-120b', 'openai/gpt-oss-20b', 'qwen/qwen3.8-27b']:
        payload['model'] = model
        try:
            r = requests.post(url, headers=headers, json=payload, timeout=12)
            if r.status_code == 200:
                raw_text = r.json()['choices'][0]['message']['content'].strip()
                raw_text = re.sub(r"^```(?:json)?\s*", "", raw_text)
                raw_text = re.sub(r"\s*```$", "", raw_text)
                parsed = json_repair.repair_json(raw_text, return_objects=True)
                res = {}
                if isinstance(parsed, list):
                    for item in parsed:
                        if isinstance(item, dict) and 'id' in item:
                            val = item.get('uzbek') or item.get('translated') or item.get('text') or ''
                            res[int(item['id'])] = val
                            res[str(item['id'])] = val
                if len(res) >= len(texts_with_ids) * 0.7:
                    return res
        except Exception as e:
            print(f"Error calling {model}: {e}")
            pass
    return {}

for f_idx, fname in enumerate(target_files, 1):
    fpath = os.path.join(folder, fname)
    if not os.path.exists(fpath):
        print(f"File not found: {fpath}")
        continue
    prs = Presentation(fpath)
    
    items_to_translate = []
    
    for s_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = p.text.strip()
                    if t and len(t) >= 2:
                        items_to_translate.append((p, t))
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            for p in cell.text_frame.paragraphs:
                                t = p.text.strip()
                                if t and len(t) >= 2:
                                    items_to_translate.append((p, t))
                                    
    print(f"\n[{f_idx}/{len(target_files)}] {fname} ({len(items_to_translate)} paragraphs)...", flush=True)
    
    batch_size = 15
    for i in range(0, len(items_to_translate), batch_size):
        chunk = items_to_translate[i:i+batch_size]
        batch_input = [{"id": idx, "text": p_tuple[1]} for idx, p_tuple in enumerate(chunk)]
        
        tr_res = translate_texts(batch_input)
        for idx, (p, orig_text) in enumerate(chunk):
            tr_val = tr_res.get(idx) or tr_res.get(str(idx))
            if tr_val and tr_val.strip() and tr_val.strip() != orig_text.strip():
                # Clean up any leftover cyrillic artifact if any
                clean_tr = tr_val.replace('З', "Z").replace('з', 'z').replace('ы', 'y')
                p.text = clean_tr
                
    prs.save(fpath)
    print(f"  -> Saved {fname} successfully!", flush=True)

print("\nFinished refining all 8 files!", flush=True)
