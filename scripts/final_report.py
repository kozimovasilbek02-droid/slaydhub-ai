import os
import glob
import re
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"
files = sorted([f for f in glob.glob(os.path.join(folder, "*.pptx")) if not os.path.basename(f).startswith("~$")])

cyrillic_pattern = re.compile(r'[\u0400-\u04FF]')
real_ru_words = [r'\bchto\b', r'\bkak\b', r'\besli\b', r'\bkogda\b', r'\bpotomu\b', r'\bdlya\b', r'\bposle\b', r'\bkotory[^\s]*\b', r'\buchenik[^\s]*\b', r'\burok[^\s]*\b', r'\bzadanie\b', r'\beto\b', r'\bvremya\b', r'\bveshestv[^\s]*\b', r'\bspasibo\b', r'\bvnimanie\b']
real_watermarks = ['antonenkova', 'allppt', 'slidesgo', 'presentationgo', 'pedsovet', 'infourok', 'myshared', 'prezentacii', 'videouroki', 'viki.rdf.ru', 'uroki.net', 'uchportal', 'kopilkaurokov']

cyr_found = 0
ru_found = 0
wm_found = 0
quote_bug_found = 0

for f in files:
    fname = os.path.basename(f)
    prs = Presentation(f)
    for s_idx, s in enumerate(prs.slides):
        for shape in s.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text
                if cyrillic_pattern.search(txt):
                    cyr_found += 1
                for rw in real_ru_words:
                    if re.search(rw, txt, re.I):
                        if not txt.startswith('http') and not 'Koterova' in txt and not 'Axriyeva' in txt and not 'Smutnoye vremya' in txt:
                            ru_found += 1
                for wm in real_watermarks:
                    if wm in txt.lower():
                        wm_found += 1
                if len(txt) >= 5 and len(txt) % 2 == 1 and all(c == "'" for c in txt[::2]):
                    quote_bug_found += 1

print("\n" + "="*50)
print("YAKUNIY 100% TEKSHIRUV NATIJASI:")
print(f"Jami taqdimotlar soni: {len(files)}")
print(f"Qo'shtirnoq xatosi (Quote bug): {quote_bug_found}")
print(f"Kirill harflari xatosi: {cyr_found}")
print(f"Ruscha jumlalar xatosi: {ru_found}")
print(f"Suv belgilari/Reklamalar: {wm_found}")
print("="*50)
