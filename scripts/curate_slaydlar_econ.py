import os
import sys
import re
import json
import shutil
import hashlib
import zipfile
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')

POWERPOINT_ROOT = Path(r"C:\Users\user\Desktop\Antigravity\Power Point")
SOFF_ROOT = Path(r"C:\Users\user\Desktop\Antigravity\Soff.uz")
ECON_DIR = SOFF_ROOT / "Slaydlar_ECON"
SARA_DIR = ECON_DIR / "01_Sara_Slaydlar_ECON"
KARANTIN_DIR = ECON_DIR / "99_Karantin_Tekshirish_Uchun"
TEMP_CONV = SOFF_ROOT / "Temp_PPT_Convert" / "econ_ppt"

SARA_DIR.mkdir(parents=True, exist_ok=True)
KARANTIN_DIR.mkdir(parents=True, exist_ok=True)
TEMP_CONV.mkdir(parents=True, exist_ok=True)

import win32com.client
from pptx import Presentation

def fix_extensionless_files():
    print("--- 1. Nomsiz/kengaytmasiz fayllarni tekshirish va to'g'rilash ---")
    for f in ECON_DIR.glob("**/*"):
        if f.is_file() and f.suffix.lower() not in ['.ppt', '.pptx', '.md', '.json', '.py']:
            with open(f, 'rb') as fp:
                hdr = fp.read(8)
            if hdr[:4] == b'PK\x03\x04':
                new_p = f.with_suffix('.pptx')
                shutil.move(str(f), str(new_p))
                print(f"  Fix .pptx: {f.name} -> {new_p.name}")
            elif hdr[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                new_p = f.with_suffix('.ppt')
                shutil.move(str(f), str(new_p))
                print(f"  Fix .ppt: {f.name} -> {new_p.name}")

def convert_ppt_to_pptx():
    print("\n--- 2. Eski .ppt formatlarni zamonaviy .pptx ga o'tkazish ---")
    ppt_files = [f for f in ECON_DIR.glob("**/*.ppt") if not f.name.startswith("~$") and f.suffix.lower() == ".ppt"]
    print(f"Jami konvertatsiya qilinishi kerak bo'lgan .ppt fayllar: {len(ppt_files)} ta")
    
    if not ppt_files:
        return {}

    converted_map = {}
    ppt_app = None
    try:
        ppt_app = win32com.client.Dispatch("PowerPoint.Application")
    except Exception as e:
        print(f"PowerPoint COM ishga tushmadi: {e}")
        return {}

    for idx, ppt_path in enumerate(ppt_files, 1):
        target_pptx = TEMP_CONV / f"{ppt_path.stem}.pptx"
        if target_pptx.exists() and target_pptx.stat().st_size > 1000:
            converted_map[ppt_path] = target_pptx
            continue
        try:
            print(f"  [{idx}/{len(ppt_files)}] Konvertatsiya: {ppt_path.name}")
            pres = ppt_app.Presentations.Open(str(ppt_path.resolve()), WithWindow=False)
            pres.SaveAs(str(target_pptx.resolve()), 24) # 24 = ppSaveAsOpenXMLPresentation (.pptx)
            pres.Close()
            converted_map[ppt_path] = target_pptx
        except Exception as err:
            print(f"    Xatolik ({ppt_path.name}): {err}")
            
    if ppt_app:
        try:
            ppt_app.Quit()
        except Exception:
            pass

    print(f"Muvaffaqiyatli konvertatsiya qilindi: {len(converted_map)}/{len(ppt_files)}")
    return converted_map

def get_pptx_metrics(pptx_path):
    try:
        prs = Presentation(pptx_path)
        slides_c = len(prs.slides)
        total_words = 0
        texts = []
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        t = p.text.strip()
                        if t:
                            texts.append(t)
                            total_words += len(t.split())
        return slides_c, total_words, texts
    except Exception as e:
        return 0, 0, []

def file_hash(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()

def main():
    fix_extensionless_files()
    converted_map = convert_ppt_to_pptx()

    print("\n--- 3. Barcha taqdimotlarni tahlil qilish va saralash ---")
    all_pres = []
    for sub in sorted(ECON_DIR.iterdir()):
        if sub.is_dir() and sub.name not in ["01_Sara_Slaydlar_ECON", "99_Karantin_Tekshirish_Uchun"]:
            for f in sub.glob("*"):
                if f.is_file() and f.suffix.lower() in [".ppt", ".pptx"] and not f.name.startswith("~$"):
                    all_pres.append((sub.name, f))

    print(f"Jami tahlil qilinayotgan fayllar: {len(all_pres)} ta")

    sara_list = []
    karantin_list = []
    seen_hashes = {}
    seen_stems = set()

    for idx, (sub_name, orig_path) in enumerate(all_pres, 1):
        actual_pptx = orig_path
        is_converted = False
        if orig_path.suffix.lower() == ".ppt":
            actual_pptx = converted_map.get(orig_path)
            is_converted = True

        if not actual_pptx or not actual_pptx.exists():
            karantin_list.append({
                "file": orig_path.name,
                "folder": sub_name,
                "reason": "PPT fayl XML formatga konvertatsiya bo'lmadi yoki shikastlangan",
                "slides": 0,
                "words": 0,
                "orig_path": str(orig_path)
            })
            continue

        slides_c, words_c, texts = get_pptx_metrics(actual_pptx)
        f_h = file_hash(actual_pptx)
        name_lower = orig_path.stem.lower()

        # Filtrlash mezonlari
        quarantine_reason = None

        if f_h in seen_hashes:
            quarantine_reason = f"To'liq takroriy nusxa (Dublikat: {seen_hashes[f_h]})"
        elif "kopiya" in name_lower or "копия" in name_lower or "_copy" in name_lower:
            quarantine_reason = "Nusxa/takroriy fayl (kopiya)"
        elif slides_c < 8:
            quarantine_reason = f"Slaydlar soni juda kam ({slides_c} ta < 8 ta mezon)"
        elif words_c < 150:
            quarantine_reason = f"Matn miqdori juda kam ({words_c} ta so'z < 150 ta mezon)"
        elif any(q in name_lower for q in ["viktorina", "pole_chudes", "krossvord", "igra"]):
            quarantine_reason = "O'yin yoki viktorina taqdimoti (akademik nazariyaga ega emas)"

        if quarantine_reason:
            karantin_list.append({
                "file": orig_path.name,
                "folder": sub_name,
                "reason": quarantine_reason,
                "slides": slides_c,
                "words": words_c,
                "orig_path": str(orig_path),
                "actual_pptx": str(actual_pptx)
            })
        else:
            seen_hashes[f_h] = orig_path.name
            sara_list.append({
                "file": orig_path.name,
                "folder": sub_name,
                "slides": slides_c,
                "words": words_c,
                "orig_path": str(orig_path),
                "actual_pptx": str(actual_pptx),
                "is_converted": is_converted
            })

    print(f"\n📊 Tahlil Natijalari:")
    print(f"  ✅ Sara taqdimotlar (01_Sara_Slaydlar_ECON): {len(sara_list)} ta")
    print(f"  ⚠️ Karantin (99_Karantin_Tekshirish_Uchun): {len(karantin_list)} ta")

    # 4. Fayllarni nusxalash va joylashtirish
    print("\n--- 4. Saralangan fayllarni papkalarga xavfsiz joylashtirish ---")
    for item in sara_list:
        src = Path(item["actual_pptx"])
        # Formatni doimo .pptx qilamiz
        dst_name = Path(item["file"]).stem + ".pptx"
        dst = SARA_DIR / dst_name
        shutil.copy2(src, dst)

    for item in karantin_list:
        src = Path(item["orig_path"])
        dst = KARANTIN_DIR / src.name
        shutil.copy2(src, dst)

    # 5. Karantin hisobotini yozish
    report_md = f"""# Slaydlar_ECON — Karantin va Saralash Hisoboti

Ushbu hisobot `C:\\Users\\user\\Desktop\\Antigravity\\Soff.uz\\Slaydlar_ECON` papkasidagi barcha taqdimotlar tahlili natijalarini va saralash mezonlarini o'z ichiga oladi.
**Nol yo'qotish qoidasi (Zero Data Loss):** Hech bir fayl o'chirilmadi! Barcha filtrlangan fayllar ushbu karantin papkasida to'liq saqlanmoqda.

## Saralash Mezonlari
1. **Slaydlar soni:** Kamida 8 ta slayd.
2. **Matn hajmi:** Kamida 150 ta so'z (mazmunsiz, bo'sh slaydlar ajratildi).
3. **Dublikatlar:** Fayl heshi yoki nomi bo'yicha takroriy nusxalar (`- kopiya`).
4. **Mavzu sifati:** O'yin, krossvord, viktorinalar karantinga olindi.

## Statistika
- **Jami tahlil qilingan fayllar:** {len(all_pres)} ta
- **Sara taqdimotlar (01_Sara_Slaydlar_ECON):** {len(sara_list)} ta ({len(sara_list)/len(all_pres)*100:.1f}%)
- **Karantinga olingan fayllar:** {len(karantin_list)} ta ({len(karantin_list)/len(all_pres)*100:.1f}%)

---

## Karantinga Olingan Fayllar Ro'yxati

| № | Fayl Nomi | Asl Yo'nalishi | Slaydlar | So'zlar | Sabab |
|---|---|---|---|---|---|
"""
    for idx, item in enumerate(karantin_list, 1):
        report_md += f"| {idx} | `{item['file']}` | {item['folder']} | {item['slides']} ta | {item['words']} ta | {item['reason']} |\n"

    report_path = KARANTIN_DIR / "KARANTIN_HISOBOTI.md"
    with open(report_path, "w", encoding="utf-8") as rf:
        rf.write(report_md)
    print(f"Karantin hisoboti yaratildi: {report_path}")

    # Sara fayllar ro'yxatini json qilib saqlaymiz
    sara_manifest = SOFF_ROOT / "manifest_sara_econ.json"
    with open(sara_manifest, "w", encoding="utf-8") as sm:
        json.dump(sara_list, sm, ensure_ascii=False, indent=2)
    print(f"Sara manifest saqlandi: {sara_manifest}")

if __name__ == "__main__":
    main()
