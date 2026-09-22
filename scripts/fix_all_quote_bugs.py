import os
import glob
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"
files = sorted(glob.glob(os.path.join(folder, "*.pptx")))

def unquote_text(txt):
    if not txt or len(txt) < 3:
        return txt
    
    # Check if entire string is bugged with quotes at every even index
    cur = txt
    while len(cur) >= 3 and len(cur) % 2 == 1 and all(c == "'" for c in cur[::2]):
        cur = cur[1:-1:2]
    
    # Check line by line
    lines = cur.split('\n')
    new_lines = []
    for line in lines:
        l = line
        while len(l) >= 3 and len(l) % 2 == 1 and all(c == "'" for c in l[::2]):
            l = l[1:-1:2]
        new_lines.append(l)
    
    return '\n'.join(new_lines)

total_files_fixed = 0
total_paragraphs_fixed = 0

def process_file(fpath):
    global total_files_fixed, total_paragraphs_fixed
    fname = os.path.basename(fpath)
    prs = Presentation(fpath)
    file_modified = False
    
    def process_shape(shape):
        nonlocal file_modified
        global total_paragraphs_fixed
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                orig = p.text
                fixed = unquote_text(orig)
                if fixed != orig:
                    p.text = fixed
                    file_modified = True
                    total_paragraphs_fixed += 1
                    
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if cell.text_frame:
                        for p in cell.text_frame.paragraphs:
                            orig = p.text
                            fixed = unquote_text(orig)
                            if fixed != orig:
                                p.text = fixed
                                file_modified = True
                                total_paragraphs_fixed += 1
                                
        if shape.shape_type == 6: # Group shape
            for subshape in shape.shapes:
                process_shape(subshape)

    for slide in prs.slides:
        for shape in slide.shapes:
            process_shape(shape)
            
    if file_modified:
        prs.save(fpath)
        total_files_fixed += 1
        print(f"Fixed quotes in: {fname}")

for f in files:
    process_file(f)

print("\n" + "="*60)
print(f"Barcha fayllar tozalandi! Jami tuzatilgan fayllar: {total_files_fixed}, tuzatilgan paragraflar: {total_paragraphs_fixed}")
print("="*60)
