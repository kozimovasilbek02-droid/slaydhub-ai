import os
import glob
import re
import json
import sys
from pptx import Presentation

# Ensure stdout handles utf-8
sys.stdout.reconfigure(encoding='utf-8')

folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"
files = glob.glob(os.path.join(folder, "*.pptx"))

print(f"Total files found: {len(files)}")

cyrillic_pattern = re.compile(r'[\u0400-\u04FF]')
latin_pattern = re.compile(r'[a-zA-Z]')

watermark_keywords = [
    'antonenkova', 'allppt', 'slidesgo', 'presentationgo', 'shkola', 
    'o\'qituvchi', 'учитель', 'maktab', 'pedsovet', 'infourok', 'myshared',
    'prezentacii', 'videouroki', 'viki.rdf.ru', 'uroki.net', 'uchportal', 'kopilkaurokov'
]

results = []
total_slides = 0
pure_uzbek_count = 0
cyrillic_leak_count = 0
heavy_cyrillic_count = 0
watermark_count = 0
bad_typography_count = 0

for idx, f in enumerate(sorted(files)):
    fname = os.path.basename(f)
    try:
        prs = Presentation(f)
        slide_count = len(prs.slides)
        total_slides += slide_count
        
        file_cyr = 0
        file_lat = 0
        file_watermarks = []
        file_cyr_samples = []
        file_font_issues = []
        image_count = 0
        
        for s_idx, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if shape.shape_type == 13: # Picture
                    image_count += 1
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if not text:
                        continue
                    
                    cyr_matches = cyrillic_pattern.findall(text)
                    lat_matches = latin_pattern.findall(text)
                    
                    file_cyr += len(cyr_matches)
                    file_lat += len(lat_matches)
                    
                    if len(cyr_matches) > 3 and len(file_cyr_samples) < 5:
                        clean_samp = text.replace('\n', ' ')[:80]
                        file_cyr_samples.append(f"Slide {s_idx+1}: {clean_samp}")
                        
                    for kw in watermark_keywords:
                        if kw in text.lower():
                            file_watermarks.append({
                                "slide": s_idx + 1,
                                "keyword": kw,
                                "snippet": text[:60]
                            })
                            
                    # Check font size consistency
                    font_sizes = []
                    for p in shape.text_frame.paragraphs:
                        for r in p.runs:
                            if r.font.size:
                                font_sizes.append(r.font.size.pt)
                    if font_sizes:
                        if max(font_sizes) - min(font_sizes) > 14 and len(set(font_sizes)) > 2:
                            file_font_issues.append((s_idx+1, f"Mixed sizes: {min(font_sizes)}pt - {max(font_sizes)}pt"))
                            
        is_heavy_cyr = file_cyr > 150 and file_cyr > file_lat * 0.4
        is_cyr_leak = file_cyr > 5 and not is_heavy_cyr
        is_pure_uz = file_cyr <= 5
        
        if is_heavy_cyr:
            heavy_cyrillic_count += 1
            status = "HEAVY_CYRILLIC"
        elif is_cyr_leak:
            cyrillic_leak_count += 1
            status = "CYRILLIC_LEAK"
        else:
            pure_uzbek_count += 1
            status = "PURE_UZBEK"
            
        if file_watermarks:
            watermark_count += 1
            
        if file_font_issues:
            bad_typography_count += 1
            
        item = {
            "path": f,
            "filename": fname,
            "slides": slide_count,
            "cyrillic_chars": file_cyr,
            "latin_chars": file_lat,
            "status": status,
            "watermarks": file_watermarks,
            "cyr_samples": file_cyr_samples,
            "font_issues": len(file_font_issues),
            "images": image_count
        }
        results.append(item)
        
    except Exception as e:
        print(f"Error reading {fname}: {e}")

print("==================================================")
print("AUDIT REPORT FOR 01_Tayyor_Sara_Slaydlar_UZ")
print("==================================================")
print(f"Total Presentations: {len(results)}")
print(f"Total Slides: {total_slides}")
print(f"Pure Uzbek (Latin): {pure_uzbek_count} files")
print(f"Files with minor Cyrillic leaks: {cyrillic_leak_count} files")
print(f"Files untranslated / Heavy Cyrillic: {heavy_cyrillic_count} files")
print(f"Files with watermarks / author footers: {watermark_count} files")
print(f"Files with inconsistent typography: {bad_typography_count} files")
print("==================================================")

out_path = r"C:\Users\user\Desktop\Antigravity\Soff.uz\audit_01_tayyor_sara.json"
with open(out_path, "w", encoding="utf-8") as out:
    json.dump(results, out, ensure_ascii=False, indent=2)

print(f"\nAudit saved to {out_path}")
