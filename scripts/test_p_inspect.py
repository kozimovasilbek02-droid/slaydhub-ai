import os, re, sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"
fname = "Fizika - Elementar zarralar fizikasining paydo bo'lishi (11-sinf).pptx"
fpath = os.path.join(folder, fname)
prs = Presentation(fpath)
for s_idx, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if re.search(r'[\u0400-\u04FF]', p.text):
                    print(f"Slide {s_idx+1}: {repr(p.text)}")
