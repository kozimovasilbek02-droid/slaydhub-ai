import os
import glob
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"
files = sorted(glob.glob(os.path.join(folder, "*.pptx")))

for fpath in files:
    prs = Presentation(fpath)
    modified = False
    
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    txt = p.text
                    if set(txt.strip()).issubset({"'", "\n", " "}):
                        p.text = ""
                        modified = True
                    elif txt.startswith("'\n") or txt.endswith("\n'"):
                        p.text = txt.replace("'\n", "\n").replace("\n'", "\n")
                        modified = True
                        
    if modified:
        prs.save(fpath)
        print(f"Cleaned empty newline artifacts in: {os.path.basename(fpath)}")

print("\nAll empty artifacts cleaned!")
