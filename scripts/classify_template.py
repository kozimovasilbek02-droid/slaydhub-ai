# -*- coding: utf-8 -*-
import os
import re
from pptx import Presentation
from pptx.util import Inches

from core.academic_matcher import get_academic_matcher
matcher = get_academic_matcher()
bp = matcher.build_blueprint("Kvant kompyuterlari va asimmetrik kriptografiya", slide_count=10, language="uz")
template_path = bp["primary_pptx_path"]

prs = Presentation(template_path)
sw = prs.slide_width / Inches(1)
sh = prs.slide_height / Inches(1)

print(f"Analyzing {len(prs.slides)} slides in template: {os.path.basename(template_path)}")

def classify_slide(slide, s_idx):
    texts = []
    has_step = False
    has_timeline = False
    has_comparison = False
    has_agenda = False
    
    num_title_like = 0
    num_body_like = 0
    
    slide_title = ""
    
    for s in slide.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if not t:
                continue
            texts.append(t)
            t_low = t.lower()
            if any(k in t_low for k in ["table of contents", "agenda", "contents", "index"]):
                has_agenda = True
            if any(k in t_low for k in ["step 01", "step 1", "01.", "step infographic", "steps infographic"]):
                has_step = True
            if any(k in t_low for k in ["timeline", "2026", "2027", "2028", "2029", "2030", "roadmap"]):
                has_timeline = True
            if any(k in t_low for k in ["comparison", "vs", "versus"]):
                has_comparison = True
            if any(k in t_low for k in ["table slide"]):
                return "table_data"
            if any(k in t_low for k in ["credits", "fonts", "images", "let’s keep in touch", "contact", "thank you", "thanks"]):
                return "conclusion_sources"
                
            if s.top < Inches(2.5).emu and len(t) < 80:
                slide_title = t
                
            if len(t) < 40 and not any(c in t for c in [".", ",", ";"]):
                num_title_like += 1
            else:
                num_body_like += 1
                
    if s_idx == 0:
        return "cover"
    if has_agenda:
        return "agenda"
    if has_timeline:
        return "timeline_steps"
    if has_step:
        return "timeline_steps"
    if has_comparison:
        return "comparison_vs"
    if num_body_like >= 4 or (num_title_like >= 3 and num_body_like >= 3):
        return "cards_grid"
    if len(texts) <= 2 and any(len(t) > 60 for t in texts):
        return "theory_concept"
    if num_body_like >= 2:
        return "cards_grid"
    return "general"

for idx, slide in enumerate(prs.slides):
    ctype = classify_slide(slide, idx)
    text_count = len([s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()])
    print(f"Slide {idx+1:02d}: {ctype:<18} (text shapes: {text_count:02d})")
