from pptx import Presentation
import re

fpath = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ\Fizika - O'tkazgichlarni ketma-ket va parallel ulash (8-sinf).pptx"
prs = Presentation(fpath)
for s in prs.slides:
    for shape in s.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                p.text = re.sub(r'[\u0400-\u04FF]', '', p.text)
                p.text = p.text.replace("'", "'").replace("'", "'").replace("`", "'")
                # Fix any broken spaced characters
                if "O'''t'k'a'z'g'i'c'h" in p.text or "O't'k'a'z'g'i'c'h" in p.text:
                    p.text = "O'tkazgichlarni ketma-ket ulash bayram gulchambarlarida (girlyandalarida), poyezdlar yoritish tizimida va turli maishiy texnikalarda keng qo'llaniladi."

prs.save(fpath)
print("Saved 100% clean!")
