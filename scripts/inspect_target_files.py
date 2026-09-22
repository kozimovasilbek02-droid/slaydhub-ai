import os
import glob
from pptx import Presentation
import sys

sys.stdout.reconfigure(encoding='utf-8')

folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"
target_files = [
    "Fizika - Sila Ampera.pptx",
    "Fizika - Zarozhdenie f.pptx",
    "Fizika - УРОК ФИЗИКИ В 10 КЛАССЕ.pptx",
    "Fizika - УРОК ФИЗИКИ В 9 КЛАССЕ.pptx",
    "Geografiya - klass atmosfernoe davlenie.pptx",
    "Adabiyotshunoslik - Adabiyotda miforealizm yo'nalishi.pptx",
    "Fizika - Harorat va uning o'lchov birliklari.pptx"
]

for fname in target_files:
    fpath = os.path.join(folder, fname)
    if os.path.exists(fpath):
        prs = Presentation(fpath)
        print(f"\n====================\nFILE: {fname} ({len(prs.slides)} slides)")
        for idx in range(min(4, len(prs.slides))):
            slide = prs.slides[idx]
            texts = [sh.text_frame.text.replace('\n', ' ') for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            print(f"  Slide {idx+1}: {' | '.join(texts)[:120]}")
