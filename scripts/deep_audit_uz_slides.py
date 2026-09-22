import os
import glob
import re
import json
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"
files = sorted(glob.glob(os.path.join(folder, "*.pptx")))

cyrillic_pattern = re.compile(r'[\u0400-\u04FF]')
latin_ru_words = [
    'chto', 'kak', 'esli', 'kogda', 'potomu', 'dlya', 'posle', 'kotory', 'uchenik',
    'klass', 'urok', 'zadacha', 'zadanie', 'tema', 'cel', 'metod', 'opredelenie',
    'eto', 'ponyati', 'svoystv', 'silы', 'vremya', 'veshestv', 'protsess'
]
watermark_keywords = [
    'antonenkova', 'allppt', 'slidesgo', 'presentationgo', 'shkola', 
    'o\'qituvchi', 'учитель', 'maktab', 'pedsovet', 'infourok', 'myshared',
    'prezentacii', 'videouroki', 'viki.rdf.ru', 'uroki.net', 'uchportal', 'kopilkaurokov'
]

needs_filename_fix = []
needs_text_translation = []
needs_watermark_removal = []
clean_files = []

for f in files:
    fname = os.path.basename(f)
    prs = Presentation(f)
    
    # 1. Filename check
    has_cyr_name = bool(cyrillic_pattern.search(fname))
    has_ru_latin_name = any(re.search(rf'\b{w}\b', fname.lower()) for w in ['sila', 'klass', 'urok', 'zarozhdenie', 'atmosfernoe'])
    if has_cyr_name or has_ru_latin_name:
        needs_filename_fix.append(fname)
        
    # 2. Slide text check
    cyr_count = 0
    ru_latin_count = 0
    wm_count = 0
    image_count = 0
    
    for s_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.shape_type == 13:
                image_count += 1
            if shape.has_text_frame:
                txt = shape.text_frame.text
                cyr_matches = cyrillic_pattern.findall(txt)
                cyr_count += len(cyr_matches)
                txt_low = txt.lower()
                for w in latin_ru_words:
                    if re.search(rf'\b{w}', txt_low):
                        ru_latin_count += 1
                for wm in watermark_keywords:
                    if wm in txt_low:
                        wm_count += 1
                        
    is_dirty_text = cyr_count > 10 or ru_latin_count >= 3
    if is_dirty_text:
        needs_text_translation.append({
            "file": fname,
            "cyr_count": cyr_count,
            "ru_latin_count": ru_latin_count,
            "slides": len(prs.slides)
        })
    if wm_count > 0:
        needs_watermark_removal.append({
            "file": fname,
            "wm_count": wm_count
        })
    if not is_dirty_text and wm_count == 0 and not has_cyr_name and not has_ru_latin_name:
        clean_files.append(fname)

print(f"Total files audited: {len(files)}")
print(f"100% Clean Files: {len(clean_files)}")
print(f"Files needing Filename Translation: {len(needs_filename_fix)}")
for fn in needs_filename_fix:
    print(f"  - {fn}")

print(f"\nFiles needing Text Translation / Russian Cleanup: {len(needs_text_translation)}")
for item in needs_text_translation:
    print(f"  - {item['file']} (Cyr: {item['cyr_count']}, Ru-Lat: {item['ru_latin_count']}, Slides: {item['slides']})")

print(f"\nFiles needing Watermark / Ad Removal: {len(needs_watermark_removal)}")
for item in needs_watermark_removal:
    print(f"  - {item['file']} (Watermarks: {item['wm_count']})")
