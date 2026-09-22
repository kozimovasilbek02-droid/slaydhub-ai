import sys
from pathlib import Path
sys.path.insert(0, str(Path(r"C:\Users\user\Desktop\Antigravity\Power Point")))
from backend.core.pptx_processor import PPTXProcessor
from pptx import Presentation

test_in = r'C:\Users\user\Desktop\Antigravity\Soff.uz\Slaydlar_EN\01_Sara_Slaydlar_EN\CompilerTech_01-Intro.pptx'
test_out = r'C:\Users\user\Desktop\Antigravity\Soff.uz\Temp_PPT_Convert\test_apply.pptx'

data = PPTXProcessor.extract_presentation_data(test_in)
sample_map = {
    's1_sh103_p0': 'Kompilyatorlar 101',
    's1_sh104_p0': 'Kirish',
    's2_sh109_p0': "Kurs haqida umumiy ma'lumot",
    's2_sh110_p0': "Kirish (ushbu ma'ruza)",
    's2_sh110_p1': "Kompilyator frontend qismi"
}

PPTXProcessor.apply_translations_and_export(
    original_pptx_path=test_in,
    translations_map=sample_map,
    output_pptx_path=test_out,
    auto_fit=True,
    target_script='latin',
    clean_watermarks=False
)

prs = Presentation(test_out)
for idx, s in enumerate(list(prs.slides)[:2]):
    print(f"Slide {idx+1}:")
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if p.text.strip():
                    print('  P:', repr(p.text.strip()))
