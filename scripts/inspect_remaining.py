import os, glob, re, sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')
folder = r"C:\Users\user\Desktop\Antigravity\Soff.uz\01_Tayyor_Sara_Slaydlar_UZ"

wm_files = [
    "Biologiya - Oqilona foydalanish.pptx",
    "Fizika - Eruditlar tanlovi (9-sinf).pptx",
    "Fizika - Fizika va matematika bilimdonlari bellashuvi (7-sinf).pptx",
    "Fizika - Harakat tezligi trenajeri (9-sinf).pptx",
    "Fizika - O'tkazgichlarni ketma-ket va parallel ulash (8-sinf).pptx",
    "Informatika - Scratch muhitida loyiha yaratish.pptx",
    "Iqtisodiyot - O'tmish va Hozirgi Kunning Elektr Mutaxassisligi Sahifalari.pptx",
    "Matematika - Iskandariyalik Yevklid hayoti va merosi.pptx",
    "Tarix - Yevropa madaniy makoni va Rus madaniyati (6-sinf).pptx",
    "Tarjimashunoslik - Tarjima fanining shakllanishi va taraqqiyoti.pptx"
]

watermark_keywords = [
    'antonenkova', 'allppt', 'slidesgo', 'presentationgo', 'shkola', 
    'o\'qituvchi', 'учитель', 'maktab', 'pedsovet', 'infourok', 'myshared',
    'prezentacii', 'videouroki', 'viki.rdf.ru', 'uroki.net', 'uchportal', 'kopilkaurokov'
]

print("--- WATERMARKS INSPECTION ---")
for wf in wm_files:
    fpath = os.path.join(folder, wf)
    if not os.path.exists(fpath):
        continue
    prs = Presentation(fpath)
    print(f"\n=== {wf} ===")
    for s_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text
                for wm in watermark_keywords:
                    if wm in txt.lower():
                        print(f"  [Slide {s_idx+1}] Matched '{wm}': {repr(txt.strip()[:100])}")

print("\n--- TEXT TRANSLATION INSPECTION ---")
cyr_files = [
    "Adabiyotshunoslik - Adabiyotda miforealizm yo'nalishi.pptx",
    "Fizika - Elektr maydoni ishi va potensiali (10-sinf).pptx",
    "Fizika - Elementar zarralar fizikasining paydo bo'lishi (11-sinf).pptx",
    "Fizika - Energiya saqlanish qonuni bo'yicha masalalar (10-sinf).pptx",
    "Fizika - Ichki energiya va uni o'zgartirish usullari (10-sinf).pptx",
    "Fizika - Moddiy nuqta va sanoq sistemasi (9-sinf).pptx",
    "Geografiya - Atmosfera bosimi va barometrlar (7-sinf).pptx",
    "Iqtisodiyot - Xodimlar motivatsiyasi.pptx",
    "Menejment - Nizolar va stresslarni boshqarish.pptx",
    "Tilshunoslik - Qiyosiy lingvomadaniyatshunoslik asoslari.pptx",
    "Tilshunoslik - Til manzarasida german tillari va tasnifi.pptx"
]

cyrillic_pattern = re.compile(r'[\u0400-\u04FF]+')
for cf in cyr_files:
    fpath = os.path.join(folder, cf)
    if not os.path.exists(fpath):
        continue
    prs = Presentation(fpath)
    print(f"\n=== {cf} ===")
    for s_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text
                cyrs = cyrillic_pattern.findall(txt)
                if cyrs:
                    print(f"  [Slide {s_idx+1}] Cyrillic: {cyrs} in {repr(txt.strip()[:80])}")
