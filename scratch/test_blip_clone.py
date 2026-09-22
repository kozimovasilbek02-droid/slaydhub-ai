import sys
import os
import io
from pptx import Presentation

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from core.splicer_guardian import SplicerGuardian
from core.slide_preview import get_presentation_previews

src_prs = Presentation("output/5m_presentation_precision.pptx")
target_prs = Presentation()
src_slide = src_prs.slides[0]
target_slide = target_prs.slides.add_slide(target_prs.slide_layouts[6])

src_shp = [s for s in src_slide.shapes if s.name == "Picture 4"][0]
print("Found Picture 4, cloning...")
cloned = SplicerGuardian.safe_clone_shape_with_rels(src_shp._element, src_slide.part, target_slide.part)
target_slide.shapes._spTree.append(cloned)
SplicerGuardian.assign_unique_shape_ids(target_slide)

out_file = "scratch/test_blip_cloned.pptx"
target_prs.save(out_file)
print(f"Saved {out_file} successfully!")

# Check if target_prs package actually contains the image part in the zip!
import zipfile
with zipfile.ZipFile(out_file, 'r') as z:
    media_files = [f for f in z.namelist() if 'media/' in f]
    print("Media files in saved PPTX zip:", media_files)
    assert len(media_files) > 0, "No media files found in target zip! Image was not copied into target package!"

# Render preview via PowerPoint COM
previews = get_presentation_previews(out_file)
print("Rendered preview:", previews)
print("SUCCESS!")
