import glob
from pptx import Presentation

for f in glob.glob("output/**/*.pptx", recursive=True):
    try:
        prs = Presentation(f)
        for s_idx, s in enumerate(prs.slides):
            for shp in s.shapes:
                blips = shp._element.xpath('.//*[local-name()="blip"]')
                if blips:
                    print("Found blip in:", f, "Slide:", s_idx, "Shape:", shp.name)
                    exit(0)
    except Exception:
        pass
print("None found")
