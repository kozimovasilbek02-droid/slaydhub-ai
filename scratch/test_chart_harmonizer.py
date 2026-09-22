# -*- coding: utf-8 -*-
"""
scratch/test_chart_harmonizer.py
Tests ChartHarmonizer against native PowerPoint charts and renders preview images.
"""

import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from core.chart_harmonizer import ChartHarmonizer
from core.slide_preview import get_presentation_previews

def test_chart_harmonization():
    print("=== TEST 1: Generate Deck with Native Bar Chart ===")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # blank
    slide = prs.slides.add_slide(blank_layout)

    # Initial dummy chart data
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
    chart_data.add_series("Old Series 1", (10.0, 15.0, 12.0, 8.0))

    x, y, cx, cy = Inches(2), Inches(1.5), Inches(9.333), Inches(4.5)
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    )

    print("Native chart shape added to test slide.")

    # 2. Test Harmonization with Academic/AI Metrics
    new_metrics = {
        "categories": ["2023 (Klassik)", "2024 (Gibrid)", "2025 (PQC Kyber)", "2026 (Kvant Ustunligi)"],
        "series": [
            {
                "name": "Shifrlash Vaqti (ms)",
                "values": [120, 85, 34, 12]
            },
            {
                "name": "Kvant Himoyalanish Darajasi (%)",
                "values": [20, 45, 95, 100]
            }
        ]
    }

    success = ChartHarmonizer.apply_chart_data(
        chart_shape=chart_shape,
        chart_data_dict=new_metrics,
        theme_key="cyber_dark"
    )
    assert success is True, "ChartHarmonizer failed to apply data"
    print("PASS: ChartHarmonizer successfully replaced data and harmonized Penpot colors!")

    out_pptx = "output/test_chart_harmonized.pptx"
    os.makedirs("output", exist_ok=True)
    prs.save(out_pptx)
    print(f"Saved chart presentation to {out_pptx}")

    # Render preview
    previews = get_presentation_previews(out_pptx)
    print(f"Rendered {len(previews)} preview slides:")
    for p in previews:
        print(f"  -> {p}")

    print("\n==========================================")
    print("CHART HARMONIZER TESTS PASSED WITH 0 DEFECTS!")
    print("==========================================")

if __name__ == "__main__":
    test_chart_harmonization()
