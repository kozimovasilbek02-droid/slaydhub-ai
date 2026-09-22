# -*- coding: utf-8 -*-
"""
scratch/test_spatial_cleaner.py
Unit test suite for SpatialProximityCleaner (Round 9):
- Tests pruning of orphaned badge circles near unassigned cards
- Tests pruning of dead connector arrows
- Tests preservation of badges belonging to assigned cards
- Tests SafeSize protection for major backdrop panels
"""

import sys
import os
import io
from pathlib import Path

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from core.spatial_proximity_cleaner import SpatialProximityCleaner

def run_tests():
    print("🧪 Running SpatialProximityCleaner Unit Tests (Round 9)...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1. Large Slide Background Panel (SafeSize > 25% area)
    bg_panel = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(12.333), Inches(6.5))
    bg_panel.fill.solid()
    bg_panel.fill.fore_color.rgb = RGBColor(248, 250, 252) # Light slate

    # 2. Step 1: Assigned Card (with badge circle)
    card1 = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2.5), Inches(3), Inches(3))
    card1.text_frame.text = "Bosqich 1: Kvant Tahlili\nAsosiy ma'lumotlar to'plami."
    badge1 = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.OVAL, Inches(2), Inches(1.8), Inches(1), Inches(1))
    badge1.text_frame.text = "01"

    # 3. Step 2: Unassigned Card (to be pruned)
    card2 = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6), Inches(2.5), Inches(3), Inches(3))
    card2.text_frame.text = "Bosqich 2: Dummy Placeholder\nHech qanday ma'lumot yo'q."
    badge2 = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.OVAL, Inches(7), Inches(1.8), Inches(1), Inches(1)) # Orphaned badge
    badge2.text_frame.text = "02"

    # 4. Dead Connector Arrow between Card 1 and Card 2
    arrow12 = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RIGHT_ARROW, Inches(4.3), Inches(3.5), Inches(1.4), Inches(0.6))

    initial_count = len(slide.shapes)
    print(f"Boshlang'ich shakllar soni: {initial_count}")
    assert initial_count == 6

    # Simulate card2 being pruned by SemanticShapeCleaner
    pruned_bb = SpatialProximityCleaner.get_bbox(card2)
    assigned_bb = SpatialProximityCleaner.get_bbox(card1)
    SpatialProximityCleaner.prune_shape(card2) # Card 2 text shape pruned

    # Run Spatial Proximity Cleaner
    pruned_artifacts = SpatialProximityCleaner.prune_orphaned_artifacts(
        slide=slide,
        pruned_bboxes=[pruned_bb],
        assigned_bboxes=[assigned_bb],
        assigned_shape_ids={id(card1)},
        sw_pt=13.333 * 72.0,
        sh_pt=7.5 * 72.0
    )
    print(f"Fazoviy proksimiti orqali qirqilgan qoldiqlar soni: {pruned_artifacts}")
    assert pruned_artifacts == 2, f"Expected 2 artifacts pruned (badge2 and arrow12), got {pruned_artifacts}"

    remaining_shapes = list(slide.shapes)
    print(f"Tozalashdan keyingi shakllar soni: {len(remaining_shapes)}")
    assert len(remaining_shapes) == 3, f"Expected 3 shapes remaining (bg_panel, card1, badge1), got {len(remaining_shapes)}"

    # Verify background panel is preserved
    assert remaining_shapes[0] == bg_panel, "Background panel must be preserved!"

    # Verify assigned card1 and its badge1 are preserved
    remaining_names = [s.name for s in remaining_shapes]
    print(f"Qolgan shakllar nomlari: {remaining_names}")
    assert any("Oval" in s.name and "01" in s.text_frame.text for s in remaining_shapes if s.has_text_frame), "Badge 1 must be preserved!"

    # Save output deck
    os.makedirs("output", exist_ok=True)
    out_path = "output/test_spatial_cleaned.pptx"
    prs.save(out_path)
    print(f"✅ Muvaffaqiyatli saqlandi: {out_path}")
    print("🎉 Barcha testlar 100% PASS!")

if __name__ == "__main__":
    run_tests()
