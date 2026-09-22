# -*- coding: utf-8 -*-
"""
scratch/test_security_purger.py
Unit test suite verifying Claude Code's Round 7 challenge:
- Invisible ghost action buttons
- Spammy promotional hyperlinks
- Legitimate shapes preservation
"""

import sys
import os
import io
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.security_purger import SecurityPurger

def run_tests():
    print("🧪 Running SecurityPurger Unit Tests (Round 7)...")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1. Fullscreen ghost action button (PresentationGO style)
    ghost1 = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(7.5))
    ghost1.fill.background()
    ghost1.line.fill.background()
    ghost1.click_action.hyperlink.address = "https://www.presentationgo.com/free-template"

    # 2. Corner ghost button (SlideEgg style)
    ghost2 = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(6.5), Inches(1.2), Inches(0.8))
    ghost2.fill.background()
    ghost2.line.fill.background()
    ghost2.click_action.hyperlink.address = "https://slideegg.com/promo-deal"

    # 3. Textbox with valid text and spam hyperlink in run
    legit_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1.5))
    p = legit_box.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Q1 2026 Moliyaviy Tahlil: "
    r2 = p.add_run()
    r2.text = "Batafsil ma'lumot (SlidesCarnival)"
    r2.hyperlink.address = "https://slidescarnival.com/track?id=451"

    # 4. Massive transparent blocker overlay without hyperlink
    blocker = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(8), Inches(5))
    blocker.fill.background()
    blocker.line.fill.background()

    # 5. Legitimate colored card with content (MUST PRESERVE)
    card = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(3), Inches(4), Inches(2))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(56, 189, 248) # Sky blue
    card.line.color.rgb = RGBColor(255, 255, 255)
    card.text_frame.text = "Xavfsizlik moduli to'liq ishlamoqda"

    initial_shape_count = len(slide.shapes)
    print(f"Boshlang'ich shakllar soni: {initial_shape_count}")
    assert initial_shape_count == 5, f"Expected 5 shapes, got {initial_shape_count}"

    # Run Security Purge
    stats = SecurityPurger.purge_presentation_security(prs)
    print(f"Audit Natijasi: {stats}")

    # Assertions
    assert stats["ghosts_pruned"] >= 3, f"Expected at least 3 ghost shapes pruned, got {stats['ghosts_pruned']}"
    assert stats["hyperlinks_defused"] >= 1, f"Expected at least 1 spam hyperlink defused, got {stats['hyperlinks_defused']}"

    # Verify remaining shapes on slide
    remaining_shapes = list(slide.shapes)
    print(f"Tozalashdan keyingi shakllar soni: {len(remaining_shapes)}")
    assert len(remaining_shapes) == 2, f"Expected 2 shapes remaining (legit_box and card), got {len(remaining_shapes)}"

    # Verify legit_box content and that its spam hyperlink was defused
    assert "Q1 2026 Moliyaviy Tahlil" in legit_box.text_frame.text
    run_hlinks = legit_box._element.xpath('.//a:hlinkClick')
    print(f"legit_box ichidagi havolalar soni: {len(run_hlinks)}")
    assert len(run_hlinks) == 0, "Spam hyperlink in run should have been defused!"

    # Verify legitimate card is intact
    assert "Xavfsizlik moduli to'liq ishlamoqda" in card.text_frame.text

    # Save test output
    os.makedirs("output", exist_ok=True)
    test_out = "output/test_security_purged.pptx"
    prs.save(test_out)
    print(f"✅ Muvaffaqiyatli saqlandi: {test_out}")
    print("🎉 Barcha testlar 100% PASS!")

if __name__ == "__main__":
    run_tests()
