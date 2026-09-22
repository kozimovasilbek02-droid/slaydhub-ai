# -*- coding: utf-8 -*-
"""
scratch/test_theme_harmonizer.py
Tests ThemeHarmonizer against real PowerPoint templates and exports slide preview images.
"""

import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from pptx import Presentation
from core.theme_harmonizer import ThemeHarmonizer
from core.penpot_themes import PenpotThemeManager

from core.slide_preview import get_presentation_previews

def test_harmonizer():
    candidates = [
        "output/action_plan_s2_flawless.pptx",
        "output/test_ideal_native.pptx",
        "output/test_morph_mc.pptx",
        "test_arch.pptx"
    ]
    template_path = None
    for c in candidates:
        if os.path.exists(c):
            template_path = c
            break

    if not template_path:
        print("No candidates found!")
        return


    print(f"Testing ThemeHarmonizer on: {template_path}")
    prs = Presentation(template_path)
    sw_pt = prs.slide_width.pt
    sh_pt = prs.slide_height.pt

    themes_to_test = ["quantum_deep", "academic_warm", "cyber_dark"]
    output_files = []

    for theme_key in themes_to_test:
        prs_copy = Presentation(template_path)
        total_harmonized = 0
        for idx, slide in enumerate(prs_copy.slides):
            count = ThemeHarmonizer.harmonize_slide_accents(
                slide, theme_key=theme_key, slide_width_pt=sw_pt, slide_height_pt=sh_pt
            )
            total_harmonized += count

        out_pptx = f"output/test_harmonized_{theme_key}.pptx"
        os.makedirs("output", exist_ok=True)
        prs_copy.save(out_pptx)
        print(f"Theme '{theme_key}': {total_harmonized} accent shapes harmonized. Saved to {out_pptx}")
        output_files.append((theme_key, out_pptx))

    # Render preview for the first one
    first_theme, first_file = output_files[0]
    preview_paths = get_presentation_previews(first_file)
    print(f"Rendered previews for {first_theme}: {preview_paths}")


if __name__ == "__main__":
    test_harmonizer()
