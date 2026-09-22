# -*- coding: utf-8 -*-
"""
scratch/test_contrast_harmonizer.py
Unit tests for ContrastHarmonizer (Round 10).
Validates:
1. WCAG 2.1 Relative Luminance & Contrast Ratio calculation.
2. Auto-inversion of text color on light vs dark recolored shapes.
3. Penpot theme dynamic text palette extraction.
4. XML text frame contrast inoculation.
"""

import os
import sys
import io

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from core.contrast_harmonizer import ContrastHarmonizer


def test_contrast_harmonizer():
    print("🧪 Running ContrastHarmonizer Unit Tests (Round 10)...")

    # 1. Luminance calculations
    lum_black = ContrastHarmonizer.calculate_luminance("#000000")
    lum_white = ContrastHarmonizer.calculate_luminance("#FFFFFF")
    print(f"Luminance Black: {lum_black:.4f}, White: {lum_white:.4f}")
    assert abs(lum_black - 0.0) < 0.001, "Black luminance must be ~0"
    assert abs(lum_white - 1.0) < 0.001, "White luminance must be ~1"

    # 2. Contrast Ratios
    cr_max = ContrastHarmonizer.calculate_contrast_ratio("#000000", "#FFFFFF")
    cr_min = ContrastHarmonizer.calculate_contrast_ratio("#FFFFFF", "#FFFFFF")
    print(f"Contrast Ratio (Black vs White): {cr_max:.2f}:1")
    print(f"Contrast Ratio (White vs White): {cr_min:.2f}:1")
    assert abs(cr_max - 21.0) < 0.1, "Max contrast ratio must be 21:1"
    assert abs(cr_min - 1.0) < 0.01, "Identical colors ratio must be 1:1"

    # 3. High Contrast Text Decisions
    text_on_dark = ContrastHarmonizer.get_high_contrast_text_color("0F172A", "tech_minimal")
    text_on_white = ContrastHarmonizer.get_high_contrast_text_color("FFFFFF", "tech_minimal")
    text_on_yellow = ContrastHarmonizer.get_high_contrast_text_color("F59E0B", "tech_minimal")
    text_on_navy = ContrastHarmonizer.get_high_contrast_text_color("0A3B75", "tech_minimal")

    print(f"Text on Slate 900 (#0F172A): #{text_on_dark} (expected FFFFFF)")
    print(f"Text on White (#FFFFFF): #{text_on_white} (expected dark text)")
    print(f"Text on Gold (#F59E0B): #{text_on_yellow} (expected dark text)")
    print(f"Text on Deep Navy (#0A3B75): #{text_on_navy} (expected FFFFFF)")

    assert text_on_dark == "FFFFFF"
    assert text_on_white == "0F172A"
    assert text_on_yellow == "0F172A"
    assert text_on_navy == "FFFFFF"

    # 4. Theme Text Palette Extraction
    colors_emerald = ContrastHarmonizer.get_theme_text_colors("emerald_science", is_dark_bg=False)
    colors_cyber = ContrastHarmonizer.get_theme_text_colors("cyber_dark", is_dark_bg=True)

    print(f"Emerald Science Title Color: {ContrastHarmonizer.rgb_to_hex(colors_emerald['title'])}")
    print(f"Emerald Science Lead-In Color: {ContrastHarmonizer.rgb_to_hex(colors_emerald['lead_in'])}")
    print(f"Cyber Dark Title Color: {ContrastHarmonizer.rgb_to_hex(colors_cyber['title'])}")

    assert ContrastHarmonizer.rgb_to_hex(colors_emerald["title"]) == "064E3B"
    assert ContrastHarmonizer.rgb_to_hex(colors_emerald["lead_in"]) == "059669"
    assert ContrastHarmonizer.rgb_to_hex(colors_cyber["title"]) == "FFFFFF"

    # 5. In-Memory Shape Text Frame Inoculation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    shape.text = "STEP 01: KVANT ALGORITMI"

    # Inoculate with dark navy background -> Text MUST turn to white
    ContrastHarmonizer.ensure_text_contrast(shape, "0F172A", "tech_minimal")
    run_color = shape.text_frame.paragraphs[0].runs[0].font.color.rgb
    print(f"Inoculated Text Run Color on #0F172A: {ContrastHarmonizer.rgb_to_hex(run_color)}")
    assert ContrastHarmonizer.rgb_to_hex(run_color) == "FFFFFF"

    # Inoculate with pure white background -> Text MUST turn to theme dark #0F172A
    ContrastHarmonizer.ensure_text_contrast(shape, "FFFFFF", "tech_minimal")
    run_color2 = shape.text_frame.paragraphs[0].runs[0].font.color.rgb
    print(f"Inoculated Text Run Color on #FFFFFF: {ContrastHarmonizer.rgb_to_hex(run_color2)}")
    assert ContrastHarmonizer.rgb_to_hex(run_color2) == "0F172A"

    os.makedirs("output", exist_ok=True)
    prs.save("output/test_contrast_inoculated.pptx")
    print("✅ Muvaffaqiyatli saqlandi: output/test_contrast_inoculated.pptx")
    print("🎉 Barcha ContrastHarmonizer testlari 100% PASS!\n")


if __name__ == "__main__":
    test_contrast_harmonizer()
