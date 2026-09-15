import os
import argparse
from pptx import Presentation
from pptx.util import Inches

def convert_images_to_slides(input_folder=".", output_file="presentation.pptx", aspect_ratio="16:9"):
    prs = Presentation()
    
    # Set Slide Dimensions
    if aspect_ratio == "16:9":
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    else: # 4:3
        prs.slide_width = Inches(10.0)
        prs.slide_height = Inches(7.5)
        
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    supported_exts = (".png", ".jpg", ".jpeg", ".webp", ".bmp")
    
    if os.path.isfile(input_folder):
        image_files = [input_folder]
    else:
        image_files = sorted([
            os.path.join(input_folder, f) 
            for f in os.listdir(input_folder) 
            if f.lower().endswith(supported_exts)
        ])

    if not image_files:
        print(f"No image files found in {input_folder}")
        return

    print(f"Found {len(image_files)} image(s). Creating 1:1 PowerPoint presentation ({aspect_ratio})...")

    for idx, img_path in enumerate(image_files, start=1):
        print(f"  Adding Slide {idx}: {os.path.basename(img_path)}")
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            img_path, 
            Inches(0), 
            Inches(0), 
            width=prs.slide_width, 
            height=prs.slide_height
        )

    os.makedirs(os.path.dirname(os.path.abspath(output_file)), exist_ok=True)
    prs.save(output_file)
    print(f"SUCCESS: Created presentation at {output_file}")
    return output_file

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Convert images to full-bleed 1:1 PowerPoint slides")
    parser.add_argument("input_folder", nargs="?", default=".", help="Folder containing images or single image path")
    parser.add_argument("-o", "--output", default="output/full_slides.pptx", help="Output .pptx path")
    parser.add_argument("--ratio", choices=["16:9", "4:3"], default="16:9", help="Slide aspect ratio")
    args = parser.parse_args()
    convert_images_to_slides(args.input_folder, args.output, args.ratio)
