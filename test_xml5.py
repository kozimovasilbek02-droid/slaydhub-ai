from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
shape = slide.shapes.title
shape.text = "Hello Ozbekiston"

# Ensure rPr exists
tf = shape.text_frame
r = tf.paragraphs[0].runs[0]
rPr = r._r.get_or_add_rPr()

latin = rPr.find('./a:latin')
if latin is None:
    latin = OxmlElement('a:latin')
    rPr.insert(0, latin)
latin.set('typeface', 'Montserrat')

cs = rPr.find('./a:cs')
if cs is None:
    cs = OxmlElement('a:cs')
    rPr.insert(1, cs)
cs.set('typeface', 'Montserrat')

print(rPr.xml)
