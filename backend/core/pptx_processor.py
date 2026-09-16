# -*- coding: utf-8 -*-
import os
import re
import io
import zipfile
from typing import Dict, Any, List, Optional
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

def safe_load_presentation(pptx_path: str) -> Presentation:
    """
    Safely loads a Presentation, auto-patching template.main+xml content types
    (common in PresentationGO, SlidesMania, or POTX templates).
    """
    try:
        return Presentation(pptx_path)
    except ValueError as e:
        if "template.main+xml" in str(e):
            with zipfile.ZipFile(pptx_path, "r") as zin:
                buf = io.BytesIO()
                with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zout:
                    for item in zin.infolist():
                        data = zin.read(item.filename)
                        if item.filename == "[Content_Types].xml":
                            data = data.replace(
                                b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                                b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
                            )
                        zout.writestr(item, data)
                buf.seek(0)
            return Presentation(buf)
        raise


from backend.core.transliteration import ensure_script
from backend.core.font_manager import font_manager
from backend.core.thumbnail_generator import ThumbnailGenerator
from backend.core.image_translator import ImageTranslator

def sanitize_control_chars(text: str) -> str:
    if not text:
        return ""
    t = re.sub(r'_x[0-9a-fA-F]{4}_', ' ', text)
    t = re.sub(r'[\u0001-\u0008\u000b\u000c\u000e-\u001f\u007f]', ' ', t)
    t = re.sub(r'[ \t]+', ' ', t)
    return t.strip()

class PPTXProcessor:
    WATERMARK_PATTERNS = [
        r"presentationgo", r"ypppt", r"slidescarnival", r"slidesmania",
        r"slidesgo", r"poweredtemplate", r"slidemodel", r"designed with",
        r"free templates?", r"questions or need help", r"visit our faq",
        r"更多精品", r"ppt模板", r"ppt背景", r"by:\s*", r"\.com",
        r"51ppt", r"优品ppt", r"free-powerpoint-templates-design", r"freeppt",
        r"\ballppt(\.com)?\b", r"free\s*ppt(\s*templates?)?",
        r"free\s*powerpoint\s*templates?",
        r"clean\s+text\s+slide\s+for\s+your\s+presentation",
        r"allppt\s+(layout|tartibi)",
        r"http[s]?://\S*allppt\S*",
        r"http[s]?://\S*free-powerpoint-templates\S*",
        r"http[s]?://\S*presentationgo\S*",
        r"bepul\s+ppt\s+shablonlar",
        r"\b(учитель|преподаватель|автор|составитель|разработчик|выполнил[а]?|подготовил[а]?|o['’`]?qituvchi(si)?|muallif(i)?|tayyorladi|bajaruvchi|teacher|author|prepared\s+by)\b",
        r"\b(школ[аые]|сош|гимнази[яи]|лице[йи]|колледж|институт|университет|доу|детский\s*сад|maktab(i)?|litsey|kollej|bog['’`]?cha|school|lyceum|college|university)\b",
        r"\b(истори[яи]|математик[аи]|биологи[яи]|физик[аи]|литератур[аы]|geografiy[ai]|tarix|ona\s+tili|adabiyot)\s+(o['’`]?qituvchisi|учитель|преподаватель)\b",
        r"\b(infourok|multiurok|nsportal|kopilkaurokov|urok\.1sept|videouroki|pedsovet|myshared|ppt4web|zanimatika)\b",
        r"[А-ЯA-Z]\.\s*[А-ЯA-Z]\.\s*[А-ЯA-Zа-яa-z]{3,}"
    ]

    AD_SLIDE_PATTERNS = [
        r"slidescarnival", r"slidesgo", r"presentationgo", r"slidesmania",
        r"poweredtemplate", r"slidemodel", r"this presentation template is free",
        r"this template is free for everyone", r"ushbu taqdimot shabloni",
        r"uses the following free fonts", r"instructions for use",
        r"fonts & colors used", r"alternative resources", r"happy designing",
        r"ijodingizga zafarlar", r"credits\s*:", r"visit slidescarnival",
        r"visit slidesgo", r"更多精品", r"ppt模板", r"ppt背景", r"51ppt", r"优品ppt", r"ypppt",
        r"minnatdorchilik", r"pexels, pixabay", r"terms of use",
        r"editable icons", r"free icons", r"customizable icons", r"free fonts online",
        r"fully editable shapes?", r"fully editable icon", r"icon sets?:?\s*[a-z]?",
        r"png\s+images?", r"place\s+your\s+picture", r"place\s+your\s+image",
        r"vector\s+icons?", r"icon\s+pack", r"free\s+vector"
    ]

    @staticmethod
    def _is_real_text(text: str) -> bool:
        if not text:
            return False
        t = sanitize_control_chars(text)
        if not t:
            return False
        if len(t) == 1 and ord(t[0]) >= 0xE000:
            return False
        return True

    @staticmethod
    def _is_watermark_text(txt: str) -> bool:
        if not txt:
            return False
        t = sanitize_control_chars(txt).lower()
        return any(re.search(pat, t) for pat in PPTXProcessor.WATERMARK_PATTERNS)

    @staticmethod
    def _is_watermark_recursive(shape, depth: int = 0, max_depth: int = 10) -> bool:
        if depth >= max_depth:
            return False
        # 1. Matnli reklama tekshiruvi
        if shape.has_text_frame and PPTXProcessor._is_watermark_text(shape.text_frame.text):
            return True
        # 2. XML tekshiruvi (Group, Picture descr, r:embed havolalar va logo rasmlari)
        try:
            xml = shape._element.xml
            if re.search(r"(allppt|free-powerpoint-templates|presentationgo|slidesgo|slidestime|slidegeeks|templatewise|slideteam|poweredtemplate)", xml, re.I):
                return True
        except Exception:
            pass
        # 3. Guruh ichidagi shakllarni rekursiv tekshirish
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                for sub in shape.shapes:
                    if PPTXProcessor._is_watermark_recursive(sub, depth + 1, max_depth):
                        return True
            except Exception:
                pass
        return False

    @staticmethod
    def _is_ad_slide(slide, slide_index: int = 2, total_slides: int = 10) -> bool:
        """Taqdimot oxiridagi SlidesCarnival/Slidesgo/Freepik reklama va resurs slaydlarini aniqlash."""
        # Birinchi yoki ikkinchi slayd (muqova/reja) hech qachon reklama slaydi sifatida o'chirilmaydi
        if slide_index <= 2:
            return False
        # Kontent slaydlarini tasodifiy o'chirib yubormaslik uchun faqat oxirgi qismdagi (65%+) slaydlar tekshiriladi
        if total_slides > 5 and slide_index < total_slides * 0.65:
            return False

        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                texts.append(sanitize_control_chars(sh.text_frame.text))
            elif sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            texts.append(sanitize_control_chars(cell.text_frame.text))
        combined = " ".join(texts).lower()
        if not combined.strip():
            return False

        # 1. Aniq reklama va resurs shablonlari
        for pat in PPTXProcessor.AD_SLIDE_PATTERNS:
            if re.search(pat, combined, re.I):
                return True

        # 2. Resurslar, shriftlar va ranglar kombinatsiyasi (shablon sahifasi)
        has_shrift = any(w in combined for w in ["shrift", "font", "typography"])
        has_rang = any(w in combined for w in ["rang", "color", "palette"])
        has_resurs = any(w in combined for w in ["resurs", "resource", "canva", "slidescarnival", "slidesgo"])
        if (has_shrift and has_rang) and has_resurs:
            return True

        # 3. Minnatdorchilik va foto kreditlari kombinatsiyasi
        has_credit = any(w in combined for w in ["minnatdorchilik", "credits", "tashakkur", "acknowledgement", "attribution"])
        has_provider = any(w in combined for w in ["pexels", "pixabay", "unsplash", "freepik", "flaticon", "slidescarnival", "slidesgo"])
        if has_credit and has_provider:
            return True

        return False

    @staticmethod
    def _delete_slide(prs: Presentation, index: int):
        """Taqdimotdan berilgan indeksdagi slaydni butunlay xavfsiz o'chirish."""
        try:
            sldIdLst = prs.slides._sldIdLst
            sldId = sldIdLst[index]
            rId = sldId.rId
            prs.part.drop_rel(rId)
            sldIdLst.remove(sldId)
        except Exception:
            pass

    @staticmethod
    def clean_presentation_watermarks(prs: Presentation) -> int:
        removed = 0
        sw = prs.slide_width
        sh_h = prs.slide_height

        # 1. Taqdimot oxiridagi barcha reklama/resurs slaydlarini to'liq o'chirish
        # Agar taqdimotda "Thank you" / "E'tiboringiz uchun rahmat" slaydi bo'lsa, undan keyingi barcha shablon ilovalarini o'chirish
        total_s = len(prs.slides)
        thank_you_idx = -1
        for idx in range(total_s - 1, max(0, int(total_s * 0.5)), -1):
            s = prs.slides[idx]
            txts = [p.text.strip().lower() for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs if p.text.strip()]
            comb = " ".join(txts)
            if any(re.search(p, comb) for p in [r"\bthank\s*you\b", r"\be['’`]?tiboringiz\s+uchun\s+rahmat\b", r"\bspasibo\s+za\s+vnimanie\b"]):
                thank_you_idx = idx
                break

        if thank_you_idx != -1 and thank_you_idx < len(prs.slides) - 1:
            for s_idx in range(len(prs.slides) - 1, thank_you_idx, -1):
                PPTXProcessor._delete_slide(prs, s_idx)
                removed += 1

        slide_count = len(prs.slides)
        for s_idx in range(slide_count - 1, 0, -1):
            slide = prs.slides[s_idx]
            if PPTXProcessor._is_ad_slide(slide, slide_index=s_idx + 1, total_slides=slide_count):
                PPTXProcessor._delete_slide(prs, s_idx)
                removed += 1

        # 2. Clean Slide Masters and Layouts
        for master in prs.slide_masters:
            for sh in list(master.shapes):
                if PPTXProcessor._is_watermark_recursive(sh):
                    try:
                        sh._element.getparent().remove(sh._element)
                        removed += 1
                    except Exception:
                        pass
                elif sh.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER:
                    if sh.left is not None and sh.top is not None:
                        if (sh.left > sw * 0.80 and sh.top < sh_h * 0.15) or sh.top > sh_h * 0.85:
                            try:
                                sh._element.getparent().remove(sh._element)
                                removed += 1
                            except Exception:
                                pass

            for layout in master.slide_layouts:
                for sh in list(layout.shapes):
                    if PPTXProcessor._is_watermark_recursive(sh):
                        try:
                            sh._element.getparent().remove(sh._element)
                            removed += 1
                        except Exception:
                            pass
                    elif sh.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER:
                        if sh.left is not None and sh.top is not None:
                            if (sh.left > sw * 0.80 and sh.top < sh_h * 0.15) or sh.top > sh_h * 0.85:
                                try:
                                    sh._element.getparent().remove(sh._element)
                                    removed += 1
                                    pass
                                except Exception:
                                    pass

        # 3. Clean Slide Level Watermarks, Corner Logo Badges and Footer Links
        for s_idx, slide in enumerate(prs.slides, 1):
            for sh in list(slide.shapes):
                # 1-slaydda asosiy sarlavha (Title / Title Placeholder) va sarlavha osti qutilari o'chirilmasligi kerak!
                # ALLPPT kabi shablonlarda 1-slayddagi sarlavha 'Free PPT Templates' yoki 'Insert the title' deb yozilgan bo'ladi.
                if s_idx == 1 and sh.has_text_frame and (sh.width or 0) >= sw * 0.35:
                    txt = sh.text_frame.text.strip().lower()
                    if not any(u in txt for u in ["http://", "https://", "www.allppt"]):
                        continue

                # A. Har tomonlama reklama va watermark tekshiruvi (matn, XML, descr, logo rasmlari, guruhlar)
                if PPTXProcessor._is_watermark_recursive(sh):
                    try:
                        sh._element.getparent().remove(sh._element)
                        removed += 1
                        continue
                    except Exception:
                        pass

                # B. Burchaklardagi logo nishonlari (masalan, yuqori o'ngdagi ALLPPT.com pilli) va footer havolalari
                try:
                    l, t, w, h = sh.left, sh.top, sh.width, sh.height
                    if l is not None and t is not None and w is not None and h is not None and sw and sh_h:
                        # Yuqori o'ng burchakdagi reklama nishoni (L > 80%, T < 15%, W < 25%, H < 12%)
                        if l > sw * 0.80 and t < sh_h * 0.15 and w < sw * 0.25 and h < sh_h * 0.12:
                            txt = sh.text_frame.text.strip() if sh.has_text_frame else ""
                            if PPTXProcessor._is_watermark_text(txt) or len(txt) < 15:
                                sh._element.getparent().remove(sh._element)
                                removed += 1
                                continue
                        # Pastki footer watermark / muallif / o'qituvchi / maktab (T > 82%, H < 20%)
                        if t > sh_h * 0.82 and h < sh_h * 0.20:
                            txt = sh.text_frame.text.strip().lower() if sh.has_text_frame else ""
                            if PPTXProcessor._is_watermark_text(txt) or any(p in txt for p in [
                                "http", "www", "free", "allppt", "template", "design", ".com",
                                "antonenkova", "o'qituvchi", "o‘qituvchi", "maktab", "maktabi",
                                "учитель", "преподаватель", "школа", "сош", "лицей", "гимназия",
                                "muallif", "автор", "tayyorladi", "подготовил", "выполнил"
                            ]) or re.search(r"[a-zа-я]\.[a-zа-я]\.\s*[a-zа-я]{3,}", txt) or re.search(r"\d{2}\.\d{2}\.\d{4}", txt):
                                sh._element.getparent().remove(sh._element)
                                removed += 1
                                continue
                except Exception:
                    pass

        return removed

    @staticmethod
    def _extract_font_meta(paragraph) -> Dict[str, Any]:
        font_name = "Calibri"
        font_size_pt = 16.0
        is_bold = False
        is_italic = False
        font_color = "#333333"

        if paragraph.runs:
            first_r = paragraph.runs[0]
            if first_r.font:
                if first_r.font.name:
                    font_name = first_r.font.name
                if first_r.font.size and first_r.font.size.pt:
                    font_size_pt = first_r.font.size.pt
                is_bold = bool(first_r.font.bold)
                is_italic = bool(first_r.font.italic)
                try:
                    if first_r.font.color and first_r.font.color.rgb:
                        font_color = f"#{first_r.font.color.rgb}"
                except Exception:
                    pass
        elif paragraph.font:
            if paragraph.font.name:
                font_name = paragraph.font.name
            if paragraph.font.size and paragraph.font.size.pt:
                font_size_pt = paragraph.font.size.pt
            is_bold = bool(paragraph.font.bold)
            is_italic = bool(paragraph.font.italic)

        try:
            font_manager.ensure_font_available(font_name)
        except Exception:
            pass

        return {
            "font_name": font_name,
            "font_size_pt": round(font_size_pt, 1),
            "is_bold": is_bold,
            "is_italic": is_italic,
            "font_color": font_color
        }

    @staticmethod
    def extract_presentation_data(pptx_path: str) -> Dict[str, Any]:
        if not os.path.exists(pptx_path):
            raise FileNotFoundError(f"PPTX topilmadi: {pptx_path}")
        prs = safe_load_presentation(pptx_path)
        slides_data = []
        total_items_count = 0
        slide_width = prs.slide_width
        slide_height = prs.slide_height

        for s_idx, slide in enumerate(prs.slides, start=1):
            if PPTXProcessor._is_ad_slide(slide, slide_index=s_idx):
                continue
            slide_items = []
            PPTXProcessor._extract_shapes_recursive(
                shapes=slide.shapes,
                slide_index=s_idx,
                slide_width=slide_width,
                slide_height=slide_height,
                items_list=slide_items
            )
            total_items_count += len(slide_items)
            slides_data.append({
                "slide_index": len(slides_data) + 1,
                "slide_id": f"slide_{len(slides_data) + 1}",
                "orig_slide_index": s_idx,
                "items_count": len(slide_items),
                "items": slide_items
            })

        return {
            "slide_width_pt": slide_width.pt if slide_width else 960,
            "slide_height_pt": slide_height.pt if slide_height else 540,
            "slides_count": len(slides_data),
            "total_items": total_items_count,
            "slides": slides_data
        }

    @staticmethod
    def _extract_shapes_recursive(shapes, slide_index: int, slide_width, slide_height, items_list: list, prefix: str = "", depth: int = 0, max_depth: int = 10):
        if depth >= max_depth:
            return
        for sh_idx, shape in enumerate(shapes):
            sh_id_str = f"{prefix}sh{getattr(shape, 'shape_id', sh_idx)}"
            box_info = {"left": 5, "top": 5, "width": 90, "height": 20}
            try:
                if shape.left is not None and shape.top is not None and slide_width and slide_height:
                    box_info = {
                        "left": round((shape.left / slide_width) * 100, 2),
                        "top": round((shape.top / slide_height) * 100, 2),
                        "width": round((shape.width / slide_width) * 100, 2),
                        "height": round((shape.height / slide_height) * 100, 2)
                    }
            except Exception:
                pass

            # 1. Group shapes
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                try:
                    PPTXProcessor._extract_shapes_recursive(shape.shapes, slide_index, slide_width, slide_height, items_list, prefix=f"{sh_id_str}_g", depth=depth + 1, max_depth=max_depth)
                except Exception:
                    pass
                continue

            # 2. Table
            if shape.has_table:
                table = shape.table
                for r_idx, row in enumerate(table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        if cell.text_frame:
                            for p_idx, p in enumerate(cell.text_frame.paragraphs):
                                p_text = sanitize_control_chars(p.text)
                                if PPTXProcessor._is_real_text(p_text) and not PPTXProcessor._is_watermark_text(p_text):
                                    f_meta = PPTXProcessor._extract_font_meta(p)
                                    item_id = f"s{slide_index}_{sh_id_str}_tbl_r{r_idx}_c{c_idx}_p{p_idx}"
                                    items_list.append({
                                        "id": item_id,
                                        "slide_index": slide_index,
                                        "shape_name": f"Jadval [Qator {r_idx+1}, Ustun {c_idx+1}]",
                                        "item_type": "table_cell",
                                        "original_text": p_text,
                                        "translated_text": p_text,
                                        "font_name": f_meta["font_name"],
                                        "font_size_pt": f_meta["font_size_pt"],
                                        "is_bold": f_meta["is_bold"],
                                        "is_italic": f_meta["is_italic"],
                                        "font_color": f_meta["font_color"],
                                        "box": box_info
                                    })
                continue

            # 3. Regular Shape / TextBox
            if shape.has_text_frame:
                tf = shape.text_frame
                has_tf_text = False
                for p_idx, p in enumerate(tf.paragraphs):
                    p_text = sanitize_control_chars(p.text)
                    if PPTXProcessor._is_real_text(p_text) and not PPTXProcessor._is_watermark_text(p_text):
                        f_meta = PPTXProcessor._extract_font_meta(p)
                        item_id = f"s{slide_index}_{sh_id_str}_p{p_idx}"
                        items_list.append({
                            "id": item_id,
                            "slide_index": slide_index,
                            "shape_name": shape.name or f"Shakl {sh_idx+1}",
                            "item_type": "title" if p_idx == 0 and len(p_text) < 40 else "body",
                            "original_text": p_text,
                            "translated_text": p_text,
                            "font_name": f_meta["font_name"],
                            "font_size_pt": f_meta["font_size_pt"],
                            "is_bold": f_meta["is_bold"],
                            "is_italic": f_meta["is_italic"],
                            "font_color": f_meta["font_color"],
                            "box": box_info
                        })
                        has_tf_text = True
                if has_tf_text:
                    continue

            # 4. Native PowerPoint Charts
            if shape.has_chart:
                try:
                    c_elem = shape.chart._element
                    for t_idx, t_node in enumerate(c_elem.xpath('.//c:title//a:t')):
                        t_text = sanitize_control_chars(t_node.text) if t_node.text else ""
                        if PPTXProcessor._is_real_text(t_text) and not PPTXProcessor._is_watermark_text(t_text):
                            item_id = f"s{slide_index}_{sh_id_str}_chtitle_{t_idx}"
                            items_list.append({
                                "id": item_id,
                                "slide_index": slide_index,
                                "shape_name": f"Diagramma Sarlavhasi ({shape.name or f'Shakl {sh_idx+1}'})",
                                "item_type": "title",
                                "original_text": t_text,
                                "translated_text": t_text,
                                "font_name": "Calibri",
                                "font_size_pt": 18.0,
                                "is_bold": True,
                                "is_italic": False,
                                "font_color": "#333333",
                                "box": box_info
                            })
                    for cat_idx, v_node in enumerate(c_elem.xpath('.//c:cat//c:pt//c:v')):
                        c_text = sanitize_control_chars(v_node.text) if v_node.text else ""
                        if PPTXProcessor._is_real_text(c_text) and not PPTXProcessor._is_watermark_text(c_text):
                            item_id = f"s{slide_index}_{sh_id_str}_chcat_{cat_idx}"
                            items_list.append({
                                "id": item_id,
                                "slide_index": slide_index,
                                "shape_name": f"Diagramma Toifasi [{c_text}]",
                                "item_type": "body",
                                "original_text": c_text,
                                "translated_text": c_text,
                                "font_name": "Calibri",
                                "font_size_pt": 14.0,
                                "is_bold": False,
                                "is_italic": False,
                                "font_color": "#333333",
                                "box": box_info
                            })
                    for ser_idx, v_node in enumerate(c_elem.xpath('.//c:ser//c:tx//c:v')):
                        s_text = sanitize_control_chars(v_node.text) if v_node.text else ""
                        if PPTXProcessor._is_real_text(s_text) and not PPTXProcessor._is_watermark_text(s_text):
                            item_id = f"s{slide_index}_{sh_id_str}_chser_{ser_idx}"
                            items_list.append({
                                "id": item_id,
                                "slide_index": slide_index,
                                "shape_name": f"Diagramma Seriyasi [{s_text}]",
                                "item_type": "body",
                                "original_text": s_text,
                                "translated_text": s_text,
                                "font_name": "Calibri",
                                "font_size_pt": 14.0,
                                "is_bold": False,
                                "is_italic": False,
                                "font_color": "#333333",
                                "box": box_info
                            })
                except Exception:
                    pass
                continue

            # 5. SmartArt, DrawingML & Custom Shapes XML Fallback
            try:
                elem = getattr(shape, "_element", None)
                if elem is not None:
                    for p_idx, p_elem in enumerate(elem.xpath('.//a:p')):
                        t_nodes = p_elem.xpath('.//a:t')
                        if t_nodes:
                            p_text = sanitize_control_chars("".join(t.text or "" for t in t_nodes))
                            if PPTXProcessor._is_real_text(p_text) and not PPTXProcessor._is_watermark_text(p_text):
                                item_id = f"s{slide_index}_{sh_id_str}_xmlp_{p_idx}"
                                if not any(it["id"] == item_id for it in items_list):
                                    items_list.append({
                                        "id": item_id,
                                        "slide_index": slide_index,
                                        "shape_name": shape.name or f"Diagramma Blok {sh_idx+1}",
                                        "item_type": "body",
                                        "original_text": p_text,
                                        "translated_text": p_text,
                                        "font_name": "Calibri",
                                        "font_size_pt": 16.0,
                                        "is_bold": False,
                                        "is_italic": False,
                                        "font_color": "#333333",
                                        "box": box_info
                                    })
            except Exception:
                pass

    @staticmethod
    def apply_translations_and_export(
        original_pptx_path: str,
        translations_map: Dict[str, str],
        output_pptx_path: str,
        auto_fit: bool = True,
        target_script: str = "latin",
        clean_watermarks: bool = True,
        presentation_title: str = "",
        embed_thumbnail: bool = False
    ) -> str:
        if not os.path.exists(original_pptx_path):
            raise FileNotFoundError(f"Original PPTX topilmadi: {original_pptx_path}")

        prs = safe_load_presentation(original_pptx_path)

        # 1. Matnli shakllar va SmartArt diagrammalarini tarjima qilish (Asl slayd indekslari bo'yicha)
        for s_idx, slide in enumerate(prs.slides, start=1):
            PPTXProcessor._apply_to_shapes_recursive(
                shapes=slide.shapes,
                slide_index=s_idx,
                translations_map=translations_map,
                auto_fit=auto_fit,
                target_script=target_script
            )
            # Ustma-ust tushishlar va gorizontal/vertikal noaniqliklarni avtomatik bartaraf etish
            PPTXProcessor._optimize_slide_layout(slide, prs.slide_width, prs.slide_height)

        # 2. Rasm shaklidagi slaydlar va diagramma rasmlarini AI (Gemini Vision / Nano) orqali tarjima qilish
        try:
            img_translator = ImageTranslator()
            image_cache: Dict[bytes, bytes] = {}
            for slide in prs.slides:
                # Slayd ichidagi barcha rasmlarni a:blip orqali to'liq qamrab olish
                blip_rIds = slide._element.xpath('.//a:blip/@r:embed')
                for rId in set(blip_rIds):
                    try:
                        img_part = slide.part.related_part(rId)
                        if hasattr(img_part, 'blob') and img_part.blob:
                            orig_blob = img_part.blob
                            if len(orig_blob) > 2000:
                                if orig_blob in image_cache:
                                    new_blob = image_cache[orig_blob]
                                else:
                                    new_blob = img_translator.analyze_and_translate_image(orig_blob, target_script=target_script)
                                    image_cache[orig_blob] = new_blob
                                
                                if new_blob and new_blob != orig_blob:
                                    img_part._blob = new_blob
                    except Exception:
                        pass
        except Exception as img_err:
            print(f"[PPTXProcessor] Rasm tarjimasida ogohlantirish: {img_err}")

        # 3. Reklama, suvbelgilar va resurs slaydlarini tarjima yakunlangach tozalash
        if clean_watermarks:
            PPTXProcessor.clean_presentation_watermarks(prs)

        # 1-slayd sarlavhasini kafolatlash (agar shablon sarlavhasiz yoki placeholder bo'lsa)
        if len(prs.slides) > 0 and presentation_title:
            s1 = prs.slides[0]
            sw = prs.slide_width
            
            # Keng matn qutilarini topish (kengligi kamida slaydning 35% qismini egallagan)
            candidates = []
            for sh in s1.shapes:
                if sh.has_text_frame and (sh.width or 0) >= sw * 0.35:
                    candidates.append((sh.top or 0, sh))
            
            candidates.sort(key=lambda x: x[0])
            
            if candidates:
                # 1-kandidat: Asosiy sarlavha (Title)
                title_sh = candidates[0][1]
                if title_sh.text_frame.paragraphs:
                    p_title = title_sh.text_frame.paragraphs[0]
                    if p_title.runs:
                        p_title.runs[0].text = presentation_title
                        for r in p_title.runs[1:]:
                            r.text = ""
                        p_title.runs[0].font.size = Pt(38) if len(presentation_title) > 30 else Pt(44)
                        p_title.runs[0].font.bold = True
                    else:
                        p_title.text = presentation_title
                        p_title.font.size = Pt(38) if len(presentation_title) > 30 else Pt(44)
                        p_title.font.bold = True
                    p_title.alignment = PP_ALIGN.CENTER
                
                # 2-kandidat: Sarlavha osti (Subtitle)
                if len(candidates) >= 2:
                    sub_sh = candidates[1][1]
                    if sub_sh.text_frame.paragraphs:
                        p_sub = sub_sh.text_frame.paragraphs[0]
                        if p_sub.runs:
                            p_sub.runs[0].text = "Taqdimot materiali"
                            for r in p_sub.runs[1:]:
                                r.text = ""
                            p_sub.runs[0].font.size = Pt(18)
                        else:
                            p_sub.text = "Taqdimot materiali"
                            p_sub.font.size = Pt(18)
                        p_sub.alignment = PP_ALIGN.CENTER
            else:
                for sh in s1.shapes:
                    if sh.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and sh.has_text_frame:
                        sh.text_frame.text = presentation_title
                        for p in sh.text_frame.paragraphs:
                            p.alignment = PP_ALIGN.CENTER
                            p.font.size = Pt(36)
                            p.font.bold = True
                        break

        os.makedirs(os.path.dirname(os.path.abspath(output_pptx_path)), exist_ok=True)
        prs.save(output_pptx_path)

        # 1-slayd preview rasmini generatsiya qilish va .pptx arxivi ichiga embed qilish (ixtiyoriy)
        if embed_thumbnail:
            try:
                ThumbnailGenerator.embed_thumbnail_into_pptx(output_pptx_path)
            except Exception as e:
                print(f"[PPTXProcessor] Thumbnail yaratishda ogohlantirish: {e}")

        return output_pptx_path

    @staticmethod
    def _apply_to_shapes_recursive(shapes, slide_index: int, translations_map: Dict[str, str], auto_fit: bool, target_script: str, prefix: str = "", depth: int = 0, max_depth: int = 10):
        if depth >= max_depth:
            return
        for sh_idx, shape in enumerate(shapes):
            sh_id_str = f"{prefix}sh{getattr(shape, 'shape_id', sh_idx)}"

            # 1. Group
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                try:
                    PPTXProcessor._apply_to_shapes_recursive(shape.shapes, slide_index, translations_map, auto_fit, target_script, prefix=f"{sh_id_str}_g", depth=depth + 1, max_depth=max_depth)
                except Exception:
                    pass
                continue

            # 2. Table
            if shape.has_table:
                table = shape.table
                for r_idx, row in enumerate(table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        if cell.text_frame:
                            cell.text_frame.word_wrap = True
                            cell.text_frame.margin_left = Inches(0.01)
                            cell.text_frame.margin_right = Inches(0.01)
                            cell.text_frame.margin_top = Inches(0.01)
                            cell.text_frame.margin_bottom = Inches(0.01)
                            for p_idx, p in enumerate(cell.text_frame.paragraphs):
                                p_orig = sanitize_control_chars(p.text)
                                if PPTXProcessor._is_real_text(p_orig) and not PPTXProcessor._is_watermark_text(p_orig):
                                    item_id = f"s{slide_index}_{sh_id_str}_tbl_r{r_idx}_c{c_idx}_p{p_idx}"
                                    trans = translations_map.get(item_id)
                                    if trans:
                                        PPTXProcessor._set_paragraph_text_safe(p, ensure_script(sanitize_control_chars(trans), target_script), auto_fit=auto_fit, shape=shape)
                            PPTXProcessor._harmonize_text_frame_typography(cell.text_frame)
                continue

            # 3. Regular Shape / TextBox
            if shape.has_text_frame:
                tf = shape.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.TOP  # Ustma-ust tushishning oldini oladi (matn yuqoriga emas, pastga kengayadi)
                tf.margin_left = Inches(0.01)
                tf.margin_right = Inches(0.01)
                tf.margin_top = Inches(0.01)
                tf.margin_bottom = Inches(0.01)
                for p_idx, p in enumerate(tf.paragraphs):
                    p_orig = sanitize_control_chars(p.text)
                    if PPTXProcessor._is_real_text(p_orig) and not PPTXProcessor._is_watermark_text(p_orig):
                        item_id = f"s{slide_index}_{sh_id_str}_p{p_idx}"
                        trans = translations_map.get(item_id)
                        if trans:
                            PPTXProcessor._set_paragraph_text_safe(p, ensure_script(sanitize_control_chars(trans), target_script), auto_fit=auto_fit, shape=shape)
                PPTXProcessor._harmonize_text_frame_typography(tf)

            # 4. Native Charts
            if shape.has_chart:
                try:
                    c_elem = shape.chart._element
                    for t_idx, t_node in enumerate(c_elem.xpath('.//c:title//a:t')):
                        item_id = f"s{slide_index}_{sh_id_str}_chtitle_{t_idx}"
                        trans = translations_map.get(item_id)
                        if trans:
                            t_node.text = ensure_script(sanitize_control_chars(trans), target_script)
                    for cat_idx, v_node in enumerate(c_elem.xpath('.//c:cat//c:pt//c:v')):
                        item_id = f"s{slide_index}_{sh_id_str}_chcat_{cat_idx}"
                        trans = translations_map.get(item_id)
                        if trans:
                            v_node.text = ensure_script(sanitize_control_chars(trans), target_script)
                    for ser_idx, v_node in enumerate(c_elem.xpath('.//c:ser//c:tx//c:v')):
                        item_id = f"s{slide_index}_{sh_id_str}_chser_{ser_idx}"
                        trans = translations_map.get(item_id)
                        if trans:
                            v_node.text = ensure_script(sanitize_control_chars(trans), target_script)
                except Exception:
                    pass

            # 5. SmartArt, DrawingML & Custom Shapes XML Fallback
            try:
                elem = getattr(shape, "_element", None)
                if elem is not None:
                    for p_idx, p_elem in enumerate(elem.xpath('.//a:p')):
                        item_id = f"s{slide_index}_{sh_id_str}_xmlp_{p_idx}"
                        trans = translations_map.get(item_id)
                        if trans:
                            t_nodes = p_elem.xpath('.//a:t')
                            if t_nodes:
                                safe_trans = ensure_script(sanitize_control_chars(trans), target_script)
                                t_nodes[0].text = safe_trans
                                for t_extra in t_nodes[1:]:
                                    t_extra.text = ""
            except Exception:
                pass

    @staticmethod
    def _harmonize_text_frame_typography(tf):
        """
        Matn qutisi (Text Frame) ichidagi barcha bandlar (ro'yxat va paragraflar)
        uchun yagona va bir xil shrift o'lchamini va intervalini ta'minlaydi.
        """
        non_empty = [p for p in tf.paragraphs if p.text and p.text.strip()]
        if len(non_empty) <= 1:
            return

        sizes = []
        fonts = []
        for p in non_empty:
            for r in p.runs:
                if r.font and r.font.size and r.text and len(r.text.strip()) > 2:
                    try:
                        sizes.append(r.font.size.pt)
                    except Exception:
                        pass
                if r.font and r.font.name:
                    fonts.append(r.font.name)

        if not sizes:
            return

        # Eng mos keluvchi (eng kichik yoki barqaror) shrift o'lchamini tanlash
        target_size_pt = min(sizes)
        target_font_name = max(set(fonts), key=fonts.count) if fonts else "Calibri"

        for p in non_empty:
            try:
                p.font.size = Pt(target_size_pt)
                p.font.name = target_font_name
                p.space_after = Pt(6)
                p.line_spacing = 1.15
                for r in p.runs:
                    if r.font:
                        r.font.size = Pt(target_size_pt)
                        r.font.name = target_font_name
            except Exception:
                pass

    @staticmethod
    def _set_paragraph_text_safe(paragraph, new_text: str, auto_fit: bool = True, shape = None):
        # Remove any lingering <a:br> elements that cause control character _x000B_ or trailing line jumps
        try:
            for br in list(paragraph._p.xpath('.//a:br')):
                br.getparent().remove(br)
        except Exception:
            pass

        if not paragraph.runs:
            paragraph.text = sanitize_control_chars(new_text)
            return

        clean_val = sanitize_control_chars(new_text)
        safe_text = re.sub(r"([A-Za-zА-Яа-яЎўҒғҚқҲҳ])['`’‘ʼʻ]([A-Za-zА-Яа-яЎўҒғҚқҲҳ])", r"\1'\2", clean_val)

        # Lorem Ipsum va shablonning o'zini o'zi maqtash reklama matnlarini avtomatik o'zbekchalashtirish
        low_t = safe_text.lower()
        if "lorem ipsum" in low_t or any(pat in low_t for pat in [
            "ushbu shablon vaqtingiz",
            "obro'yingizni tejashiga",
            "obro‘yingizni tejashiga",
            "obroyingizni tejashiga",
            "tinglovchilaringizni hayratda qoldiring",
            "tomoshabinlarni hayratda qoldiring",
            "shablonlarimiz yordamida hisobot",
            "oddiy portfolio",
            "impress your audience",
            "this template will",
            "save your time, money",
            "easy to change colors, photos",
            "get a modern powerpoint presentation",
            "presentation designed",
            "simple portfolio"
        ]):
            if len(safe_text) <= 35:
                safe_text = "Asosiy tahliliy ko'rsatkichlar"
            else:
                safe_text = "Ushbu bo'limda taqdimot mavzusi yuzasidan batafsil ma'lumotlar, asosiy ko'rsatkichlar va tahliliy xulosalar keltiriladi."

        # Matn tanasining asl parametrlarini olish uchun body_run'ni aniqlash
        body_run = None
        for r in paragraph.runs:
            if r.text and len(r.text.strip()) > 3:
                body_run = r
                break
        if not body_run and paragraph.runs:
            body_run = paragraph.runs[-1]
        if not body_run and paragraph.runs:
            body_run = paragraph.runs[0]

        first_run = paragraph.runs[0]
        orig_font_name = body_run.font.name if (body_run.font and body_run.font.name) else (first_run.font.name if first_run.font else None)
        orig_font_size = body_run.font.size if (body_run.font and body_run.font.size) else (first_run.font.size if first_run.font else None)
        orig_bold = body_run.font.bold if (body_run.font and body_run.font.bold is not None) else None
        orig_italic = body_run.font.italic if (body_run.font and body_run.font.italic is not None) else None
        orig_color = None
        try:
            if body_run.font and body_run.font.color and body_run.font.color.type is not None:
                if body_run.font.color.rgb:
                    orig_color = body_run.font.color.rgb
        except Exception:
            pass

        first_run.text = safe_text

        # 100% Consistent Typography: Shrift o'lchami va uslubini barcha bandlar uchun bir xil saqlash
        if orig_font_name and first_run.font:
            first_run.font.name = orig_font_name
        if orig_font_size and first_run.font:
            first_run.font.size = orig_font_size
        if orig_bold is not None and first_run.font:
            first_run.font.bold = orig_bold
        if orig_italic is not None and first_run.font:
            first_run.font.italic = orig_italic
        if orig_color and first_run.font and first_run.font.color:
            try:
                first_run.font.color.rgb = orig_color
            except Exception:
                pass

        # Ortiqcha qo'shimcha runlarni tozalash
        if len(paragraph.runs) > 1:
            for r in paragraph.runs[1:]:
                r.text = ""

    @staticmethod
    def _optimize_slide_layout(slide, slide_width, slide_height):
        """
        Slayd ichidagi matn va shakllarning ustma-ust tushishi (overlap)
        hamda gorizontal/vertikal noaniqliklarni intellektual to'g'irlaydi.
        """
        sw = slide_width
        sh_h = slide_height

        # 1. Vertikal dekorativ shakllar / ustunlarni aniqlash (masalan, chapdagi vertikal ajratuvchi chiziq)
        v_bars = []
        for sh in slide.shapes:
            try:
                sh_w = getattr(sh, 'width', None)
                sh_h_val = getattr(sh, 'height', None)
                if sh_w is not None and sh_h_val is not None:
                    if sh_h_val >= sh_h * 0.65 and sh_w <= sw * 0.25 and getattr(sh, 'shape_type', None) in [MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM]:
                        v_bars.append(sh)
            except Exception:
                pass

        all_tbs = [sh for sh in slide.shapes if getattr(sh, 'has_text_frame', False)]

        for tb in all_tbs:
            try:
                tf = tb.text_frame
            except Exception:
                continue
            txt = tf.text.strip()
            if not txt:
                continue

            # Shablonning qo'shimcha sun'iy so'zlarini tozalash (masalan: "Kun tartibi uslubi" -> "Kun tartibi")
            cleaned_txt = re.sub(r"\b(uslubi|stili)\b", "", txt, flags=re.IGNORECASE).strip()
            if cleaned_txt and cleaned_txt != txt and len(cleaned_txt) >= 3:
                try:
                    p0 = tf.paragraphs[0]
                    p0.text = cleaned_txt
                    txt = cleaned_txt
                except Exception:
                    pass

            first_p = tf.paragraphs[0]
            size_pt = 16.0
            if first_p.runs and first_p.runs[0].font and first_p.runs[0].font.size:
                try:
                    size_pt = first_p.runs[0].font.size.pt
                except Exception:
                    pass

            tb_left = getattr(tb, 'left', None)
            tb_top = getattr(tb, 'top', None)
            tb_w = getattr(tb, 'width', None)
            tb_h = getattr(tb, 'height', None)
            if tb_left is None or tb_top is None or tb_w is None or tb_h is None:
                continue

            # A. Gorizontal ustun to'qnashuvini tuzatish (Sarlavha vertikal chiziq/grafika ustiga chiqib qolmasligi)
            for bar in v_bars:
                b_left = getattr(bar, 'left', None)
                b_w = getattr(bar, 'width', None)
                if b_left is None or b_w is None:
                    continue
                bar_right = b_left + b_w
                # Agar sarlavha chiziqdan oldin yoki ichida boshlanib, o'ng tomonga cho'zilgan bo'lsa
                if tb_top < sh_h * 0.35 and tb_left < bar_right and (tb_left + tb_w) > bar_right + Inches(1.5):
                    # O'ng tarafdagi kontent bloklarini topamiz
                    right_shapes = [s for s in slide.shapes if getattr(s, 'left', None) is not None and s.left >= bar_right and id(s) != id(tb)]
                    if right_shapes:
                        min_right = min(s.left for s in right_shapes)
                        tb.left = min_right
                        tb.width = max(Inches(3), sw - tb.left - Inches(0.4))
                        tb_left = tb.left
                        tb_w = tb.width

            # B. Vertikal to'qnashuv (yuqoridagi rasm/ikonka bilan ustma-ust tushish)
            shapes_above = []
            for other in slide.shapes:
                if id(other) != id(tb):
                    o_left = getattr(other, 'left', None)
                    o_top = getattr(other, 'top', None)
                    o_w = getattr(other, 'width', None)
                    o_h = getattr(other, 'height', None)
                    if o_left is not None and o_top is not None and o_w is not None and o_h is not None:
                        # Gorizontal kesishish
                        if not (o_left + o_w <= tb_left or o_left >= tb_left + tb_w):
                            other_bottom = o_top + o_h
                            if other_bottom <= tb_top + Inches(0.2) and (tb_top - other_bottom) < Inches(0.8):
                                shapes_above.append(other)

            if shapes_above:
                tf.vertical_anchor = MSO_ANCHOR.TOP
                if "\n" not in txt and len(txt) <= 25 and size_pt >= 24.0:
                    box_w_pt = tb.width.pt if tb.width else 400
                    max_pt = (box_w_pt - 25) / (len(txt) * 0.65)
                    target_pt = min(size_pt, max(16.0, max_pt))
                    if target_pt < size_pt:
                        first_p.runs[0].font.size = Pt(round(target_pt, 1))

        # 1. Sarlavha / Banner shakllarining balandligi va matn sig'ishini optimallashtirish
        header_shapes = []
        for tb in all_tbs:
            tb_top = getattr(tb, 'top', None)
            tb_h = getattr(tb, 'height', None)
            tb_w = getattr(tb, 'width', None)
            if tb_top is not None and tb_h is not None and tb_w is not None:
                # Yuqori sarlavha paneli (Top 22%, Width > 35%)
                if tb_top < sh_h * 0.22 and tb_w >= sw * 0.35:
                    header_shapes.append(tb)
                    tf = tb.text_frame
                    txt = tf.text.strip()
                    if txt:
                        char_count = len(txt)
                        # Agar sarlavha 30 belgidan oshsa yoki quti balandligi kichik bo'lsa
                        if char_count > 30 or "\n" in txt or tb_h < Inches(1.2):
                            target_header_size = 19.0 if char_count > 45 else 22.0
                            for p in tf.paragraphs:
                                p.line_spacing = 1.0
                                p.space_after = Pt(0)
                                p.space_before = Pt(0)
                                for r in p.runs:
                                    if r.font:
                                        r.font.size = Pt(target_header_size)
                        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                        tf.margin_top = Inches(0.02)
                        tf.margin_bottom = Inches(0.02)

        # 2. Sarlavha ostidagi kontent bloklari sarlavha ustiga minib qolmasligini ta'minlash
        if header_shapes:
            max_header_bottom = max(h.top + h.height for h in header_shapes)
            for other in slide.shapes:
                if id(other) not in [id(h) for h in header_shapes]:
                    o_top = getattr(other, 'top', None)
                    if o_top is not None and o_top < max_header_bottom + Inches(0.05) and o_top >= max_header_bottom - Inches(0.40):
                        try:
                            other.top = int(max_header_bottom + Inches(0.08))
                        except Exception:
                            pass

        # 3. Bitta slaydda bir nechta kartochka / bloklar bo'lsa (masalan, 13-slayddagi 8 ta ko'k quti):
        card_shapes = []
        for sh in slide.shapes:
            if getattr(sh, 'has_text_frame', False) and id(sh) not in [id(h) for h in header_shapes]:
                sh_w = getattr(sh, 'width', None)
                sh_h_val = getattr(sh, 'height', None)
                sh_t = getattr(sh, 'top', None)
                if sh_w and sh_h_val and sh_t and sh_t > sh_h * 0.15:
                    if sw * 0.15 <= sh_w <= sw * 0.70 and sh_h_val <= sh_h * 0.35:
                        if sh.text_frame.text.strip():
                            card_shapes.append(sh)

        if len(card_shapes) >= 3:
            card_sizes = []
            for c_sh in card_shapes:
                txt = c_sh.text_frame.text.strip()
                c_w_pt = c_sh.width.pt if c_sh.width else 200
                fit_sz = min(15.0, max(10.0, (c_w_pt / max(1, len(txt))) * 1.7))
                if len(txt) > 25:
                    fit_sz = min(fit_sz, 12.0)
                card_sizes.append(fit_sz)

            unified_card_font = min(card_sizes) if card_sizes else 13.0
            unified_card_font = max(11.0, min(14.0, unified_card_font))

            for c_sh in card_shapes:
                tf_c = c_sh.text_frame
                tf_c.word_wrap = True
                tf_c.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf_c.margin_left = Inches(0.02)
                tf_c.margin_right = Inches(0.02)
                tf_c.margin_top = Inches(0.02)
                tf_c.margin_bottom = Inches(0.02)
                for p in tf_c.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    p.line_spacing = 1.05
                    p.space_after = Pt(0)
                    p.space_before = Pt(0)
                    for r in p.runs:
                        if r.font:
                            r.font.size = Pt(unified_card_font)

        # 4. Qayta qolgan reklama va mualliflik footer matnli shakllarini butunlay o'chirish
        for sh in list(slide.shapes):
            if sh.has_text_frame and PPTXProcessor._is_watermark_text(sh.text_frame.text):
                try:
                    sh._element.getparent().remove(sh._element)
                except Exception:
                    pass
