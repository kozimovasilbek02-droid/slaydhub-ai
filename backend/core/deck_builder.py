# -*- coding: utf-8 -*-
"""
Native Vector PPTX Deck Builder
-------------------------------
JSON spetsifikatsiyasi asosida 100% tahrirlanuvchi (native editable vector)
yuqori sifatli PowerPoint taqdimotlarini noldan chizuvchi dvigatel.
"""

import os
import re
from typing import Dict, List, Any, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


PALETTES = {
    "uzbek_blue": {
        "primary": RGBColor(11, 61, 145),      # Deep Uzbek Blue
        "secondary": RGBColor(0, 150, 136),    # Emerald Teal
        "accent": RGBColor(212, 175, 55),      # Gold
        "dark_text": RGBColor(33, 37, 41),     # Dark Slate
        "muted_text": RGBColor(100, 116, 139), # Muted Gray
        "light_bg": RGBColor(248, 250, 252),   # Crisp Light Gray-Blue
        "card_bg": RGBColor(255, 255, 255),    # White
        "card_border": RGBColor(226, 232, 240)
    },
    "emerald_teal": {
        "primary": RGBColor(15, 118, 110),     # Teal Dark
        "secondary": RGBColor(13, 148, 136),   # Teal Accent
        "accent": RGBColor(245, 158, 11),      # Warm Amber
        "dark_text": RGBColor(15, 23, 42),
        "muted_text": RGBColor(100, 116, 139),
        "light_bg": RGBColor(240, 253, 250),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(204, 251, 241)
    },
    "dark_slate": {
        "primary": RGBColor(30, 41, 59),       # Slate Dark
        "secondary": RGBColor(99, 102, 241),   # Indigo
        "accent": RGBColor(6, 182, 212),       # Cyan
        "dark_text": RGBColor(15, 23, 42),
        "muted_text": RGBColor(100, 116, 139),
        "light_bg": RGBColor(241, 245, 249),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(226, 232, 240)
    },
    "modern_purple": {
        "primary": RGBColor(107, 33, 168),     # Purple
        "secondary": RGBColor(139, 92, 246),   # Violet
        "accent": RGBColor(245, 158, 11),      # Amber
        "dark_text": RGBColor(30, 27, 75),
        "muted_text": RGBColor(100, 116, 139),
        "light_bg": RGBColor(250, 245, 255),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(243, 232, 255)
    },
    "crimson_ruby": {
        "primary": RGBColor(153, 27, 27),      # Crimson
        "secondary": RGBColor(234, 88, 12),    # Orange Accent
        "accent": RGBColor(217, 119, 6),       # Amber Gold
        "dark_text": RGBColor(24, 24, 27),
        "muted_text": RGBColor(113, 113, 122),
        "light_bg": RGBColor(254, 242, 242),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(254, 226, 226)
    }
}


class DeckBuilder:
    """JSON spetsifikatsiyasidan 16:9 formatidagi vektorli PPTX fayl quruvchi sinf."""

    @classmethod
    def build(cls, spec: Dict[str, Any], output_path: str) -> str:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6]

        theme_name = spec.get("theme", "uzbek_blue")
        colors = PALETTES.get(theme_name, PALETTES["uzbek_blue"])
        slides_data = spec.get("slides", [])
        total_slides = len(slides_data)

        for s_idx, s_info in enumerate(slides_data, 1):
            slide = prs.slides.add_slide(blank_layout)
            stype = s_info.get("type", "cards")

            # 1. Background
            cls._set_background(slide, colors["light_bg"])

            # 2. Render Slide Type
            if stype == "hero":
                cls._render_hero(slide, s_info, colors)
            elif stype == "agenda":
                cls._render_header(slide, s_info.get("title", "Mundarija"), s_info.get("badge", "REJA"), colors)
                cls._render_agenda(slide, s_info, colors)
                cls._render_footer(slide, s_idx, total_slides, colors)
            elif stype == "stats":
                cls._render_header(slide, s_info.get("title", ""), s_info.get("badge", ""), colors)
                cls._render_stats(slide, s_info, colors)
                cls._render_footer(slide, s_idx, total_slides, colors)
            elif stype == "comparative":
                cls._render_header(slide, s_info.get("title", ""), s_info.get("badge", ""), colors)
                cls._render_comparative(slide, s_info, colors)
                cls._render_footer(slide, s_idx, total_slides, colors)
            elif stype == "conclusion":
                cls._render_header(slide, s_info.get("title", "Xulosa"), s_info.get("badge", "YAKUNIY XULOSA"), colors)
                cls._render_conclusion(slide, s_info, colors)
                cls._render_footer(slide, s_idx, total_slides, colors)
            else:  # cards
                cls._render_header(slide, s_info.get("title", ""), s_info.get("badge", ""), colors)
                cls._render_cards(slide, s_info, colors)
                cls._render_footer(slide, s_idx, total_slides, colors)

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        prs.save(output_path)
        return output_path

    @classmethod
    def _set_background(cls, slide, color: RGBColor):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        slide.shapes._spTree.remove(bg._element)
        slide.shapes._spTree.insert(2, bg._element)

    @classmethod
    def _render_header(cls, slide, title_text: str, badge_text: str, colors: Dict[str, RGBColor]):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_badge = tf.paragraphs[0]
        p_badge.text = (badge_text or "TAHLILIY BO'LIM").upper()
        p_badge.font.size = Pt(11)
        p_badge.font.bold = True
        p_badge.font.color.rgb = colors["secondary"]
        p_badge.font.name = 'Arial'

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = colors["primary"]
        p_title.font.name = 'Arial'
        p_title.space_before = Pt(3)

    @classmethod
    def _render_footer(cls, slide, current_slide: int, total_slides: int, colors: Dict[str, RGBColor]):
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = colors["card_border"]
        line.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        tf = tb.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"Slayd {current_slide} / {total_slides}"
        p.font.size = Pt(10)
        p.font.color.rgb = colors["muted_text"]
        p.font.name = 'Arial'
        p.alignment = PP_ALIGN.RIGHT

    @classmethod
    def _render_hero(cls, slide, data: Dict[str, Any], colors: Dict[str, RGBColor]):
        # Decorative colored accent top bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(1.5), Inches(0.06))
        bar.fill.solid()
        bar.fill.fore_color.rgb = colors["accent"]
        bar.line.fill.background()

        # Category Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(3.2), Inches(0.45))
        badge.fill.solid()
        badge.fill.fore_color.rgb = colors["primary"]
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_b = tf_b.paragraphs[0]
        p_b.text = (data.get("badge") or data.get("category") or "ILMIY TAQDIMOT").upper()
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = RGBColor(255, 255, 255)
        p_b.font.name = 'Arial'
        p_b.alignment = PP_ALIGN.CENTER

        # Main Title & Subtitle
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.733), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_t = tf.paragraphs[0]
        p_t.text = data.get("title", "Taqdimot Sarlavhasi")
        p_t.font.size = Pt(34)
        p_t.font.bold = True
        p_t.font.color.rgb = colors["primary"]
        p_t.font.name = 'Arial'

        p_s = tf.add_paragraph()
        p_s.text = data.get("subtitle", "")
        p_s.font.size = Pt(16)
        p_s.font.color.rgb = colors["muted_text"]
        p_s.font.name = 'Arial'
        p_s.space_before = Pt(20)

    @classmethod
    def _render_agenda(cls, slide, data: Dict[str, Any], colors: Dict[str, RGBColor]):
        items = data.get("items", [])
        top_y = Inches(1.6)
        item_h = Inches(0.9)
        gap = Inches(0.15)

        for i, item in enumerate(items[:5]):
            cur_y = top_y + i * (item_h + gap)
            # Card background
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), cur_y, Inches(11.733), item_h)
            card.fill.solid()
            card.fill.fore_color.rgb = colors["card_bg"]
            card.line.color.rgb = colors["card_border"]
            card.line.width = Pt(1)

            # Number badge
            num_b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), cur_y + Inches(0.15), Inches(0.6), Inches(0.6))
            num_b.fill.solid()
            num_b.fill.fore_color.rgb = colors["secondary"]
            num_b.line.fill.background()
            p_n = num_b.text_frame.paragraphs[0]
            p_n.text = item.get("num", f"{i+1:02d}")
            p_n.font.size = Pt(14)
            p_n.font.bold = True
            p_n.font.color.rgb = RGBColor(255, 255, 255)
            p_n.alignment = PP_ALIGN.CENTER

            # Text box
            tb = slide.shapes.add_textbox(Inches(1.8), cur_y + Inches(0.12), Inches(10.5), Inches(0.65))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

            p_t = tf.paragraphs[0]
            p_t.text = item.get("title", "")
            p_t.font.size = Pt(14)
            p_t.font.bold = True
            p_t.font.color.rgb = colors["dark_text"]

            p_d = tf.add_paragraph()
            p_d.text = item.get("desc", "")
            p_d.font.size = Pt(11)
            p_d.font.color.rgb = colors["muted_text"]
            p_d.space_before = Pt(2)

    @classmethod
    def _render_cards(cls, slide, data: Dict[str, Any], colors: Dict[str, RGBColor]):
        cards = data.get("cards", [])[:3]
        card_w = Inches(3.68)
        card_h = Inches(4.7)
        gap = Inches(0.34)
        start_x = Inches(0.8)
        start_y = Inches(1.7)

        for i, c in enumerate(cards):
            cur_x = start_x + i * (card_w + gap)

            # Outer Card
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cur_x, start_y, card_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = colors["card_bg"]
            card.line.color.rgb = colors["card_border"]
            card.line.width = Pt(1)

            # Top colored accent border
            top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cur_x + Inches(0.2), start_y + Inches(0.2), Inches(1.0), Inches(0.04))
            top_bar.fill.solid()
            top_bar.fill.fore_color.rgb = colors["secondary"]
            top_bar.line.fill.background()

            # Content box
            tb = slide.shapes.add_textbox(cur_x + Inches(0.25), start_y + Inches(0.35), card_w - Inches(0.5), card_h - Inches(0.5))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

            p_t = tf.paragraphs[0]
            p_t.text = c.get("title", f"Aspekt {i+1}")
            p_t.font.size = Pt(15)
            p_t.font.bold = True
            p_t.font.color.rgb = colors["primary"]

            # Point 1 box
            p1 = tf.add_paragraph()
            p1.text = f"• {c.get('p1', '')}"
            p1.font.size = Pt(12)
            p1.font.color.rgb = colors["dark_text"]
            p1.space_before = Pt(16)

            # Point 2 box
            p2 = tf.add_paragraph()
            p2.text = f"• {c.get('p2', '')}"
            p2.font.size = Pt(12)
            p2.font.color.rgb = colors["dark_text"]
            p2.space_before = Pt(14)

    @classmethod
    def _render_stats(cls, slide, data: Dict[str, Any], colors: Dict[str, RGBColor]):
        metrics = data.get("metrics", [])[:3]
        box_w = Inches(3.68)
        box_h = Inches(4.7)
        gap = Inches(0.34)
        start_x = Inches(0.8)
        start_y = Inches(1.7)

        for i, m in enumerate(metrics):
            cur_x = start_x + i * (box_w + gap)

            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cur_x, start_y, box_w, box_h)
            box.fill.solid()
            box.fill.fore_color.rgb = colors["card_bg"]
            box.line.color.rgb = colors["card_border"]
            box.line.width = Pt(1)

            tb = slide.shapes.add_textbox(cur_x + Inches(0.3), start_y + Inches(0.5), box_w - Inches(0.6), box_h - Inches(0.8))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

            # Big Stat Value
            p_v = tf.paragraphs[0]
            p_v.text = m.get("value", "100%")
            p_v.font.size = Pt(40)
            p_v.font.bold = True
            p_v.font.color.rgb = colors["secondary"]
            p_v.alignment = PP_ALIGN.CENTER

            # Label
            p_l = tf.add_paragraph()
            p_l.text = m.get("label", "")
            p_l.font.size = Pt(16)
            p_l.font.bold = True
            p_l.font.color.rgb = colors["primary"]
            p_l.alignment = PP_ALIGN.CENTER
            p_l.space_before = Pt(16)

            # Desc
            p_d = tf.add_paragraph()
            p_d.text = m.get("desc", "")
            p_d.font.size = Pt(12)
            p_d.font.color.rgb = colors["muted_text"]
            p_d.alignment = PP_ALIGN.CENTER
            p_d.space_before = Pt(14)

    @classmethod
    def _render_comparative(cls, slide, data: Dict[str, Any], colors: Dict[str, RGBColor]):
        col_w = Inches(5.65)
        col_h = Inches(4.7)
        gap = Inches(0.43)
        start_x = Inches(0.8)
        start_y = Inches(1.7)

        # Left Column
        box1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_x, start_y, col_w, col_h)
        box1.fill.solid()
        box1.fill.fore_color.rgb = colors["card_bg"]
        box1.line.color.rgb = colors["card_border"]

        tb1 = slide.shapes.add_textbox(start_x + Inches(0.3), start_y + Inches(0.3), col_w - Inches(0.6), col_h - Inches(0.6))
        tf1 = tb1.text_frame
        tf1.word_wrap = True
        p1_h = tf1.paragraphs[0]
        p1_h.text = data.get("col1_title", "1-Yo'nalish")
        p1_h.font.size = Pt(17)
        p1_h.font.bold = True
        p1_h.font.color.rgb = colors["primary"]

        for pt in data.get("col1_points", []):
            p = tf1.add_paragraph()
            p.text = f"• {pt}"
            p.font.size = Pt(12)
            p.font.color.rgb = colors["dark_text"]
            p.space_before = Pt(12)

        # Right Column
        box2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, start_x + col_w + gap, start_y, col_w, col_h)
        box2.fill.solid()
        box2.fill.fore_color.rgb = colors["card_bg"]
        box2.line.color.rgb = colors["card_border"]

        tb2 = slide.shapes.add_textbox(start_x + col_w + gap + Inches(0.3), start_y + Inches(0.3), col_w - Inches(0.6), col_h - Inches(0.6))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        p2_h = tf2.paragraphs[0]
        p2_h.text = data.get("col2_title", "2-Yo'nalish")
        p2_h.font.size = Pt(17)
        p2_h.font.bold = True
        p2_h.font.color.rgb = colors["secondary"]

        for pt in data.get("col2_points", []):
            p = tf2.add_paragraph()
            p.text = f"• {pt}"
            p.font.size = Pt(12)
            p.font.color.rgb = colors["dark_text"]
            p.space_before = Pt(12)

    @classmethod
    def _render_conclusion(cls, slide, data: Dict[str, Any], colors: Dict[str, RGBColor]):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.733), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = colors["card_bg"]
        card.line.color.rgb = colors["card_border"]

        tb = slide.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.933), Inches(4.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p_h = tf.paragraphs[0]
        p_h.text = "Asosiy Xulosalar va Amaliy Takliflar"
        p_h.font.size = Pt(18)
        p_h.font.bold = True
        p_h.font.color.rgb = colors["primary"]

        for pt in data.get("points", []):
            p = tf.add_paragraph()
            p.text = f"✔ {pt}"
            p.font.size = Pt(13)
            p.font.color.rgb = colors["dark_text"]
            p.space_before = Pt(14)

        p_ty = tf.add_paragraph()
        p_ty.text = data.get("thank_you", "E'tiboringiz uchun rahmat!")
        p_ty.font.size = Pt(16)
        p_ty.font.bold = True
        p_ty.font.color.rgb = colors["secondary"]
        p_ty.space_before = Pt(24)
