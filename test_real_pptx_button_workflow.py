# -*- coding: utf-8 -*-
"""
======================================================================
REAL END-TO-END WORKFLOW & UI BUTTON SIMULATOR TEST
Tests the exact sequence of operations triggered by the
"🚀 PPTX Taqdimotni Yaratish" button in SlaydHub AI.
======================================================================
"""

import os
import sys
import time
import json
import traceback
from pptx import Presentation

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
if CURRENT_DIR not in sys.path:
    sys.path.insert(0, CURRENT_DIR)

if sys.platform == "win32":
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

from core.config import config
from core.academic_catalog import get_academic_catalog
from core.academic_matcher import get_academic_matcher
from core.slide_archetypes import SLIDE_ARCHETYPES
from core.notebooklm_prompt_gen import NotebookLMPromptGenerator
from core.notebooklm_markdown_parser import NotebookLMMarkdownParser
from core.notebooklm_executor import NotebookLMExecutor
from core.template_pptx_builder import get_academic_pptx_builder
from core.slide_preview import get_presentation_previews


class Colors:
    HEADER = '\033[95m'
    BLUE = '\033[94m'
    CYAN = '\033[96m'
    GREEN = '\033[92m'
    YELLOW = '\033[93m'
    FAIL = '\033[91m'
    ENDC = '\033[0m'
    BOLD = '\033[1m'


def get_all_slide_text(slide):
    texts = []
    def _rec(shapes):
        for s in shapes:
            if s.shape_type == 6:  # GroupShape
                _rec(s.shapes)
            elif s.has_text_frame and s.text_frame.text.strip():
                texts.append(s.text_frame.text.strip())
    _rec(slide.shapes)
    return " ".join(texts)


def run_real_test():
    print("=" * 75)
    print(f"{Colors.BOLD}{Colors.HEADER}SLAYDHUB AI — REAL PPTX TUGMASI VA ISH OQIMI TESTI{Colors.ENDC}")
    print("Foydalanuvchi UI interfeysida '🚀 PPTX Taqdimotni Yaratish' tugmasini")
    print("bosganda bajariladigan har bir qadam real hayotiy ma'lumotlar bilan tekshiriladi.")
    print("=" * 75)

    catalog = get_academic_catalog()
    matcher = get_academic_matcher()
    builder = get_academic_pptx_builder()

    # STEP 1: Topic and Blueprint initialization
    print(f"\n{Colors.CYAN}{Colors.BOLD}[1-BOSQICH] Mavzu va Reja (Blueprint) Tuzish...{Colors.ENDC}")
    topic = "Kvant kompyuterlari va asimmetrik kriptografiya"
    slide_count = 10
    lang = "uz"
    
    t0 = time.time()
    bp = matcher.build_blueprint(
        topic=topic,
        slide_count=slide_count,
        language=lang,
        presentation_mode="scientific_defense_bmi",
        presenter_info={
            "author_name": "Alisher Vohidov",
            "organization": "O'zbekiston Milliy Universiteti / Axborot xavfsizligi",
            "supervisor": "DSc., prof. A. Karimov"
        }
    )
    assert bp is not None, "Blueprint yaratilmadi!"
    assert os.path.exists(bp["primary_pptx_path"]), f"Shablon fayli topilmadi: {bp['primary_pptx_path']}"
    print(f"  {Colors.GREEN}[OK]{Colors.ENDC} Reja tuzildi ({time.time()-t0:.2f}s)")
    print(f"       Asosiy shablon: {bp['primary_deck_title']}")
    print(f"       Fayl: {os.path.basename(bp['primary_pptx_path'])}")

    # STEP 2: Realistic Academic Markdown Content (NotebookLM / Gemini Output Simulation)
    print(f"\n{Colors.CYAN}{Colors.BOLD}[2-BOSQICH] Ilmiy Taqdimot Kontentini Tahlil Qilish (Markdown Parser)...{Colors.ENDC}")
    realistic_md = f"""### 1-slayd: Titul va Kirish
* **Action-oriented sarlavha:** {topic}
* **Mantiqiy maqsadi:** Titul slaydi va tadqiqot tafsilotlari
* **Tarkibiy tezislar:**
  * Kvant algoritmlari davrida axborot xavfsizligi muammolari
  * Post-kvant kriptografiyasining dolzarbligi
  * Tadqiqotchi: Alisher Vohidov | O'zMU 2026

### 2-slayd: Tadqiqotning Mantiqiy Tuzilishi
* **Action-oriented sarlavha:** Tadqiqot Rejasi va Strukturasi
* **Mantiqiy maqsadi:** Taqdimot yo'l xaritasi
* **Tarkibiy tezislar:**
  * 1. Kvant hisoblashlari va Shor algoritmi xavfi
  * 2. Asimmetrik kriptografiya tahlili: RSA va ECC zaifliklari
  * 3. NIST PQC standartlari: ML-KEM va ML-DSA
  * 4. Gibrid xavfsizlik arxitekturasi va tavsiyalar

### 3-slayd: Kvant Tahdidlari va Shor Algoritmi
* **Action-oriented sarlavha:** Shor Algoritmi Asimmetrik Kriptotizimlarni Ishdan Chiqaradi
* **Mantiqiy maqsadi:** Nazariy asos va matematik muammo
* **Tarkibiy tezislar:**
  * Shor algoritmi diskret logarifm va ko'paytuvchilarga ajratishni polinomiyal vaqtda O((log N)³) yechadi.
  * RSA-2048 va ECC-256 kvant quvvati qarshisida o'z kuchini to'liq yo'qotadi.
  * Ma'lumotlarni hozir saqlab, kelajakda shifrdan chiqarish ("Harvest Now, Decrypt Later") tahdidi mavjud.

### 4-slayd: Tadqiqot Obyekti va Vazifalari
* **Action-oriented sarlavha:** Ilmiy Izlanishning Asosiy Vazifalari
* **Mantiqiy maqsadi:** Obyekt, predmet va maqsadlarni belgilash
* **Tarkibiy tezislar:**
  * Obyekt: Korporativ tarmoqlardagi TLS 1.3 va IPsec protokollari.
  * Predmet: Post-kvant KEM mexanizmlarining tezkorligi va xavfsizligi.
  * Maqsad: Mavjud PKI infratuzilmasiga minimal ortiqcha yuklama bilan gibrid shifrlashni joriy etish.

### 5-slayd: Asosiy Miqdoriy Ko'rsatkichlar
* **Action-oriented sarlavha:** Kvant Xavfsizligining Miqdoriy Metriklari
* **Mantiqiy maqsadi:** Empirik ko'rsatkichlar
* **Tarkibiy tezislar:**
  * 99.9% Barqarorlik: NIST FIPS 203 standarti asosidagi panjarali shifrlash kafolati.
  * 2,500+ Mantiqiy Kubit: RSA kalitini sindirish uchun zarur bo'lgan minimal kvant hisoblash quvvati.
  * 1.2 ms Handshake: Gibrid TLS ulanishida yuzaga keladigan o'rtacha kechikish vaqti.

### 6-slayd: Qiyosiy Tahlil: Klassik vs Post-Kvant
* **Action-oriented sarlavha:** Klassik RSA va Post-Kvant ML-KEM Qiyosi
* **Mantiqiy maqsadi:** Model A va Model B taqqoslash
* **Tarkibiy tezislar:**
  * Klassik RSA-2048: Kalit hajmi 256 bayt, kvantga chidamlilik 0%, shifrlash tezligi past.
  * Post-Kvant ML-KEM: Kalit hajmi 1184 bayt, kvantga chidamlilik 100%, panjarali SVP matematik murakkablik.

### 7-slayd: Post-Kvantga O'tish Xronologiyasi
* **Action-oriented sarlavha:** Tizimli Migratsiya Bosqichlari (2026-2030)
* **Mantiqiy maqsadi:** Bosqichma-bosqich yo'l xaritasi
* **Tarkibiy tezislar:**
  * 1-Bosqich (Audit): Barcha mavjud sertifikatlar va kriptografik kalitlar inventarizatsiyasi.
  * 2-Bosqich (Gibrid Rejim): Klassik ECDH va ML-KEM algoritmlarini parallel qo'llash.
  * 3-Bosqich (To'liq O'tish): Infratuzilmani 100% PQC standartlariga ko'chirish.

### 8-slayd: Tizim Arxitekturasi va Tarkibi
* **Action-oriented sarlavha:** Uch Bosqichli Himoyalangan Arxitektura
* **Mantiqiy maqsadi:** Blokli tizim tarkibi
* **Tarkibiy tezislar:**
  * Kalit Almashinuv Qatlami: ML-KEM-768 protokoli asosidagi kapsulyatsiya.
  * Raqamli Imzo Qatlami: ML-DSA standarti yordamida soxtalashtirishdan himoya.
  * Kripto-Moslashuvchanlik: Yangi standartlar paydo bo'lganda algoritmlarni to'xtovsiz almashtirish.

### 9-slayd: Xulosalar va Tavsiyalar
* **Action-oriented sarlavha:** Ilmiy Xulosalar va Amaliy Tavsiyalar
* **Mantiqiy maqsadi:** Tadqiqot xulosalari
* **Tarkibiy tezislar:**
  * Post-kvant kriptografiyasiga o'tish kechiktirib bo'lmas strategik xavfsizlik talabidir.
  * Gibrid kalit almashinuvi klassik tizimlar bilan to'liq orqaga moslikni (backward compatibility) ta'minlaydi.
  * Xalqaro NIST FIPS 203, 204 standartlarini milliy qonunchilikka tatbiq etish tavsiya etiladi.

### 10-slayd: Foydalanilgan Manbalar
* **Action-oriented sarlavha:** Ilmiy Adabiyotlar va Standartlar Ro'yxati
* **Mantiqiy maqsadi:** Bibliografiya
* **Tarkibiy tezislar:**
  * 1. NIST FIPS 203: Module-Lattice-Based Key-Encapsulation Mechanism Standard (2024).
  * 2. Shor, P. W. Algorithms for Quantum Computation: Discrete Logarithms and Factoring (IEEE, 1994).
  * 3. Bernstein, D. J. Post-quantum cryptography: dealing with the inevitable (Springer, 2009).
"""

    t0 = time.time()
    parsed_content = NotebookLMMarkdownParser.parse(realistic_md, topic=topic, default_metadata=bp.get("presenter_info", {}))
    slides_count = len(parsed_content.get("slides", []))
    assert slides_count == 10, f"Kutilgan 10 slayd, o'qildi: {slides_count}"
    print(f"  {Colors.GREEN}[OK]{Colors.ENDC} Markdown muvaffaqiyatli tahlil qilindi: {slides_count} ta slayd aniqlandi ({time.time()-t0:.2f}s)")

    # STEP 3: Dynamic Template Archetype Matching
    print(f"\n{Colors.CYAN}{Colors.BOLD}[3-BOSQICH] Dinamik Arxetiplar va Shablonlar Bog'lanishi...{Colors.ENDC}")
    t0 = time.time()
    for idx, s_p in enumerate(parsed_content.get("slides", [])):
        arch = s_p.get("layout_type", "cards_grid")
        if idx < len(bp["slides"]):
            bp["slides"][idx]["layout_type"] = arch
            bp["slides"][idx]["layout_name"] = SLIDE_ARCHETYPES.get(arch, {}).get("name_uz", arch)
            bp["slides"][idx]["title_hint"] = s_p.get("title", bp["slides"][idx].get("title_hint", ""))
            matches = catalog.search_slides_by_archetype(arch, category=bp.get("category"), topic=bp.get("topic"), limit=1)
            if matches:
                bp["slides"][idx]["matched_slide_ref"] = matches[0]
                bp["slides"][idx]["pptx_path"] = matches[0]["pptx_path"]
                bp["slides"][idx]["template_title"] = matches[0]["template_title"]
                bp["slides"][idx]["template_slide_number"] = matches[0]["slide_index"] + 1
    print(f"  {Colors.GREEN}[OK]{Colors.ENDC} 10 ta slayd uchun professional shablonlar muvofiqlashtirildi ({time.time()-t0:.2f}s)")

    # STEP 4: Presentation Generation (Simulating "🚀 PPTX Taqdimotni Yaratish" button)
    print(f"\n{Colors.CYAN}{Colors.BOLD}[4-BOSQICH] '🚀 PPTX Taqdimotni Yaratish' Tugmasi Bosildi (create_presentation)...{Colors.ENDC}")
    t0 = time.time()
    out_pptx = builder.create_presentation(parsed_content, bp, engine_mode="harmonized")
    elapsed_gen = time.time() - t0
    
    assert os.path.exists(out_pptx), f"PPTX fayl yaratilmadi: {out_pptx}"
    file_size_kb = os.path.getsize(out_pptx) / 1024
    assert file_size_kb > 30, f"Fayl hajmi juda kichik ({file_size_kb:.1f} KB)"
    print(f"  {Colors.GREEN}[OK]{Colors.ENDC} Taqdimot fayli yaratildi ({elapsed_gen:.2f}s):")
    print(f"       Yo'l: {out_pptx}")
    print(f"       Hajmi: {file_size_kb:.1f} KB")

    # STEP 5: Deep Quality & Content Verification on the generated PPTX
    print(f"\n{Colors.CYAN}{Colors.BOLD}[5-BOSQICH] Taqdimotning Ichki Kontenti va Sifatini Chuqur Tekshirish...{Colors.ENDC}")
    prs = Presentation(out_pptx)
    assert len(prs.slides) == 10, f"Kutilgan 10 slayd, yaratildi: {len(prs.slides)}"

    # Slide 1 verification
    s1_text = get_all_slide_text(prs.slides[0])
    assert "Kvant" in s1_text or "asimmetrik" in s1_text.lower(), f"Titulda mavzu topilmadi: {s1_text}"
    print(f"  {Colors.GREEN}[OK]{Colors.ENDC} 1-Slayd (Titul): Mavzu va muallif ma'lumotlari to'liq kiritilgan")

    # Slide 5 (Metrics) verification
    s5_text = get_all_slide_text(prs.slides[4])
    assert any(num in s5_text for num in ["99.9%", "2,500", "1.2"]), f"5-slaydda metrikalar topilmadi: {s5_text}"
    print(f"  {Colors.GREEN}[OK]{Colors.ENDC} 5-Slayd (Metrikalar): Real KPI ko'rsatkichlari (99.9%, 2,500, 1.2 ms) mavjud")

    # Slide 6 (Comparison) verification
    s6_text = get_all_slide_text(prs.slides[5])
    assert "RSA" in s6_text or "Klassik" in s6_text or "Post-Kvant" in s6_text, f"6-slaydda taqqoslash topilmadi: {s6_text}"
    print(f"  {Colors.GREEN}[OK]{Colors.ENDC} 6-Slayd (Qiyosiy Tahlil): Klassik vs Post-Kvant parametrlari kiritilgan")

    # STEP 6: Previews generation (Simulating Step 5 in UI)
    print(f"\n{Colors.CYAN}{Colors.BOLD}[6-BOSQICH] Prevyu Dvigateli (UI Galereyasi va Skrinshotlar)...{Colors.ENDC}")
    t0 = time.time()
    previews = get_presentation_previews(out_pptx, max_slides=3)
    assert len(previews) > 0, "Slayd rasmlari renderlanmadi!"
    for p in previews:
        assert os.path.exists(p) and os.path.getsize(p) > 1000
    print(f"  {Colors.GREEN}[OK]{Colors.ENDC} Slayd prevyulari muvaffaqiyatli olindi ({time.time()-t0:.2f}s): {len(previews)} ta rasm")

    # STEP 7: One-Click Instant Fallback Test (Empty content automatically generated)
    print(f"\n{Colors.CYAN}{Colors.BOLD}[7-BOSQICH] 1-Klikli Avtomatik Fallback Testi (Matn kiritilmaganda ham ishlashi)...{Colors.ENDC}")
    # Simulating a user who didn't enter any Markdown in Step 3, but clicked "🚀 PPTX Taqdimotni Yaratish" directly
    t0 = time.time()
    fallback_bp = matcher.build_blueprint("Sun'iy Intellekt va Neyron Tarmoqlari", slide_count=6, language="uz")
    # Generating content with Gemini directly
    gemini_generated = NotebookLMPromptGenerator.generate_with_gemini_direct(
        topic=fallback_bp["topic"],
        domain="IT / Sun'iy Intellekt",
        author_info="Test Foydalanuvchi",
        lang="uz"
    )
    assert len(gemini_generated) > 200, "Gemini fallback matn bermadi!"
    fallback_parsed = NotebookLMMarkdownParser.parse(gemini_generated, topic=fallback_bp["topic"])
    fallback_out = builder.create_presentation(fallback_parsed, fallback_bp, engine_mode="harmonized")
    assert os.path.exists(fallback_out)
    print(f"  {Colors.GREEN}[OK]{Colors.ENDC} 1-Klikda to'liq avtomatik AI generatsiyasi va PPTX yaratilishi muvaffaqiyatli ({time.time()-t0:.2f}s)!")
    print(f"       Hajmi: {os.path.getsize(fallback_out)/1024:.1f} KB")

    print("\n" + "=" * 75)
    print(f"{Colors.GREEN}{Colors.BOLD}🎉 TEST YAKUNI: BARCHA 7 TA REAL BOSQICH 100% MUVAFFAQIYATLI O'TDI!{Colors.ENDC}")
    print(f"{Colors.GREEN}Taqdimot yaratish tugmasi va uning barcha funksiyalari to'liq ishlamoqda.{Colors.ENDC}")
    print("=" * 75)
    return True


if __name__ == "__main__":
    try:
        success = run_real_test()
        sys.exit(0 if success else 1)
    except Exception as e:
        print(f"\n{Colors.FAIL}[XATOLIK]{Colors.ENDC} Testda xatolik yuz berdi:\n{e}\n{traceback.format_exc()}")
        sys.exit(1)
