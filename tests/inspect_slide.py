from pptx import Presentation

prs = Presentation('output/academic_studio/Kvant_kompyuterlari_va_asimmetrik_kripto_Customized.pptx')
for idx, s in enumerate(prs.slides):
    found = False
    for sh in s.shapes:
        if sh.has_text_frame and 'HNDL' in sh.text_frame.text:
            found = True
            break
    if found:
        print(f"Slide {idx+1} contains HNDL! Layout: '{s.slide_layout.name}'")
        print(f"Total placeholders on slide: {len(s.placeholders)}")
        for i, p in enumerate(s.placeholders):
            txt = p.text_frame.text if p.has_text_frame else ""
            print(f"  Placeholder {i}: name='{p.name}', text='{txt[:40]}'")
        print(f"Total shapes on slide: {len(s.shapes)}")
        for i, shp in enumerate(s.shapes):
            txt = shp.text_frame.text if shp.has_text_frame else ""
            print(f"  Shape {i}: name='{shp.name}', is_ph={shp.is_placeholder}, text='{txt[:40]}'")
