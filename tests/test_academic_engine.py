# -*- coding: utf-8 -*-
"""
Comprehensive Unit Test Suite for SlaydHub Academic Presentation Engine.
Tests Catalog, Matcher, Prompt Generator, Parser, Spatial Replacer, Topic Logic Validator,
Spatial Collision Engine, Semantic Shape Cleaner, and PPTX Builder.
"""

import os
import sys
import unittest
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

# Add project root to sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from core.academic_catalog import get_academic_catalog, AcademicCatalog, ACADEMIC_CATEGORIES
from core.academic_matcher import get_academic_matcher, AcademicMatcher
from core.notebooklm_prompt_gen import generate_notebooklm_prompt, NotebookLMPromptGenerator
from core.notebooklm_executor import NotebookLMExecutor
from core.topic_logic_validator import TopicLogicValidator
from core.pptx_text_replacer import (
    PPTXTextReplacer,
    SemanticShapeCleaner,
    SpatialCollisionEngine,
    SafeContentZoneValidator,
    DeepWatermarkPurger,
    replace_presentation_content
)
from core.template_pptx_builder import get_academic_pptx_builder


class TestAcademicCatalog(unittest.TestCase):
    """Unit tests for template indexing, searching, and query expansion."""

    @classmethod
    def setUpClass(cls):
        cls.catalog = get_academic_catalog()

    def test_catalog_indexing(self):
        """Verify that authentic PPTX templates are indexed."""
        self.assertGreater(len(self.catalog.templates), 10, "Should have indexed authentic PPTX templates")
        sample = self.catalog.templates[0]
        self.assertIn("id", sample)
        self.assertIn("title", sample)
        self.assertIn("pptx_path", sample)
        self.assertTrue(sample.get("has_pptx"), "Indexed templates must have has_pptx=True")

    def test_multilingual_expansion(self):
        """Verify that Uzbek terms expand to English search terms."""
        terms_kvant = self.catalog.expand_query_terms("kvant")
        self.assertIn("quantum", terms_kvant)

        terms_tib = self.catalog.expand_query_terms("tibbiyot")
        self.assertIn("medical", terms_tib)

        terms_ai = self.catalog.expand_query_terms("dasturlash")
        self.assertTrue(any(t in ["programming", "coding", "software"] for t in terms_ai))

    def test_search_accuracy(self):
        """Verify search finds relevant templates for academic topics."""
        for query in ["quantum", "medical", "cyber", "history", "economics"]:
            results = self.catalog.search_templates(query, limit=5)
            self.assertGreater(len(results), 0, f"Query '{query}' should return results")
            for r in results:
                self.assertTrue(os.path.exists(r["pptx_path"]), f"Template path must exist: {r['pptx_path']}")


class TestAcademicMatcher(unittest.TestCase):
    """Unit tests for topic category detection, blueprint planning, and slide layout sequencing."""

    @classmethod
    def setUpClass(cls):
        cls.matcher = get_academic_matcher()

    def test_category_detection(self):
        """Verify automatic academic discipline detection."""
        cat_phys = self.matcher.detect_category("Termodinamika va Kvant mexanikasi")
        self.assertEqual(cat_phys, "exact_sciences")

        cat_med = self.matcher.detect_category("Klinik anatomiya va yurak qon-tomir kasalliklari")
        self.assertEqual(cat_med, "medicine_health")

        cat_it = self.matcher.detect_category("Kiberxavfsizlik va kriptografik algoritmlar")
        self.assertEqual(cat_it, "engineering_it")

    def test_blueprint_generation(self):
        """Verify narrative blueprint layout sequence for different slide counts."""
        for count in [5, 8, 10, 12, 15]:
            bp = self.matcher.build_blueprint("Sun'iy intellekt arxitekturasi", slide_count=count, language="uz")
            self.assertEqual(bp["slide_count"], count)
            self.assertEqual(len(bp["slides"]), count)
            self.assertEqual(bp["slides"][0]["layout_type"], "cover")
            self.assertEqual(bp["slides"][-1]["layout_type"], "conclusion_sources")
            self.assertTrue(os.path.exists(bp["primary_pptx_path"]), "Blueprint must have valid PPTX path")

    def test_multilingual_title_hints(self):
        """Verify title hints in UZ, RU, and EN."""
        hint_uz = self.matcher.generate_slide_title_hint(2, "agenda", "Test", "uz")
        self.assertIn("Mundarija", hint_uz)

        hint_ru = self.matcher.generate_slide_title_hint(2, "agenda", "Test", "ru")
        self.assertIn("Оглавление", hint_ru)

        hint_en = self.matcher.generate_slide_title_hint(2, "agenda", "Test", "en")
        self.assertIn("Table of Contents", hint_en)


class TestTopicLogicValidator(unittest.TestCase):
    """Unit tests for topic consistency validation, text sanitization, and bullet optimization."""

    def test_sanitize_placeholder_text(self):
        dirty = "Lorem ipsum dolor sit amet presentationgo.com Kvant shifrlash asoslari"
        cleaned = TopicLogicValidator.sanitize_text_field(dirty)
        self.assertNotIn("Lorem ipsum", cleaned)
        self.assertNotIn("presentationgo", cleaned)
        self.assertIn("Kvant shifrlash asoslari", cleaned)

    def test_optimize_bullet_length(self):
        long_text = "Ushbu tadqiqotda kvant kompyuterlarining zamonaviy asimmetrik kriptografik tizimlarga ta'siri juda keng va atroflicha o'rganildi hamda amaliy tavsiyalar berildi va barcha mezonlar hisobga olindi."
        optimized = TopicLogicValidator.optimize_bullet_length(long_text, max_words=12)
        words = optimized.split()
        self.assertLessEqual(len(words), 14)

    def test_validate_and_refine_presentation(self):
        raw_data = {
            "topic": "Post-Kvant Kriptografiyasi",
            "slides": [
                {"slide_number": 1, "layout_type": "cover", "title": "Lorem Ipsum Kripto", "subtitle": "slidenest.com tavsifi"},
                {"slide_number": 2, "layout_type": "agenda", "points": ["1. Kirish va dolzarblik", "2. Asosiy maqsad va vazifalar"]}
            ]
        }
        refined = TopicLogicValidator.validate_and_refine_presentation(raw_data)
        self.assertEqual(len(refined["slides"]), 2)
        self.assertNotIn("Lorem Ipsum", refined["slides"][0]["title"])
        self.assertNotIn("slidenest.com", refined["slides"][0]["subtitle"])


class TestNotebookLMPromptGenerator(unittest.TestCase):
    """Unit tests for prompt construction and schema consistency."""

    def test_prompt_generation(self):
        matcher = get_academic_matcher()
        bp = matcher.build_blueprint("Genetik modifikatsiyalangan organizmlar", slide_count=8, language="uz")
        prompt = generate_notebooklm_prompt(bp)

        self.assertIsInstance(prompt, str)
        self.assertGreater(len(prompt), 1000)
        self.assertIn("Genetik modifikatsiyalangan organizmlar", prompt)
        self.assertIn('"layout_type": "cover"', prompt)
        self.assertIn('"layout_type": "agenda"', prompt)
        self.assertIn('"layout_type": "conclusion_sources"', prompt)


class TestNotebookLMExecutorParser(unittest.TestCase):
    """Unit tests for robust JSON parsing of LLM/NotebookLM output."""

    def test_clean_json_parsing(self):
        raw = '{"topic": "Kvant", "slides": [{"slide_number": 1, "layout_type": "cover"}]}'
        parsed = NotebookLMExecutor.parse_notebooklm_response(raw)
        self.assertEqual(parsed["topic"], "Kvant")
        self.assertEqual(len(parsed["slides"]), 1)

    def test_markdown_codeblock_parsing(self):
        raw = 'Here is the response:\n```json\n{"topic": "Test", "slides": [{"slide_number": 1}]}\n```\nHope it helps!'
        parsed = NotebookLMExecutor.parse_notebooklm_response(raw)
        self.assertEqual(parsed["topic"], "Test")

    def test_trailing_comma_repair(self):
        raw = '{"topic": "Test", "slides": [{"slide_number": 1,},],}'
        parsed = NotebookLMExecutor.parse_notebooklm_response(raw)
        self.assertEqual(parsed["topic"], "Test")

    def test_empty_input_raises(self):
        with self.assertRaises(ValueError):
            NotebookLMExecutor.parse_notebooklm_response("")


class TestPPTXTextReplacerAndSpatialEngine(unittest.TestCase):
    """Unit tests for in-place text replacement, high-legibility font scaling, and shape pruning."""

    def test_dummy_text_detection(self):
        """Verify dummy text patterns are recognized."""
        self.assertTrue(PPTXTextReplacer.is_dummy_text("Lorem ipsum dolor sit amet"))
        self.assertTrue(PPTXTextReplacer.is_dummy_text("Make a big impact with our professional slides"))
        self.assertTrue(PPTXTextReplacer.is_dummy_text("Paris"))
        self.assertTrue(PPTXTextReplacer.is_dummy_text("Berlin"))
        self.assertTrue(PPTXTextReplacer.is_dummy_text("Karla Gard"))
        self.assertTrue(PPTXTextReplacer.is_dummy_text("CFO"))
        self.assertFalse(PPTXTextReplacer.is_dummy_text("Kvant kompyuterlari va kriptografiya"))

    def test_step_number_detection(self):
        """Verify step numbers and badges."""
        self.assertTrue(PPTXTextReplacer.is_step_number("01"))
        self.assertTrue(PPTXTextReplacer.is_step_number("01."))
        self.assertTrue(PPTXTextReplacer.is_step_number("04"))
        self.assertFalse(PPTXTextReplacer.is_step_number("Mundarija va Reja"))

    def test_high_legibility_font_floors(self):
        """Verify that fonts never drop below comfortable minimum reading thresholds."""
        # Long body text in a small box
        fs_body = PPTXTextReplacer.calculate_safe_font_size(
            "Bu juda uzun ilmiy matn bo'lib, shrift o'lchami o'qishga juda qulay bo'lishi kerak.",
            width_pt=100.0,
            height_pt=40.0,
            max_initial_fs=16.0,
            min_fs=11.5,
            is_title=False
        )
        self.assertGreaterEqual(fs_body, 11.0, "Body font size must not drop below 11pt")

        # Title font size in a wide banner
        fs_title = PPTXTextReplacer.calculate_safe_font_size(
            "Kvant Kompyuterlari va Asimmetrik Kriptografiya",
            width_pt=600.0,
            height_pt=80.0,
            max_initial_fs=32.0,
            min_fs=22.0,
            is_title=True
        )
        self.assertGreaterEqual(fs_title, 22.0, "Title font size must not drop below 22pt")

    def test_bounding_box_collision_detection(self):
        """Verify 2D bounding box intersection calculation."""
        # Box 1: (0, 0, 100, 100), Box 2: (50, 50, 150, 150) -> Overlap!
        b1 = (0.0, 0.0, 100.0, 100.0)
        b2 = (50.0, 50.0, 150.0, 150.0)
        self.assertTrue(SpatialCollisionEngine.has_overlap(b1, b2))

        # Box 3: (200, 200, 300, 300) -> No overlap with Box 1
        b3 = (200.0, 200.0, 300.0, 300.0)
        self.assertFalse(SpatialCollisionEngine.has_overlap(b1, b3))

    def test_semantic_shape_pruning(self):
        """Verify XML element pruning physically removes the shape."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shp = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        self.assertEqual(len(slide.shapes), 1)

        pruned = SemanticShapeCleaner.prune_shape(shp)
        self.assertTrue(pruned)
        self.assertEqual(len(slide.shapes), 0)

    def test_in_place_slide_customization(self):
        """Verify that replacing text in a real template produces valid output with 0 dummy text and high legibility."""
        catalog = get_academic_catalog()
        template = catalog.templates[0]
        template_path = template["pptx_path"]

        sample_content = {
            "topic": "Unit Test Presentation",
            "slides": [
                {
                    "slide_number": 1,
                    "layout_type": "cover",
                    "title": "Unit Test Title",
                    "subtitle": "Unit Test Subtitle | 2026"
                },
                {
                    "slide_number": 2,
                    "layout_type": "agenda",
                    "points": ["1. Birinchi Reja", "2. Ikkinchi Reja", "3. Uchinchi Reja", "4. To'rtinchi Reja"]
                }
            ]
        }

        out_dir = os.path.join(PROJECT_ROOT, "output", "test_runs")
        os.makedirs(out_dir, exist_ok=True)
        out_file = os.path.join(out_dir, "test_customized_output.pptx")

        result_path = replace_presentation_content(template_path, sample_content, out_file)
        self.assertTrue(os.path.exists(result_path))
        self.assertGreater(os.path.getsize(result_path), 10000)

        # Inspect resulting presentation
        prs = Presentation(result_path)
        self.assertEqual(len(prs.slides), 2)
        
        # Verify slide 1 title
        slide1_texts = [s.text_frame.text.strip() for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.strip()]
        self.assertTrue(any("Unit Test Title" in t for t in slide1_texts))

        # Verify no dummy text remaining on slide 1 & 2
        for slide in prs.slides:
            for s in slide.shapes:
                if s.has_text_frame:
                    t = s.text_frame.text.strip()
                    if t and not PPTXTextReplacer.is_step_number(t):
                        self.assertFalse(PPTXTextReplacer.is_dummy_text(t), f"Dummy text found: '{t}'")


    def test_narrow_strip_rejection(self):
        """Verify that narrow shapes (< 120pt) are rejected from receiving body/card text."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # Add narrow shape (width = 80pt)
        narrow_shp = slide.shapes.add_textbox(Inches(1), Inches(1), Pt(80), Pt(100))
        self.assertFalse(SafeContentZoneValidator.is_valid_content_box(narrow_shp, sw_pt=960.0, sh_pt=540.0, is_title=False))

        # Add valid shape (width = 250pt)
        valid_shp = slide.shapes.add_textbox(Inches(1), Inches(1), Pt(250), Pt(100))
        self.assertTrue(SafeContentZoneValidator.is_valid_content_box(valid_shp, sw_pt=960.0, sh_pt=540.0, is_title=False))

    def test_deep_watermark_purger(self):
        """Verify that watermarks in slide, layout, master, and group shapes are eradicated."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        tb.text_frame.text = "Designed by presentationgo.com free templates"
        self.assertEqual(len(slide.shapes), 1)

        from core.pptx_text_replacer import DeepWatermarkPurger
        DeepWatermarkPurger.purge_presentation_watermarks(prs)
        # Either text is emptied or shape is pruned
        remaining_text = "".join([s.text_frame.text for s in slide.shapes if s.has_text_frame])
        self.assertNotIn("presentationgo", remaining_text.lower())

    def test_anti_overflow_condensation(self):
        """Verify that long multiline text in a small box does not overflow boundaries."""
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2.5), Inches(1.0))
        
        long_paragraph = (
            "Kvant hisoblashlari va post-kvant kriptografiyasi sohasida olib borilgan "
            "ko'p yillik fundamental tadqiqotlar shuni ko'rsatadiki, yangi avlod ML-KEM "
            "va ML-DSA standartlari xavfsizlikni to'liq ta'minlaydi va barcha tarmoqlarni "
            "zamonaviy kibertahdidlardan mutlaq himoya qiladi."
        )
        PPTXTextReplacer.set_shape_text_preserving_style(box, long_paragraph, default_font_size=14.0)
        
        # Resulting text should be condensed cleanly
        final_text = box.text_frame.text
        self.assertTrue(len(final_text) > 20)
        self.assertLess(len(final_text.split('\n')), 5)


class TestEndToEndPresentationBuilder(unittest.TestCase):
    """Unit tests for complete presentation pipeline."""

    def test_full_presentation_creation(self):
        matcher = get_academic_matcher()
        bp = matcher.build_blueprint("Kiberxavfsizlik va Ma'lumotlar Himoyasi", slide_count=5, language="uz")

        sample_content = {
            "topic": bp["topic"],
            "slides": [
                {"slide_number": 1, "layout_type": "cover", "title": "Kiberxavfsizlik va Himoya", "subtitle": "Amaliy qo'llanma | 2026"},
                {"slide_number": 2, "layout_type": "agenda", "points": ["1. Xavf turlari", "2. Shifrlash", "3. Xulosa"]},
                {"slide_number": 3, "layout_type": "theory_concept", "title": "Axborot Xavfsizligi", "concept_name": "Kriptografiya", "definition": "Ma'lumotlarni maxfiy saqlash san'ati.", "key_formula_or_law": "AES-256", "importance": "Bank va davlat tizimlarida."},
                {"slide_number": 4, "layout_type": "cards_grid", "cards": [{"title": "Xavf 1", "desc": "Fishing hujumlari."}, {"title": "Xavf 2", "desc": "DDoS hujumlari."}]},
                {"slide_number": 5, "layout_type": "conclusion_sources", "key_takeaways": ["Xavfsizlik eng ustuvor vazifa."], "references": ["1. ISO/IEC 27001"]}
            ]
        }

        builder = get_academic_pptx_builder()
        out_pptx = builder.create_presentation(sample_content, bp)
        self.assertTrue(os.path.exists(out_pptx))
        self.assertGreater(os.path.getsize(out_pptx), 15000)

    def test_multi_archetype_visual_renderers(self):
        """Verify specialized visual renderers for metrics, timeline steps, and vertical columns."""
        prs = Presentation()
        sw_pt, sh_pt = 720.0, 405.0
        prs.slide_width = Inches(sw_pt / 72.0)
        prs.slide_height = Inches(sh_pt / 72.0)
        scale_factor = 1.0
        left_m_pt = 45.0
        content_w_pt = sw_pt - (left_m_pt * 2)
        top_start_pt = 80.0

        # 1. Test KPI / Metric Stat Cards
        s_metric = prs.slides.add_slide(prs.slide_layouts[6])
        metric_data = {
            "layout_type": "metrics_stats",
            "title": "Asosiy Ko'rsatkichlar",
            "metrics": [
                {"value": "99.9%", "label": "Kvant Barqarorligi", "desc": "NIST FIPS 203 standarti"},
                {"value": "2,500", "label": "Kubit Talabi", "desc": "Shor algoritmi amaliyotda"},
                {"value": "1.2 ms", "label": "Kripto Tezlik", "desc": "Gibrid TLS 1.3 integratsiyasi"}
            ]
        }
        metric_ids = PPTXTextReplacer.render_cards_by_archetype(
            slide=s_metric,
            slide_data=metric_data,
            layout_type="metrics_stats",
            content_items=[],
            sw_pt=sw_pt,
            sh_pt=sh_pt,
            scale_factor=scale_factor,
            left_m_pt=left_m_pt,
            content_w_pt=content_w_pt,
            top_start_pt=top_start_pt
        )
        self.assertEqual(len(metric_ids), 3, "Should render 3 KPI metric stat cards")

        # 2. Test Timeline / Process Step Cards
        s_step = prs.slides.add_slide(prs.slide_layouts[6])
        step_data = {
            "layout_type": "timeline_steps",
            "title": "Jarayon Bosqichlari",
            "steps": [
                {"step_number": "01", "title": "Standartlashtirish", "desc": "NIST standartlari tahlili"},
                {"step_number": "02", "title": "Gibrid Protokol", "desc": "TLS 1.3 sinovlari"},
                {"step_number": "03", "title": "To'liq O'tish", "desc": "Infratuzilma migratsiyasi"}
            ]
        }
        step_ids = PPTXTextReplacer.render_cards_by_archetype(
            slide=s_step,
            slide_data=step_data,
            layout_type="timeline_steps",
            content_items=[],
            sw_pt=sw_pt,
            sh_pt=sh_pt,
            scale_factor=scale_factor,
            left_m_pt=left_m_pt,
            content_w_pt=content_w_pt,
            top_start_pt=top_start_pt
        )
        self.assertEqual(len(step_ids), 3, "Should render 3 timeline process step cards")

        # 3. Test Vertical Column Cards (cards_grid)
        s_col = prs.slides.add_slide(prs.slide_layouts[6])
        col_items = [
            ("Lattice Kriptografiya", "Shortest Vector Problem asosidagi murakkablik."),
            ("Gibrid Xavfsizlik", "Klassik va kvant algoritmlarining birgalikdagi ish faoliyati."),
            ("Infratuzilma", "PKI va sertifikatlash markazlarining to'liq muvofiqligi.")
        ]
        col_ids = PPTXTextReplacer.render_cards_by_archetype(
            slide=s_col,
            slide_data={"layout_type": "cards_grid"},
            layout_type="cards_grid",
            content_items=col_items,
            sw_pt=sw_pt,
            sh_pt=sh_pt,
            scale_factor=scale_factor,
            left_m_pt=left_m_pt,
            content_w_pt=content_w_pt,
            top_start_pt=top_start_pt
        )
        self.assertEqual(len(col_ids), 3, "Should render 3 vertical column cards")


if __name__ == "__main__":
    unittest.main()
