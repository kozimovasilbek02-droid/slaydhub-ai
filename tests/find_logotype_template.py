import os
import sys
import glob
sys.path.insert(0, os.path.abspath("."))
from pptx import Presentation
from core.config import config

print("Template dirs:", config.TEMPLATE_DIRS)
found_count = 0
for tdir in config.TEMPLATE_DIRS:
    for root, dirs, files in os.walk(tdir):
        if config.is_folder_blacklisted(root):
            continue
        for f in files:
            if f.lower().endswith(".pptx"):
                p = os.path.join(root, f)
                try:
                    prs = Presentation(p)
                    for s_idx, s in enumerate(prs.slides):
                        for shp in s.shapes:
                            txt = shp.text_frame.text if shp.has_text_frame else ''
                            if 'logotype' in txt.lower():
                                print(f"MATCH: {p} | Slide {s_idx+1} | Shape: {shp.name} | Text: '{txt}'")
                                found_count += 1
                                if found_count >= 10:
                                    break
                        if found_count >= 10:
                            break
                    for m in prs.slide_masters:
                        for shp in m.shapes:
                            txt = shp.text_frame.text if shp.has_text_frame else ''
                            if 'logotype' in txt.lower():
                                print(f"MATCH MASTER: {p} | Shape: {shp.name} | Text: '{txt}'")
                                found_count += 1
                                if found_count >= 10:
                                    break
                        for l in m.slide_layouts:
                            for shp in l.shapes:
                                txt = shp.text_frame.text if shp.has_text_frame else ''
                                if 'logotype' in txt.lower():
                                    print(f"MATCH LAYOUT: {p} | Shape: {shp.name} | Text: '{txt}'")
                                    found_count += 1
                                    if found_count >= 10:
                                        break
                except Exception:
                    pass
            if found_count >= 10:
                break
        if found_count >= 10:
            break
