from pptx import Presentation

prs = Presentation('output/academic_studio/Kvant_kompyuterlari_va_asimmetrik_kripto_Customized.pptx')
for idx, s in enumerate(prs.slides):
    for shp in s.shapes:
        txt = shp.text_frame.text if shp.has_text_frame else ''
        if 'logo' in txt.lower():
            print(f"Slide {idx+1}: Shape '{shp.name}' has logo text: '{txt}'")

for m_idx, m in enumerate(prs.slide_masters):
    for shp in m.shapes:
        txt = shp.text_frame.text if shp.has_text_frame else ''
        if 'logo' in txt.lower():
            print(f"Master {m_idx}: Shape '{shp.name}' has logo text: '{txt}'")
    for l_idx, l in enumerate(m.slide_layouts):
        for shp in l.shapes:
            txt = shp.text_frame.text if shp.has_text_frame else ''
            if 'logo' in txt.lower():
                print(f"Layout {l_idx} ({l.name}): Shape '{shp.name}' has logo text: '{txt}'")
