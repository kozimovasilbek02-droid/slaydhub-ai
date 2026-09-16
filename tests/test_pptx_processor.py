# -*- coding: utf-8 -*-
"""
tests/test_pptx_processor.py
PPTXProcessor modulining reklama slaydlarini aniqlash, suvbelgilar va tozalash testlari.
"""
import sys, pathlib, unittest
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from backend.core.pptx_processor import sanitize_control_chars, PPTXProcessor

class TestControlChars(unittest.TestCase):
    def test_sanitize_spaces(self):
        assert sanitize_control_chars("A\x0bB") == "A B"
        assert sanitize_control_chars("Test_x000B_Text") == "Test Text"

class TestWatermarksAndAds(unittest.TestCase):
    def test_watermark_detection(self):
        assert PPTXProcessor._is_watermark_text("Visit slidescarnival.com for more") is True
        assert PPTXProcessor._is_watermark_text("Designed by presentationgo") is True
        assert PPTXProcessor._is_watermark_text("Leonardo Da Vinci Biography") is False

    def test_ad_slide_detection(self):
        from pptx import Presentation
        prs = Presentation()
        # Create a mock slide with Canva resources
        s_res = prs.slides.add_slide(prs.slide_layouts[6])
        tb1 = s_res.shapes.add_textbox(0, 0, 1000000, 1000000)
        tb1.text_frame.text = "Resurslar sahifasi\nCanva taqdimotingizda ushbu dizayn resurslaridan foydalaning\nSHRIFTLAR va RANGLAR"
        assert PPTXProcessor._is_ad_slide(s_res, slide_index=3, total_slides=4) is True

        # Create a mock slide with Credits / Minnatdorchilik
        s_cred = prs.slides.add_slide(prs.slide_layouts[6])
        tb2 = s_cred.shapes.add_textbox(0, 0, 1000000, 1000000)
        tb2.text_frame.text = "Minnatdorchilik\nUshbu taqdimot shabloni quyidagilar sharofati bilan barcha uchun bepul:\nSlides Carnival taqdimot shabloni uchun\nPexels, Pixabay rasmlar uchun\nDizayn yaratishda omad yor bo'lsin!"
        assert PPTXProcessor._is_ad_slide(s_cred, slide_index=4, total_slides=4) is True

        # Create a normal content slide
        s_normal = prs.slides.add_slide(prs.slide_layouts[6])
        tb3 = s_normal.shapes.add_textbox(0, 0, 1000000, 1000000)
        tb3.text_frame.text = "Kompaniya Moliyaviy Ko'rsatkichlari\n2026-yil 3-chorak natijalari va daromadlar tahlili"
        assert PPTXProcessor._is_ad_slide(s_normal, slide_index=2, total_slides=4) is False

    def test_clean_full_presentation(self):
        from pptx import Presentation
        prs = Presentation()
        # Slide 1: Sarlavha
        s1 = prs.slides.add_slide(prs.slide_layouts[6])
        s1.shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.text = "Biznes Strategiya 2026"

        # Slide 2: Reja
        s2 = prs.slides.add_slide(prs.slide_layouts[6])
        s2.shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.text = "Asosiy vazifalar va maqsadlar"

        # Slide 3: Resurslar sahifasi (Reklama/shablon)
        s3 = prs.slides.add_slide(prs.slide_layouts[6])
        s3.shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.text = "Resurslar sahifasi\nCanva taqdimotingizda ushbu dizayn resurslaridan foydalaning\nSHRIFTLAR: ERICA ONE, DM SANS\nDIZAYN ELEMENTLARI va RANGLAR"

        # Slide 4: Minnatdorchilik (SlidesCarnival, Pexels, Pixabay)
        s4 = prs.slides.add_slide(prs.slide_layouts[6])
        s4.shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.text = "Minnatdorchilik\nUshbu taqdimot shabloni quyidagilar sharofati bilan barcha uchun bepul:\nSlides Carnival taqdimot shabloni uchun\nPexels, Pixabay rasmlar uchun\nDizayn yaratishda omad yor bo'lsin!"

        assert len(prs.slides) == 4
        removed = PPTXProcessor.clean_presentation_watermarks(prs)
        assert len(prs.slides) == 2, f"Kutilgan 2 slayd, amalda: {len(prs.slides)}"
        assert "Biznes Strategiya" in prs.slides[0].shapes[0].text_frame.text
        assert "Asosiy vazifalar" in prs.slides[1].shapes[0].text_frame.text

    def test_real_text(self):
        assert PPTXProcessor._is_real_text("") is False
        assert PPTXProcessor._is_real_text("   ") is False
        assert PPTXProcessor._is_real_text("Mavzu") is True

if __name__ == '__main__':
    unittest.main()
