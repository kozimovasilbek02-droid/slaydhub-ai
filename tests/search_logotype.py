import os
import glob
from pptx import Presentation

for pptx_file in glob.glob("output/**/*.pptx", recursive=True):
    try:
        prs = Presentation(pptx_file)
        for s_idx, s in enumerate(prs.slides):
            for shp in s.shapes:
                txt = shp.text_frame.text if shp.has_text_frame else ''
                if 'logotype' in txt.lower() or 'logo' in txt.lower():
                    print(f"File: {pptx_file} | Slide {s_idx+1} | Shape: {shp.name} | Text: '{txt}'")
        for m_idx, m in enumerate(prs.slide_masters):
            for shp in m.shapes:
                txt = shp.text_frame.text if shp.has_text_frame else ''
                if 'logotype' in txt.lower():
                    print(f"File: {pptx_file} | Master {m_idx} | Shape: {shp.name} | Text: '{txt}'")
            for l_idx, l in enumerate(m.slide_layouts):
                for shp in l.shapes:
                    txt = shp.text_frame.text if shp.has_text_frame else ''
                    if 'logotype' in txt.lower():
                        print(f"File: {pptx_file} | Layout {l_idx} | Shape: {shp.name} | Text: '{txt}'")
    except Exception as e:
        pass
